"""
tests/test_knitting_incentive.py
-----------------------------------
Engine `knitting.incentive` (người dùng chốt 2026-10-09):
    %Achieve = Σ(KNT N.W(kg) × Std.PTM) / Σ Available, tháng cộng dồn từ ngày 1 tới "As of",
    cả xưởng + từng Machine Group; Incentive = đơn giá bậc (From > , To <=) × Σ KNT N.W.

  1. Bậc đơn giá: ranh giới From (>) / To (<=), dưới bậc đầu / trên bậc cuối, validate khi lưu.
  2. File Piece Produced tháng 7 thật: %Achieve xưởng 74.66%, D-1F 79.19%, D-2F 72.84% -> 0 VND/kg; cuộn
     có Record End = ngày kết thúc file (01/08, trước 07:00) thuộc ngày sản xuất 31/07 (kẹp theo tên file).
  3. Dữ liệu dựng tay: cộng dồn MTD theo ngày, As of cắt kỳ, mỗi nhóm bậc riêng (dòng Xưởng không phải
     tổng tiền nhóm), lọc nhóm, chi tiết ngày, cuộn thiếu KNT N.W (NULL) được cảnh báo, Excel.

Chạy: python tests/test_knitting_incentive.py
"""
from __future__ import annotations

import io
import os
import sqlite3
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from modules.knitting.engines.excel_import import service as data_service  # noqa: E402
from modules.knitting.engines.incentive import service  # noqa: E402

PIECE_FIXTURE = Path(__file__).resolve().parent / "fixtures" / "sample_imports" / "CET-Piece Produced report - 01-07-2026 07;00;00 - 01-08-2026 07;00;00.csv"


def _check(label: str, actual, expected, failures: list[str]) -> None:
    ok = actual == expected
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}: actual={actual!r} expected={expected!r}")
    if not ok:
        failures.append(label)


def _expect_error(label: str, func, failures: list[str]) -> None:
    try:
        func()
        _check(label, "accepted", "ValueError", failures)
    except ValueError:
        _check(label, "ValueError", "ValueError", failures)


def _conn(path: str) -> sqlite3.Connection:
    conn = sqlite3.connect(path, isolation_level=None)
    conn.row_factory = sqlite3.Row
    return conn


def main() -> int:
    failures: list[str] = []
    os.environ.pop("DATABASE_URL", None)
    bands = [{"from_pct": a, "to_pct": b, "unit": u} for a, b, u in service.DEFAULT_BANDS]

    print("1. Bậc đơn giá")
    unit = lambda pct: service.band_for(pct, bands)["unit"]  # noqa: E731
    _check("0% -> 0", unit(0.0), 0, failures)
    _check("82.0% (<= To) -> 0", unit(82.0), 0, failures)
    _check("82.01% -> 15", unit(82.01), 15, failures)
    _check("84.5% -> 15", unit(84.5), 15, failures)
    _check("95.6% -> 130", unit(95.6), 130, failures)
    _check("100% -> 192", unit(100.0), 192, failures)
    _check("105% (trên bậc cuối) -> 192", unit(105.0), 192, failures)
    _check("khoảng cách tới bậc kế (90% -> 92%)", round(service.band_for(90.0, bands)["gap_to_next_pct"], 6), 2.0, failures)
    _check("bậc cuối không có bậc kế", service.band_for(99.0, bands)["next_from_pct"], None, failures)

    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    conn = _conn(tmp.name)
    try:
        _check("seed 9 bậc", len(service.load_bands(conn)), 9, failures)
        _expect_error("bậc không liên tục", lambda: service.save_bands(conn, [{"from_pct": 0, "to_pct": 80, "unit": 0}, {"from_pct": 81, "to_pct": 100, "unit": 10}], "t"), failures)
        _expect_error("From >= To", lambda: service.save_bands(conn, [{"from_pct": 50, "to_pct": 50, "unit": 0}], "t"), failures)
        _expect_error("đơn giá âm", lambda: service.save_bands(conn, [{"from_pct": 0, "to_pct": 100, "unit": -1}], "t"), failures)
        service.save_bands(conn, [{"from_pct": 50, "to_pct": 100, "unit": 9}, {"from_pct": 0, "to_pct": 50, "unit": 1}], "t")
        _check("lưu + sắp xếp lại", [(b["from_pct"], b["unit"]) for b in service.load_bands(conn)], [(0, 1), (50, 9)], failures)
    finally:
        conn.close()
        os.unlink(tmp.name)

    print("2. File Piece Produced tháng 7 thật")
    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    conn = _conn(tmp.name)
    try:
        data_service.import_piece_produced_file(conn, PIECE_FIXTURE.read_bytes(), PIECE_FIXTURE.name, "tester")
        f = service.normalize_filters(conn, None, None)
        _check("tháng mặc định + As of", (f["month"], f["from_date"], f["to_date"]), ("2026-07", "2026-07-01", "2026-07-31"), failures)
        rep = service.build_report(conn, f)
        groups = {g["name"]: round(g["achievement_pct"], 2) for g in rep["groups"]}
        _check("%Achieve xưởng", round(rep["workshop"]["achievement_pct"], 2), 74.66, failures)
        _check("%Achieve theo nhóm", groups, {"CETVN Block D-1F": 79.19, "CETVN Block D-2F": 72.84}, failures)
        _check("dưới 82% -> 0 VND", (rep["workshop"]["unit"], rep["workshop"]["incentive"]), (0, 0), failures)
        _check("KNT N.W tháng = cả file (87 cuộn sáng 01/08 thuộc ngày 31/07)", round(rep["workshop"]["kg"]), 262571, failures)
        _check("không có ngày tháng 8", rep["daily"][-1]["date"], "2026-07-31", failures)
        _check("không có cuộn thiếu cột", rep["missing_rolls"], 0, failures)
        _check("ngày 01/07: MTD = ngày", round(rep["daily"][0]["mtd_pct"], 1), 75.5, failures)
        _check("ngày 02/07: MTD cộng dồn", round(rep["daily"][1]["mtd_pct"], 1), 70.1, failures)
        _check("ngưỡng có thưởng", rep["threshold_pct"], 82.0, failures)
        _expect_error("As of ngoài tháng", lambda: service.normalize_filters(conn, "2026-07", "2026-08-01"), failures)
        _expect_error("tháng sai dạng", lambda: service.normalize_filters(conn, "07/2026", None), failures)
    finally:
        conn.close()
        os.unlink(tmp.name)

    print("3. Dữ liệu dựng tay")
    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    conn = _conn(tmp.name)
    try:
        service.ensure_tables(conn)
        rolls = [  # roll, group, machine, kg, std_ptm, available, record_end
            ("R1", "G1", "M1", 100, 2.0, 200, "2026-10-01"),   # G1 ngày 1: 200/200 = 100%
            ("R2", "G1", "M2", 100, 1.8, 200, "2026-10-02"),   # G1 cộng dồn: 380/400 = 95%
            ("R3", "G2", "M3", 100, 1.6, 200, "2026-10-01"),   # G2: 160/200 = 80%
            ("R4", "G2", "M3", 50, 2.0, 100, "2026-10-03"),    # G2 cộng dồn: 260/300 = 86.67%
            ("R5", "G2", "M4", None, None, 100, "2026-10-03"),  # import cũ thiếu cột -> chỉ cộng Available
            ("R6", "G1", "M1", 100, 2.0, 200, "2026-11-01"),   # tháng sau, không tính
        ]
        conn.executemany(
            "INSERT INTO knitting_piece_rolls (roll_no, machine_group, machine_code, greige_id, knt_nw_kg, std_ptm, available, record_start, record_end) VALUES (?, ?, ?, 'G', ?, ?, ?, ?, ?)",
            [(r, g, m, kg, p, a, d, d) for r, g, m, kg, p, a, d in rolls],
        )
        conn.commit()
        f = service.normalize_filters(conn, "2026-10", None)
        _check("As of = Record End mới nhất trong tháng", f["to_date"], "2026-10-03", failures)
        rep = service.build_report(conn, f)
        g = {x["name"]: x for x in rep["groups"]}
        _check("G1 95% -> 130 VND/kg × 200 kg", (round(g["G1"]["achievement_pct"], 2), g["G1"]["unit"], g["G1"]["incentive"]), (95.0, 130, 26000), failures)
        _check("G2 260/400 (cuộn thiếu cột vẫn cộng Available)", round(g["G2"]["achievement_pct"], 2), 65.0, failures)
        _check("G2 -> 0 VND", g["G2"]["incentive"], 0, failures)
        ws = rep["workshop"]
        _check("Xưởng 640/800 = 80% -> 0 (không phải tổng tiền nhóm)", (round(ws["achievement_pct"], 2), ws["incentive"]), (80.0, 0), failures)
        _check("cảnh báo cuộn thiếu cột", rep["missing_rolls"], 1, failures)
        _check("MTD theo ngày", [round(d["mtd_pct"], 2) for d in rep["daily"]], [90.0, 90.0, 80.0], failures)
        _check("MTD theo nhóm ngày 2", {k: round(v, 2) for k, v in rep["daily"][1]["mtd_pct_by_group"].items()}, {"G1": 95.0, "G2": 80.0}, failures)
        cut = service.build_report(conn, service.normalize_filters(conn, "2026-10", "2026-10-02"))
        _check("As of 02/10 cắt kỳ", (round(cut["workshop"]["achievement_pct"], 2), cut["workshop"]["unit"]), (round(540 / 600 * 100, 2), 80), failures)
        only_g1 = service.build_report(conn, service.normalize_filters(conn, "2026-10", None, "G1"))
        _check("lọc nhóm G1 -> Xưởng = G1", (round(only_g1["workshop"]["achievement_pct"], 2), [x["name"] for x in only_g1["groups"]]), (95.0, ["G1"]), failures)
        day = service.get_day_details(conn, f, "2026-10-01")
        _check("chi tiết ngày 01: 2 máy, 2 cuộn", (len(day["machines"]), len(day["rolls"])), (2, 2), failures)
        _check("chi tiết ngày 01: %Achieve", round(day["achievement_pct"], 2), 90.0, failures)
        _expect_error("ngày ngoài kỳ", lambda: service.get_day_details(conn, f, "2026-11-01"), failures)
        _check("máy thấp nhất đứng đầu", rep["machines"][0]["name"], "M4", failures)
        wb = load_workbook(io.BytesIO(service.export_excel(conn, f)))
        _check("Excel sheets", wb.sheetnames, ["Summary", "Daily", "Machines", "Bands", "Filters"], failures)
        summary = {r[0]: r for r in wb["Summary"].iter_rows(min_row=2, values_only=True)}
        _check("Excel G1 incentive", summary["G1"][6], 26000, failures)
    finally:
        conn.close()
        os.unlink(tmp.name)

    print()
    print("ALL PASS" if not failures else f"FAILED: {failures}")
    return 0 if not failures else 1


if __name__ == "__main__":
    sys.exit(main())

"""Right First Time (RFT) report — tính trực tiếp từ `batch_details` + 2 nguồn tra cứu nạp hằng
ngày (`dye_production_ops`, `dye_nc_reports` — xem `core/rft_sources_importer.py`).

Tập mẻ tính: mọi dòng `batch_details` có Dyelot kết thúc bằng "0" (mẻ gốc; mẻ làm lại đuôi
1/2/... không đếm lại mà phản ánh vào mẻ gốc qua ReworkCount). KHÔNG lọc ReDye. Mỗi lần chạy
(1 dòng batch_details) là 1 mẻ — 1 Dyelot chạy 2 lần đếm 2. Ngày = production_date theo EndTime
(mốc 07:00, fallback StartTime).

Các cột tính — đúng công thức Excel người dùng cung cấp (đã đối chiếu khớp 57/57 dòng file
`tests/fixtures/sample_imports/Copy of Batch_202696181325.xlsx`):
- STAGE: 2 ký tự đầu FormulaCode "01"/"07" -> Lab to Bulk; "08"/"09" -> Bulk to Bulk; còn lại
  (kể cả trống) -> 2nd Batch.
- MachineGroup: số sau ký tự đầu của cột MachineGroup (G600 -> 600) >= 500 -> ">=500kg", ngược lại
  "Small Machine"; không đọc được số -> không thuộc nhóm nào (Excel ra #VALUE!).
- NewBatch: Dyelot tăng ký tự cuối thêm 1 theo dãy 0-9, A-Z (C260708660 -> C260708661,
  …9 -> …A, …A -> …B — người dùng chốt, khác công thức Excel vốn ra "…10").
- DyeingRFT: NG nếu Dyelot có trong NC (đã lọc) HOẶC TotalCorrectionCnt > 0, ngược lại OK.
- ReworkCount: Rework nếu NewBatch có trong DG HOẶC Dyelot có trong NC; ngược lại Adjustment nếu
  TotalCorrectionCnt > 0; còn lại OK.
Lọc nguồn tra cứu (áp ở read time): DG = công đoạn bắt đầu "DG"; NC = công đoạn bắt đầu "DG",
Defect chứa "khác màu", Status Closed, Corrective không bắt đầu "MA".

5 tab: 3 tab Stage (rate = DyeingRFT OK / mẻ của Stage, Target là mức TỐI THIỂU) + 2 tab
Rework/Adjustment (CHỈ máy ">=500kg", rate = mẻ Rework hoặc Adjustment / mẻ máy >=500kg, Target
là mức TỐI ĐA).
"""
from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from core.brand_program_importer import ensure_brand_program_table
from core.database import execute_query, get_db, get_dialect
from core.production_time import get_production_date, production_bounds
from core.rft_sources_importer import ensure_rft_source_tables

# Slug dùng làm khoá URL/DOM — KHÔNG đổi sau khi đã có người dùng thao tác.
RFT_CATEGORY_SLUGS: dict[str, str] = {
    "lab_to_bulk": "Lab to Bulk",
    "bulk_to_bulk": "Bulk to Bulk",
    "second_batch": "2nd Batch",
    "rework": "Rework",
    "adjustment": "Adjustment",
}
RFT_CATEGORIES = tuple(RFT_CATEGORY_SLUGS.values())
STAGE_CATEGORIES: tuple[str, ...] = ("Lab to Bulk", "Bulk to Bulk", "2nd Batch")
REWORK_CATEGORIES: tuple[str, ...] = ("Rework", "Adjustment")

LARGE_MACHINE = ">=500kg"
SMALL_MACHINE = "Small Machine"
LARGE_MACHINE_MIN_KG = 500.0

# Dyeing Hub Dashboard đọc cờ này (từ thời scaffold) — luôn True vì đã có quy tắc thật.
RFT_CLASSIFICATION_READY = True

# Mỗi tab luôn chia đúng 3 loại vải chính (3 biểu đồ + 3 dòng bảng). Định nghĩa riêng ở đây,
# không import cross-engine (Vertical Slice — xem CLAUDE.md mục 3-4).
MAIN_FABRIC_TYPES: tuple[str, ...] = ("Cotton", "CVC", "Polyester")
_MAIN_FABRIC_TYPE_BY_NORM: dict[str, str] = {name.lower(): name for name in MAIN_FABRIC_TYPES}

_BRAND_PROGRAM_LABEL_SQL = "CASE WHEN COALESCE(bpm.brand, '') <> '' AND COALESCE(bpm.brand_program, '') <> '' THEN bpm.brand || ' - ' || bpm.brand_program ELSE '' END"


# ---------------------------------------------------------------------------
# Công thức từng cột (thuần, không đụng DB)
# ---------------------------------------------------------------------------


# Thứ tự ký tự cuối của Dyelot khi chạy lại mẻ: 0-9 rồi A-Z (A = 10, B = 11, ...).
_BATCH_SUFFIX_CHARS = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"


def classify_stage(formula_code: Any) -> str:
    prefix = str(formula_code or "").strip()[:2]
    if prefix in {"01", "07"}:
        return "Lab to Bulk"
    if prefix in {"08", "09"}:
        return "Bulk to Bulk"
    return "2nd Batch"


def machine_group_capacity(machine_group: Any) -> float | None:
    """Số kg trong MachineGroup (bỏ ký tự đầu): "G600" -> 600, "C1600" -> 1600."""
    text = str(machine_group or "").strip()
    try:
        return float(text[1:]) if len(text) > 1 else None
    except ValueError:
        return None


def machine_group_label(machine_group: Any) -> str | None:
    capacity = machine_group_capacity(machine_group)
    if capacity is None:
        return None
    return LARGE_MACHINE if capacity >= LARGE_MACHINE_MIN_KG else SMALL_MACHINE


def next_batch(dyelot: Any) -> str | None:
    """Dyelot tăng ký tự cuối thêm 1 theo dãy 0-9 rồi A-Z (A = 10, B = 11, ...): "…8" -> "…9",
    "…9" -> "…A", "…A" -> "…B" — độ dài Dyelot không đổi (người dùng chốt 2026-10-09, khác công
    thức Excel `RIGHT(J,1)+1` vốn ra "…10"). Ký tự cuối "Z" hoặc không phải chữ/số -> None."""
    text = str(dyelot or "").strip().upper()
    index = _BATCH_SUFFIX_CHARS.find(text[-1:]) if text else -1
    if index < 0 or index + 1 >= len(_BATCH_SUFFIX_CHARS):
        return None
    return text[:-1] + _BATCH_SUFFIX_CHARS[index + 1]


def _correction_count(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def dyeing_rft(in_nc: bool, total_correction_cnt: Any) -> str:
    return "NG" if in_nc or _correction_count(total_correction_cnt) > 0 else "OK"


def rework_count(new_batch_in_dg: bool, in_nc: bool, total_correction_cnt: Any) -> str:
    if new_batch_in_dg or in_nc:
        return "Rework"
    return "Adjustment" if _correction_count(total_correction_cnt) > 0 else "OK"


def _normalize_main_fabric_type(value: Any) -> str | None:
    return _MAIN_FABRIC_TYPE_BY_NORM.get(str(value or "").strip().lower())


def _norm_key(value: Any) -> str:
    return str(value or "").strip().lower()


# ---------------------------------------------------------------------------
# Đọc dữ liệu
# ---------------------------------------------------------------------------


def _dg_batches() -> set[str]:
    """Batch# (chuẩn hoá) đã đi qua 1 công đoạn nhuộm `DG*` trong Production Report."""
    rows = execute_query(
        "SELECT DISTINCT lower(trim(batch_no)) AS batch_key FROM dye_production_ops "
        "WHERE upper(trim(operation)) LIKE 'DG%'"
    )
    return {row["batch_key"] for row in rows if row["batch_key"]}


def is_counted_nc(row: Any) -> bool:
    """NC tính là Rework: công đoạn nhuộm `DG*`, NG khác màu, Closed, không phải thẻ MA."""
    route = str(row["operation_route"] or "").strip().upper()
    defect = str(row["defect"] or "").strip().lower()
    status = str(row["status"] or "").strip().lower()
    corrective = str(row["corrective"] or "").strip().upper()
    return route.startswith("DG") and "khác màu" in defect and status == "closed" and not corrective.startswith("MA")


def _nc_by_batch() -> dict[str, str]:
    """{Batch Ref chuẩn hoá: NC#} của các NC được tính (`is_counted_nc`)."""
    rows = execute_query(
        "SELECT nc_no, batch_ref, defect, operation_route, status, corrective FROM dye_nc_reports ORDER BY nc_no"
    )
    result: dict[str, str] = {}
    for row in rows:
        key = _norm_key(row["batch_ref"])
        if key and is_counted_nc(row) and key not in result:
            result[key] = str(row["nc_no"])
    return result


def _parse_datetime(value: Any) -> datetime | None:
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(text[:19], fmt)
        except ValueError:
            continue
    return None


def load_rft_rows(from_date: str | None = None, to_date: str | None = None) -> list[dict[str, Any]]:
    """Mọi lần chạy (Dyelot đuôi "0") có production_date trong [from_date, to_date], kèm 5 cột
    tính + nhãn Brand Program. Chưa áp filter Capacity/Machine Group/Brand Program (để tính
    danh sách lựa chọn từ CÙNG 1 lần đọc)."""
    conn = get_db()
    ensure_brand_program_table(conn)
    ensure_rft_source_tables(conn)
    conn.commit()

    record_time_sql = "COALESCE(NULLIF(b.end_time, ''), b.start_time)"
    sql = f"""
        SELECT b.id AS batch_id, b.dyelot, b.machine, b.machine_group, b.formula_code,
               b.total_correction_cnt, b.fabric_type, b.greige_code, b.customer,
               b.start_time, b.end_time, {_BRAND_PROGRAM_LABEL_SQL} AS brand_program
        FROM batch_details b
        LEFT JOIN brand_program_mapping bpm ON lower(trim(bpm.greige_code)) = lower(trim(b.greige_code))
        WHERE trim(b.dyelot) LIKE '%0'
    """
    params: list[Any] = []
    start, end = production_bounds(from_date, to_date)
    if start:
        sql += f" AND {record_time_sql} >= ?"
        params.append(start)
    if end:
        sql += f" AND {record_time_sql} < ?"
        params.append(end)
    sql += " ORDER BY b.end_time, b.dyelot"
    rows = execute_query(sql, params)

    dg_batches = _dg_batches()
    nc_by_batch = _nc_by_batch()
    result: list[dict[str, Any]] = []
    for row in rows:
        record_time = _parse_datetime(row["end_time"]) or _parse_datetime(row["start_time"])
        if record_time is None:
            continue
        dyelot = str(row["dyelot"]).strip()
        new_batch = next_batch(dyelot)
        nc_no = nc_by_batch.get(_norm_key(dyelot))
        new_batch_in_dg = new_batch is not None and _norm_key(new_batch) in dg_batches
        sources = (["DG"] if new_batch_in_dg else []) + ([nc_no] if nc_no else [])
        result.append({
            **dict(row),
            "dyelot": dyelot,
            "production_date": get_production_date(record_time),
            "stage": classify_stage(row["formula_code"]),
            "capacity": machine_group_capacity(row["machine_group"]),
            "machine_group_label": machine_group_label(row["machine_group"]),
            "new_batch": new_batch,
            "dyeing_rft": dyeing_rft(nc_no is not None, row["total_correction_cnt"]),
            "rework_count": rework_count(new_batch_in_dg, nc_no is not None, row["total_correction_cnt"]),
            "rework_source": " + ".join(sources),
            "main_fabric_type": _normalize_main_fabric_type(row["fabric_type"]),
        })
    return result


# ---------------------------------------------------------------------------
# Filter
# ---------------------------------------------------------------------------


def _parse_list(value: str | list[str] | None) -> list[str]:
    """CSV hoặc list; rỗng/'ALL' nghĩa là không lọc theo chiều đó."""
    if not value:
        return []
    values = value if isinstance(value, list) else str(value).split(",")
    result: list[str] = []
    for item in values:
        text = str(item).strip()
        if not text or text.upper() == "ALL":
            return []
        if text not in result:
            result.append(text)
    return result


def _parse_capacities(value: str | list[str] | None) -> list[float]:
    result: list[float] = []
    for text in _parse_list(value):
        try:
            number = float(text)
        except ValueError:
            continue
        if number not in result:
            result.append(number)
    return result


def _format_capacity(value: float) -> str:
    return str(int(value)) if float(value).is_integer() else str(value)


def _filtered_rows(
    capacities: str | list[str] | None, machine_groups: str | list[str] | None,
    brand_programs: str | list[str] | None, from_date: str | None, to_date: str | None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    selected_capacities = _parse_capacities(capacities)
    selected_groups = _parse_list(machine_groups)
    selected_brands = _parse_list(brand_programs)
    all_rows = load_rft_rows(from_date, to_date)
    options = {
        "available_capacities": [_format_capacity(value) for value in sorted({row["capacity"] for row in all_rows if row["capacity"] is not None})],
        "available_machine_groups": [label for label in (LARGE_MACHINE, SMALL_MACHINE) if any(row["machine_group_label"] == label for row in all_rows)],
        "available_brand_programs": sorted({row["brand_program"] for row in all_rows if row["brand_program"]}),
    }
    rows = [
        row for row in all_rows
        if (not selected_capacities or row["capacity"] in selected_capacities)
        and (not selected_groups or row["machine_group_label"] in selected_groups)
        and (not selected_brands or row["brand_program"] in selected_brands)
    ]
    filters = {
        "capacities": [_format_capacity(value) for value in selected_capacities] or "all",
        "machine_groups": selected_groups or "all",
        "brand_programs": selected_brands or "all",
        "from_date": from_date, "to_date": to_date,
    }
    return rows, {**options, "filters": filters}


def _tab_rows(category: str, rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], Any]:
    """(mẻ thuộc mẫu số của tab, hàm xác định mẻ thuộc tử số)."""
    if category in STAGE_CATEGORIES:
        return [row for row in rows if row["stage"] == category], (lambda row: row["dyeing_rft"] == "OK")
    return [row for row in rows if row["machine_group_label"] == LARGE_MACHINE], (lambda row: row["rework_count"] == category)


def target_direction(category: str) -> str:
    """"min": Target là mức tối thiểu (RFT, cao là tốt); "max": mức tối đa (Rework/Adjustment)."""
    return "max" if category in REWORK_CATEGORIES else "min"


# ---------------------------------------------------------------------------
# Target — `rft_targets`, khoá ghép (category, fabric_type)
# ---------------------------------------------------------------------------


def _ensure_targets_table(conn: Any) -> None:
    """CHỈ chạy CREATE TABLE ở SQLite — Postgres tạo qua `supabase/schema.sql`."""
    if get_dialect() == "sqlite":
        conn.execute("""
            CREATE TABLE IF NOT EXISTS rft_targets (
                category TEXT NOT NULL,
                fabric_type TEXT NOT NULL,
                target_value REAL NOT NULL DEFAULT 0,
                updated_at TEXT NOT NULL DEFAULT (datetime('now')),
                PRIMARY KEY (category, fabric_type)
            )
        """)
    conn.commit()


def get_targets(category: str) -> dict[str, float | None]:
    """{fabric_type: target} cho 3 loại vải chính; chưa cấu hình -> None (UI hiện "-")."""
    conn = get_db()
    _ensure_targets_table(conn)
    rows = execute_query("SELECT fabric_type, target_value FROM rft_targets WHERE category = ?", [category])
    result: dict[str, float | None] = {row["fabric_type"]: float(row["target_value"]) for row in rows}
    return {name: result.get(name) for name in MAIN_FABRIC_TYPES}


def set_target(category: str, fabric_type: str, target_value: float) -> dict[str, Any]:
    if category not in RFT_CATEGORIES:
        raise ValueError(f"Nhóm RFT không hợp lệ: {category}")
    if fabric_type not in MAIN_FABRIC_TYPES:
        raise ValueError(f"Fabric type không hợp lệ: {fabric_type}")
    if not 0 <= target_value <= 100:
        raise ValueError("Target phải nằm trong khoảng 0-100%.")
    conn = get_db()
    _ensure_targets_table(conn)
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn.execute(
        "INSERT INTO rft_targets (category, fabric_type, target_value, updated_at) VALUES (?, ?, ?, ?) "
        "ON CONFLICT(category, fabric_type) DO UPDATE SET target_value = excluded.target_value, updated_at = excluded.updated_at",
        (category, fabric_type, target_value, now_str),
    )
    conn.commit()
    return {"category": category, "fabric_type": fabric_type, "target_value": target_value, "updated_at": now_str}


# ---------------------------------------------------------------------------
# Pivot cho 1 tab
# ---------------------------------------------------------------------------


def _period(day: date, group_by: str) -> tuple[str, str]:
    if group_by == "month":
        return day.strftime("%Y-%m"), day.strftime("%b-%y")
    if group_by == "week":
        year, week, _ = day.isocalendar()
        return f"{year}-W{week:02d}", f"W{week:02d} {year}"
    return day.isoformat(), day.strftime("%d %b %Y")


def _rate(hit: int, total: int) -> float | None:
    return round(hit / total * 100, 1) if total else None


def _rate_line(rows: list[dict[str, Any]], is_hit: Any, period_keys: list[str], group_by: str) -> dict[str, Any]:
    counts: dict[str, list[int]] = {}
    for row in rows:
        bucket = counts.setdefault(_period(row["production_date"], group_by)[0], [0, 0])
        bucket[0] += 1
        bucket[1] += 1 if is_hit(row) else 0
    hits = sum(1 for row in rows if is_hit(row))
    return {
        "values": [_rate(counts[key][1], counts[key][0]) if key in counts else None for key in period_keys],
        "batches": [counts[key][0] if key in counts else 0 for key in period_keys],
        "total": _rate(hits, len(rows)),
        "total_batches": len(rows),
        "hit_batches": hits,
    }


def get_rft_pivot_data(
    category: str, capacities: str | list[str] | None = None,
    machine_groups: str | list[str] | None = None,
    brand_programs: str | list[str] | None = None,
    from_date: str | None = None, to_date: str | None = None, group_by: str = "date",
) -> dict[str, Any]:
    """Bảng + biểu đồ cho 1 tab theo Day/Week/Month. KPI và `chart.rate_values` (Dyeing Hub
    Dashboard đọc) gộp MỌI loại vải của tab; `rows` chia 3 loại vải chính + `total_row` gộp
    mọi loại vải (bằng KPI)."""
    if category not in RFT_CATEGORIES:
        raise ValueError(f"Nhóm RFT không hợp lệ: {category}")
    group_by = group_by if group_by in {"date", "week", "month"} else "date"
    rows, meta = _filtered_rows(capacities, machine_groups, brand_programs, from_date, to_date)
    tab_rows, is_hit = _tab_rows(category, rows)
    targets = get_targets(category)

    periods = dict(sorted({_period(row["production_date"], group_by) for row in tab_rows}))
    period_keys = list(periods)
    labels = [periods[key] for key in period_keys]

    total_line = _rate_line(tab_rows, is_hit, period_keys, group_by)
    fabric_rows = []
    for name in MAIN_FABRIC_TYPES:
        line = _rate_line([row for row in tab_rows if row["main_fabric_type"] == name], is_hit, period_keys, group_by)
        fabric_rows.append({"fabric_type": name, "target": targets.get(name), **line})

    return {
        "category": category,
        "target_direction": target_direction(category),
        "filters": {**meta["filters"], "group_by": group_by},
        "periods": labels,
        "period_keys": period_keys,
        "rows": fabric_rows,
        "total_row": {"fabric_type": "Total", **total_line},
        "chart": {
            "categories": labels,
            "values": total_line["batches"],
            "rate_values": [value if value is not None else 0.0 for value in total_line["values"]],
            "series": [{"name": row["fabric_type"], "data": row["values"]} for row in fabric_rows],
        },
        "kpis": {
            "total_batches": total_line["total_batches"],
            "hit_batches": total_line["hit_batches"],
            "rate_pct": total_line["total"] if total_line["total"] is not None else 0.0,
        },
        "unknown_machine_group_count": sum(1 for row in rows if row["machine_group_label"] is None),
        "available_capacities": meta["available_capacities"],
        "available_machine_groups": meta["available_machine_groups"],
        "available_brand_programs": meta["available_brand_programs"],
        "classification_ready": RFT_CLASSIFICATION_READY,
    }


# ---------------------------------------------------------------------------
# Export Excel — 2 pivot đúng bố cục file đối chiếu của người dùng + dữ liệu từng mẻ
# ---------------------------------------------------------------------------

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="24292F", end_color="24292F", fill_type="solid")
_BOLD = Font(bold=True)
_PCT = "0.0%"

DATA_COLUMNS: tuple[tuple[str, str], ...] = (
    ("Production Date", "production_date"), ("Dyelot", "dyelot"), ("Machine", "machine"),
    ("MachineGroup", "machine_group"), ("FabricType", "fabric_type"), ("Brand Program", "brand_program"),
    ("Customer", "customer"), ("FormulaCode", "formula_code"), ("TotalCorrectionCnt", "total_correction_cnt"),
    ("StartTime", "start_time"), ("EndTime", "end_time"), ("STAGE", "stage"),
    ("MachineGroup2", "machine_group_label"), ("DyeingRFT", "dyeing_rft"), ("NewBatch", "new_batch"),
    ("ReworkCount", "rework_count"), ("Rework Source", "rework_source"),
)


def _write_header(sheet: Any, row: int, values: list[str]) -> None:
    for col, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=col, value=value)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL


def _fabric_groups(rows: list[dict[str, Any]]) -> list[tuple[str, list[dict[str, Any]]]]:
    groups = [(name, [row for row in rows if row["main_fabric_type"] == name]) for name in MAIN_FABRIC_TYPES]
    others = [row for row in rows if row["main_fabric_type"] is None]
    if others:
        groups.append(("Other", others))
    return groups + [("Grand Total", rows)]


def _write_rft_pivot(sheet: Any, rows: list[dict[str, Any]]) -> None:
    """Pivot 1: loại vải x Stage — Batches, OK %, NG % (trên mẻ của Stage), Share % (Stage trên
    tổng mẻ của loại vải)."""
    sheet.cell(row=1, column=1, value="Dyeing RFT — OK % / NG % of batches in the stage; Share % = stage batches / fabric batches").font = _BOLD
    header = ["Fabric Type"]
    for stage in STAGE_CATEGORIES:
        header += [f"{stage} Batches", f"{stage} OK %", f"{stage} NG %", f"{stage} Share %"]
    _write_header(sheet, 3, header)
    for row_idx, (name, group) in enumerate(_fabric_groups(rows), start=4):
        sheet.cell(row=row_idx, column=1, value=name).font = _BOLD if name == "Grand Total" else Font()
        col = 2
        for stage in STAGE_CATEGORIES:
            stage_rows = [row for row in group if row["stage"] == stage]
            ok = sum(1 for row in stage_rows if row["dyeing_rft"] == "OK")
            values = [
                len(stage_rows),
                ok / len(stage_rows) if stage_rows else None,
                (len(stage_rows) - ok) / len(stage_rows) if stage_rows else None,
                len(stage_rows) / len(group) if group else None,
            ]
            for offset, value in enumerate(values):
                cell = sheet.cell(row=row_idx, column=col + offset, value=value)
                if offset:
                    cell.number_format = _PCT
            col += 4
    sheet.column_dimensions["A"].width = 14
    for index in range(2, len(header) + 1):
        sheet.column_dimensions[sheet.cell(row=3, column=index).column_letter].width = 16


def _write_rework_pivot(sheet: Any, rows: list[dict[str, Any]]) -> None:
    """Pivot 2: chỉ máy >=500kg — loại vải x ReworkCount, số mẻ + % trên TỔNG mẻ (như pivot gốc)."""
    large = [row for row in rows if row["machine_group_label"] == LARGE_MACHINE]
    sheet.cell(row=1, column=1, value="ReworkCount — MachineGroup >=500kg; % = batches / grand total batches").font = _BOLD
    results = ("OK", "Adjustment", "Rework")
    header = ["Fabric Type"] + [f"{name} Batches" for name in results] + ["Total Batches"] + [f"{name} %" for name in results] + ["Total %"]
    _write_header(sheet, 3, header)
    grand = len(large)
    for row_idx, (name, group) in enumerate(_fabric_groups(large), start=4):
        sheet.cell(row=row_idx, column=1, value=name).font = _BOLD if name == "Grand Total" else Font()
        counts = [sum(1 for row in group if row["rework_count"] == result) for result in results] + [len(group)]
        for offset, count in enumerate(counts):
            sheet.cell(row=row_idx, column=2 + offset, value=count)
            pct = sheet.cell(row=row_idx, column=2 + len(counts) + offset, value=count / grand if grand else None)
            pct.number_format = _PCT
    sheet.column_dimensions["A"].width = 14
    for index in range(2, len(header) + 1):
        sheet.column_dimensions[sheet.cell(row=3, column=index).column_letter].width = 16


def export_rft_excel(
    capacities: str | list[str] | None = None, machine_groups: str | list[str] | None = None,
    brand_programs: str | list[str] | None = None, from_date: str | None = None, to_date: str | None = None,
) -> tuple[bytes, str]:
    rows, meta = _filtered_rows(capacities, machine_groups, brand_programs, from_date, to_date)
    workbook = Workbook()
    _write_rft_pivot(workbook.active, rows)
    workbook.active.title = "Dyeing RFT"
    _write_rework_pivot(workbook.create_sheet("Rework Count"), rows)

    data_sheet = workbook.create_sheet("Data")
    _write_header(data_sheet, 1, [header for header, _ in DATA_COLUMNS])
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (_, field) in enumerate(DATA_COLUMNS, start=1):
            value = row.get(field)
            data_sheet.cell(row=row_idx, column=col_idx, value=value.isoformat() if isinstance(value, date) else value)
    data_sheet.auto_filter.ref = data_sheet.dimensions
    data_sheet.freeze_panes = "A2"
    for col_idx, (header, _) in enumerate(DATA_COLUMNS, start=1):
        data_sheet.column_dimensions[data_sheet.cell(row=1, column=col_idx).column_letter].width = max(12, len(header) + 2)

    filter_sheet = workbook.create_sheet("Filters")
    _write_header(filter_sheet, 1, ["Filter", "Value"])
    filters = meta["filters"]
    for row_idx, (label, value) in enumerate((
        ("From Date", filters["from_date"] or "All"), ("To Date", filters["to_date"] or "All"),
        ("Capacity (Kg)", filters["capacities"]), ("Machine Group", filters["machine_groups"]),
        ("Brand Program", filters["brand_programs"]), ("Batches", len(rows)),
        ("Rule", "Dyelot ending with 0; production date by EndTime (07:00 cut-off); each run counts as 1 batch"),
    ), start=2):
        filter_sheet.cell(row=row_idx, column=1, value=label)
        filter_sheet.cell(row=row_idx, column=2, value=", ".join(value) if isinstance(value, list) else value)
    filter_sheet.column_dimensions["A"].width = 18
    filter_sheet.column_dimensions["B"].width = 60

    buffer = io.BytesIO()
    workbook.save(buffer)
    period = f"{from_date or 'all'}_to_{to_date or 'all'}"
    return buffer.getvalue(), f"rft_{period}.xlsx"

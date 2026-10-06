"""
tests/test_idle_time.py
-------------------------
2026-10-05 — Engine `idle_time` (báo cáo Idle Time + Idle Entry). Kiểm tra trên DB tạm, dữ liệu
dựng tay:

  1. Ngày không trọn / mẻ vắt 07:00 / máy không chạy / ngày không có dữ liệu bị bỏ.
  2. Filter Tank (J/O) + Capacity theo Machine Master.
  3. Note bản 1 (`idle_time_notes`) được chuyển thành lần dừng source='report'.
  4. Ghép lần dừng vào khoảng idle (quy tắc người dùng chốt):
     - giờ nhập làm tròn chỉ là cơ sở -> CẢ khoảng idle (giờ chính xác từ Batch) nhận nguyên nhân;
     - khoảng idle vắt qua 07:00 ghép trên khoảng liên tục -> cả 2 ngày nhận nguyên nhân;
     - nhiều lần dừng trong 1 khoảng idle -> chia tại mốc của lần sau;
     - chồng giờ cùng máy -> lần cập nhật sau cùng thắng (overridden / partially_overridden);
     - chưa có Batch -> pending, upload Batch xong -> matched đúng giờ thật;
     - không giao khoảng nào -> khoảng gần nhất (không giới hạn độ lệch).
  5. Save trong drawer báo cáo sửa lại đúng lần dừng cùng giờ (không nhân bản); validate; xoá.
  6. Target + Excel export.

Chạy: python tests/test_idle_time.py
"""
from __future__ import annotations

import io
import os
import shutil
import sqlite3
import sys
import tempfile
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from flask import Flask  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from core.database import close_db, get_db  # noqa: E402
from modules.dyeing.engines.batch_matrix import batch_day_trend  # noqa: E402
from modules.dyeing.engines.idle_time import service  # noqa: E402

# (dyelot, machine, start, end)
BATCHES = [
    ("C100000", "M1", "2026-09-01 18:00:00", "2026-09-01 22:00:00"),
    ("C100010", "M1", "2026-09-01 23:00:00", "2026-09-02 09:00:00"),  # vắt qua 07:00
    ("C100020", "M1", "2026-09-02 10:00:00", "2026-09-03 07:00:00"),  # kết thúc đúng 07:00
    ("C100030", "M2", "2026-09-02 08:00:00", "2026-09-02 20:00:00"),
    ("C100040", "M1", "2026-09-05 08:00:00", "2026-09-05 10:00:00"),  # 03-04/09 không có dữ liệu
]
DAYS = [date(2026, 9, d) for d in range(1, 7)]


def _check(label: str, actual, expected, failures: list[str]) -> None:
    ok = actual == expected
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}: actual={actual!r} expected={expected!r}")
    if not ok:
        failures.append(label)


def _expect_error(label: str, func, failures: list[str]) -> None:
    try:
        func()
        _check(label, "accepted", "ValueError", failures)
    except ValueError:
        _check(label, "ValueError", "ValueError", failures)


def _init_db(db_path: str) -> None:
    conn = sqlite3.connect(db_path)
    conn.execute("""
        CREATE TABLE batch_details (
            id INTEGER PRIMARY KEY AUTOINCREMENT, dyelot TEXT NOT NULL, sap_lot TEXT, redye REAL DEFAULT 0,
            shade TEXT, colour_no TEXT, recipe_no TEXT, customer_color TEXT, batch_type TEXT,
            is_rework INTEGER DEFAULT 0, run_time REAL DEFAULT 0, greige_code TEXT,
            machine TEXT, fabric_type TEXT, start_time TEXT, end_time TEXT
        )
    """)
    conn.execute("CREATE TABLE machines (machine_id TEXT, machine_code TEXT, tank_type TEXT, capacity_kg REAL, domain TEXT)")
    conn.execute("CREATE TABLE users (id INTEGER PRIMARY KEY, username TEXT)")
    conn.execute("INSERT INTO users (id, username) VALUES (1, 'tester')")
    conn.executemany(
        "INSERT INTO machines (machine_id, machine_code, tank_type, capacity_kg, domain) VALUES (?, ?, ?, ?, 'dyeing')",
        [("M1", "M1", "J tank", 600), ("M2", "M2", "O tank", 1200)],
    )
    conn.executemany(
        "INSERT INTO batch_details (dyelot, machine, start_time, end_time, fabric_type, sap_lot) VALUES (?, ?, ?, ?, 'Cotton', '1000000001')",
        BATCHES,
    )
    # Note bản 1 (trước khi có idle_time_stops) — phải được chuyển sang lần dừng source='report'.
    conn.execute("""
        CREATE TABLE idle_time_notes (
            id INTEGER PRIMARY KEY AUTOINCREMENT, machine TEXT NOT NULL, production_date TEXT NOT NULL,
            gap_start TEXT NOT NULL, gap_end TEXT NOT NULL, reason TEXT, detail TEXT,
            updated_by INTEGER NOT NULL, updated_at TEXT NOT NULL, UNIQUE (machine, gap_start)
        )
    """)
    conn.execute(
        "INSERT INTO idle_time_notes (machine, production_date, gap_start, gap_end, reason, detail, updated_by, updated_at) "
        "VALUES ('M2', '2026-09-05', '2026-09-05 08:00:00', '2026-09-05 10:00:00', 'Maintenance', 'legacy', 1, '2026-10-01 00:00:00')"
    )
    conn.commit()
    conn.close()


def _recompute() -> None:
    conn = get_db()
    batch_day_trend.recompute_all(DAYS, conn)
    conn.commit()


def _rows(data: dict) -> dict[str, dict]:
    return {row["machine"]: row for row in data["rows"]}


def _add(machine: str, start: str, end: str | None, reason: str, order: int) -> int:
    """Tạo lần dừng Idle Entry rồi ép `updated_at` để kiểm soát thứ tự "sau cùng thắng"."""
    stop = service.create_stop(machine, start, end, reason, None, 1)
    conn = get_db()
    conn.execute("UPDATE idle_time_stops SET updated_at = ? WHERE id = ?", (f"2026-10-02 00:00:{order:02d}", stop["id"]))
    conn.commit()
    return stop["id"]


def _segments(machine: str, day: str) -> list[tuple]:
    return [
        (seg["start"][5:16], seg["end"][5:16], seg["stop"]["reason"] if seg["stop"] else None)
        for gap in service.get_idle_gaps(machine, day)["gaps"] for seg in gap["segments"]
    ]


def _status(from_date: str, to_date: str) -> dict[int, dict]:
    return {stop["id"]: stop for stop in service.list_stops(from_date, to_date)}


def _scenario_matrix(failures: list[str]) -> None:
    print("\n=== Kịch bản 1: ngày không trọn, mẻ vắt 07:00, máy không chạy ===")
    data = service.get_idle_matrix("2026-09-01", "2026-09-05")
    _check("ngày có dữ liệu (bỏ 03/09, 04/09)", data["days"], ["2026-09-01", "2026-09-02", "2026-09-05"], failures)
    _check("ngày không trọn", data["partial_days"], ["2026-09-01", "2026-09-05"], failures)
    _check("giờ có dữ liệu/ngày", data["day_available_hours"], {"2026-09-01": 13.0, "2026-09-02": 24.0, "2026-09-05": 2.0}, failures)
    rows = _rows(data)
    m1, m2 = rows["M1"]["cells"], rows["M2"]["cells"]
    _check("M1 01/09: chạy 4h + 8h, idle 22:00-23:00", (m1["2026-09-01"]["operating"], m1["2026-09-01"]["idle"]), (12.0, 1.0), failures)
    _check("M1 02/09: chạy 2h + 21h, idle 09:00-10:00", (m1["2026-09-02"]["operating"], m1["2026-09-02"]["idle"]), (23.0, 1.0), failures)
    _check("M2 01/09 không chạy mẻ -> idle trọn 13h", (m2["2026-09-01"]["operating"], m2["2026-09-01"]["idle"], m2["2026-09-01"]["idle_pct"]), (0.0, 13.0, 100.0), failures)
    _check("M2 02/09: 2 khoảng idle, 12h", (m2["2026-09-02"]["idle"], m2["2026-09-02"]["gap_count"]), (12.0, 2), failures)
    bad = [(row["machine"], day) for row in data["rows"] for day, cell in row["cells"].items() if abs(cell["operating"] + cell["idle"] - cell["available"]) > 0.01]
    _check("operating + idle = available ở mọi ô", bad, [], failures)
    _check("% Idle M1 cả kỳ = 2/39", rows["M1"]["idle_pct"], round(2 / 39 * 100, 2), failures)
    _check("tổng ngày 02/09", {k: data["day_totals"]["2026-09-02"][k] for k in ("operating", "idle", "available")}, {"operating": 35.0, "idle": 13.0, "available": 48.0}, failures)
    gaps = service.get_idle_gaps("M2", "2026-09-02")["gaps"]
    _check("khoảng idle M2 02/09 (khoảng đầu nối từ 01/09)", [(g["start"], g["end"], g["gap_start"], g["prev_batch"], g["next_batch"]) for g in gaps], [
        ("2026-09-02 07:00:00", "2026-09-02 08:00:00", "2026-09-01 18:00:00", None, "C100030"),
        ("2026-09-02 20:00:00", "2026-09-03 07:00:00", "2026-09-02 20:00:00", "C100030", None),
    ], failures)


def _scenario_filters(failures: list[str]) -> None:
    print("\n=== Kịch bản 2: filter Tank + Capacity ===")
    data = service.get_idle_matrix("2026-09-01", "2026-09-05")
    _check("available_tank_types", data["available_tank_types"], ["J tank", "O tank"], failures)
    _check("available_capacities", data["available_capacities"], [600.0, 1200.0], failures)
    _check("J tank -> M1", list(_rows(service.get_idle_matrix("2026-09-01", "2026-09-05", tank_types=["J tank"]))), ["M1"], failures)
    _check("Capacity 1200 -> M2", list(_rows(service.get_idle_matrix("2026-09-01", "2026-09-05", capacities=["1200"]))), ["M2"], failures)
    filtered = service.get_idle_matrix("2026-09-01", "2026-09-05", tank_types=["O tank"])
    _check("lọc không đổi giờ có dữ liệu của ngày", filtered["day_available_hours"]["2026-09-01"], 13.0, failures)


def _scenario_stops(failures: list[str]) -> None:
    print("\n=== Kịch bản 3: note bản 1 được chuyển thành lần dừng ===")
    _check("note cũ -> lần dừng 'report' trên M2 05/09", _segments("M2", "2026-09-05"), [("09-05 08:00", "09-05 10:00", "Maintenance")], failures)

    print("\n=== Kịch bản 4: ghép lần dừng vào khoảng idle ===")
    a = _add("M1", "2026-09-01 22:15", "2026-09-01 22:30", "Waiting for fabric", 1)
    _check("giờ nhập 22:15-22:30 -> CẢ khoảng idle thật 22:00-23:00", _segments("M1", "2026-09-01"), [("09-01 22:00", "09-01 23:00", "Waiting for fabric")], failures)

    _add("M2", "2026-09-02 05:00", "2026-09-02 06:00", "Maintenance", 2)
    _check("khoảng vắt 07:00: phần 01/09", _segments("M2", "2026-09-01"), [("09-01 18:00", "09-02 07:00", "Maintenance")], failures)
    _check("khoảng vắt 07:00: phần 02/09 cũng nhận nguyên nhân", _segments("M2", "2026-09-02")[0], ("09-02 07:00", "09-02 08:00", "Maintenance"), failures)

    _add("M2", "2026-09-02 20:30", "2026-09-02 22:00", "No order", 3)
    _add("M2", "2026-09-02 22:00", "2026-09-02 23:00", "Machine breakdown", 4)
    _check("2 lần dừng trong 1 khoảng -> chia tại 22:00, kéo về 2 đầu khoảng", _segments("M2", "2026-09-02")[1:], [
        ("09-02 20:00", "09-02 22:00", "No order"), ("09-02 22:00", "09-03 07:00", "Machine breakdown"),
    ], failures)

    x = _add("M1", "2026-09-02 08:30", "2026-09-02 09:30", "Waiting for water", 5)
    z = _add("M1", "2026-09-02 09:20", "2026-09-02 09:40", "Waiting for steam", 6)
    y = _add("M1", "2026-09-02 09:15", "2026-09-02 10:00", "Maintenance", 7)
    _check("chồng giờ: lần sau cùng (09:15-10:00) thắng phần chồng", _segments("M1", "2026-09-02"), [
        ("09-02 09:00", "09-02 09:15", "Waiting for water"), ("09-02 09:15", "09-02 10:00", "Maintenance"),
    ], failures)
    status = _status("2026-09-01", "2026-09-02")
    _check("trạng thái lần dừng", (status[a]["status"], status[a]["deviation_minutes"], status[x]["status"], status[x]["partially_overridden"], status[z]["status"]),
           ("matched", 0.0, "matched", True, "overridden"), failures)
    _check("lần dừng A: giờ thật đã gán", status[a]["matched"], [{"start": "2026-09-01 22:00:00", "end": "2026-09-01 23:00:00", "hours": 1.0}], failures)

    m2 = _rows(service.get_idle_matrix("2026-09-01", "2026-09-05"))["M2"]["cells"]
    _check("ô M2 01/09 + 02/09: toàn bộ idle đã có nguyên nhân", (m2["2026-09-01"]["explained"], m2["2026-09-02"]["explained"]), (13.0, 12.0), failures)

    f = _add("M1", "2026-09-05 08:30", "2026-09-05 09:00", "No order", 8)
    _check("không có khoảng idle nào của máy trong dữ liệu -> unmatched", _status("2026-09-05", "2026-09-05")[f]["status"], "unmatched", failures)
    e = _add("M1", "2026-09-06 10:00", "2026-09-06 11:00", "Waiting for previous process", 9)
    _check("chưa có Batch 06/09 -> pending", _status("2026-09-06", "2026-09-06")[e]["status"], "pending", failures)

    conn = get_db()
    conn.executemany(
        "INSERT INTO batch_details (dyelot, machine, start_time, end_time, fabric_type, sap_lot) VALUES (?, 'M1', ?, ?, 'Cotton', '1000000001')",
        [("C100060", "2026-09-06 08:00:17", "2026-09-06 10:31:42"), ("C100070", "2026-09-06 12:04:05", "2026-09-06 13:00:00")],
    )
    conn.commit()
    _recompute()
    after = _status("2026-09-06", "2026-09-06")[e]
    _check("upload Batch -> matched đúng giờ thật đến từng giây", (after["status"], after["matched"]),
           ("matched", [{"start": "2026-09-06 10:31:42", "end": "2026-09-06 12:04:05", "hours": round((3600 + 32 * 60 + 23) / 3600, 2)}]), failures)
    _check("lần dừng 05/09 gán khoảng gần nhất nhưng bị lần sau chiếm hết -> no_time", _status("2026-09-05", "2026-09-05")[f]["status"], "no_time", failures)

    # Khoảng chuyển mẻ 40 giây (13:00:00-13:00:40) KHÔNG làm ứng viên "gần nhất" (< 5 phút),
    # nhưng lần dừng GIAO trực tiếp với nó vẫn được gán.
    conn.execute(
        "INSERT INTO batch_details (dyelot, machine, start_time, end_time, fabric_type, sap_lot) VALUES ('C100080', 'M1', '2026-09-06 13:00:40', '2026-09-06 14:00:00', 'Cotton', '1000000001')"
    )
    conn.commit()
    _recompute()
    g = _add("M1", "2026-09-06 13:30", "2026-09-06 13:45", "No order", 10)
    h = _add("M1", "2026-09-06 13:00", "2026-09-06 13:01", "Machine cleaning", 11)
    status = _status("2026-09-06", "2026-09-06")
    _check("lần dừng lúc máy chạy: bỏ qua khoảng 40s cách 30 phút, lấy khoảng >= 5 phút cách 86 phút",
           status[g]["deviation_minutes"], round((13 * 60 + 30 - (12 * 60 + 4 + 5 / 60)) , 1), failures)
    _check("khoảng 40s chỉ nhận lần dừng giao trực tiếp", _segments("M1", "2026-09-06"), [
        ("09-06 10:31", "09-06 12:04", "Waiting for previous process"), ("09-06 13:00", "09-06 13:00", "Machine cleaning"),
    ], failures)
    _check("lần dừng giao khoảng ngắn -> matched", status[h]["status"], "matched", failures)

    print("\n=== Kịch bản 5: Save trong drawer, validate, xoá ===")
    service.save_from_report("M2", "2026-09-05 08:00:00", "2026-09-05 10:00:00", "No order", "updated", 1)
    count = get_db().execute("SELECT COUNT(*) AS n FROM idle_time_stops WHERE machine = 'M2' AND stop_start = '2026-09-05 08:00:00'").fetchone()["n"]
    _check("Save cùng đoạn -> sửa lần dừng cũ, không nhân bản", (count, _segments("M2", "2026-09-05")), (1, [("09-05 08:00", "09-05 10:00", "No order")]), failures)
    _expect_error("máy ngoài Machine Master bị từ chối", lambda: service.create_stop("M9", "2026-09-02 01:00", None, "No order", None, 1), failures)
    _expect_error("kết thúc trước bắt đầu bị từ chối", lambda: service.create_stop("M1", "2026-09-02 01:00", "2026-09-02 00:30", "No order", None, 1), failures)
    _expect_error("Reason ngoài danh sách bị từ chối", lambda: service.create_stop("M1", "2026-09-02 01:00", None, "Lunch", None, 1), failures)
    service.delete_stop(y)
    _check("xoá lần dừng sau cùng -> lần trước đó (09:20-09:40) được dùng lại", _segments("M1", "2026-09-02"), [
        ("09-02 09:00", "09-02 09:20", "Waiting for water"), ("09-02 09:20", "09-02 10:00", "Waiting for steam"),
    ], failures)


def _scenario_target_excel(failures: list[str]) -> None:
    print("\n=== Kịch bản 6: Target + Excel ===")
    _check("Target chưa đặt", service.get_target(), None, failures)
    service.set_target(10)
    _check("Target = 10", service.get_target(), 10.0, failures)
    _expect_error("Target > 100 bị từ chối", lambda: service.set_target(120), failures)

    book = load_workbook(io.BytesIO(service.export_idle_excel("2026-09-01", "2026-09-05")))
    sheet = book["Idle Time"]
    _check("header", [cell.value for cell in sheet[1]], ["Machine", "Tank", "Capacity (Kg)", "2026-09-01", "2026-09-02", "2026-09-05", "Total (h)", "% Idle", "Target % Idle"], failures)
    _check("dòng M1 (giờ idle)", [cell.value for cell in sheet[2]], ["M1", "J tank", 600, 1, 1, 0, 2, round(2 / 39 * 100, 2), 10], failures)
    _check("M2 01/09 (100% idle) tô màu vượt Target", sheet.cell(row=3, column=4).fill.start_color.rgb[-6:], "FFC7CE", failures)
    _check("M1 02/09 (4.2%) không tô màu", sheet.cell(row=2, column=5).fill.fill_type, None, failures)
    gaps = [[cell.value for cell in row] for row in book["Idle Gaps"].iter_rows(min_row=2)]
    _check("Idle Gaps: giờ thật + nguyên nhân + giờ người dùng nhập", [g[:10] for g in gaps if g[0] == "M1" and g[1] == "2026-09-01"], [
        ["M1", "2026-09-01", "2026-09-01 22:00:00", "2026-09-01 23:00:00", 1.0, "Waiting for fabric", None, "Idle Entry", "2026-09-01 22:15:00", "2026-09-01 22:30:00"],
    ], failures)
    operating = load_workbook(io.BytesIO(service.export_idle_excel("2026-09-01", "2026-09-05", mode="operating")))["Idle Time"]
    _check("mode operating: M1 01/09 = 12h", operating.cell(row=2, column=4).value, 12, failures)


def main() -> int:
    failures: list[str] = []
    tmp_dir = tempfile.mkdtemp()
    try:
        db_path = os.path.join(tmp_dir, "test.db")
        _init_db(db_path)
        app = Flask(__name__)
        app.config["DATABASE_PATH"] = db_path
        app.config["SQLITE_PRAGMAS"] = {}
        with app.app_context():
            get_db().row_factory = sqlite3.Row
            _recompute()
            _scenario_matrix(failures)
            _scenario_filters(failures)
            _scenario_stops(failures)
            _scenario_target_excel(failures)
            close_db()
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    print(f"\n{'=' * 60}\nKẾT QUẢ: {'TẤT CẢ KHỚP' if not failures else f'{len(failures)} CASE LỆCH'}\n{'=' * 60}")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())

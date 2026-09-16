"""Import + tra cứu bảng mapping Greige Code -> Brand Program (VD "UQ - Ht Fleece").

Nguồn dữ liệu: file Excel do khách hàng/xưởng cung cấp riêng cho từng Brand (VD
"UQ 5 program.xlsx"), map 1 Greige Code sang ĐÚNG 1 cặp (Brand, Program) — đã verify
thực nghiệm 1:1, không có Greige Code nào map lẫn 2 Program khác nhau. Đây là dữ liệu
THAM CHIẾU (reference), khác hẳn cột `program` có sẵn trong `availability_logs`
(Treatment Program — loại quy trình nhuộm, không liên quan tới khái niệm Brand Program
ở đây) — cố tình đặt tên cột `brand_program` để tránh nhầm lẫn.

Đặt ở `core/` (không phải trong 1 Engine cụ thể) vì được dùng ở 2 nơi độc lập:
- `modules/dyeing/engines/excel_import/service.py`: thêm "Brand Program Mapping" làm 1
  lựa chọn Data Type trong Modal Import chung của Dyeing Hub (cùng chỗ với Availability/
  Performance/Batch), theo đúng yêu cầu người dùng gộp chung 1 chỗ import duy nhất.
- `modules/dyeing/engines/reports/cleaning_matrix.py`: JOIN bảng `brand_program_mapping`
  theo `greige_code` (đã có sẵn trong `batch_details`) để tính filter "Brand Program"
  trên báo cáo Batch Per Day by Machine.
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from openpyxl import load_workbook

from core.database import get_db, get_dialect, insert_returning_id
from core.excel_importer import record_import_rows

BRAND_PROGRAM_HEADER_MAP = {
    "greige code": "greige_code", "greigecode": "greige_code",
    "fabric type": "fabric_type",
    "program": "brand_program",
    "cust": "brand", "customer": "brand", "brand": "brand",
    "item code*": "item_code", "item code": "item_code",
}
BRAND_PROGRAM_FIELDS = ("greige_code", "item_code", "fabric_type", "brand_program", "brand")


def _clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none", "nat"} else text


def _read_first_sheet(file_bytes: bytes) -> tuple[list[str], list[list[Any]]]:
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    try:
        rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return [], []
    return [str(value).strip() if value is not None else "" for value in rows[0]], [list(row) for row in rows[1:]]


def parse_brand_program_file(file_bytes: bytes) -> dict[str, Any]:
    """Parse file mapping Greige Code -> Brand Program. Trả {rows, errors, row_details}."""
    headers, raw_rows = _read_first_sheet(file_bytes)
    # Header trùng tên (VD "Item Code*" xuất hiện 2 lần) -> dict lấy CỘT XUẤT HIỆN SAU CÙNG,
    # chấp nhận được vì field đó (item_code) chỉ mang tính tham khảo, không dùng để lọc/join.
    normalized = {header.lower(): index for index, header in enumerate(headers)}
    indexes = {field: normalized[header] for header, field in BRAND_PROGRAM_HEADER_MAP.items() if header in normalized}
    if "greige_code" not in indexes:
        raise ValueError("File thiếu cột bắt buộc: Greige code.")

    parsed: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    row_details: list[dict[str, Any]] = []
    for row_number, values in enumerate(raw_rows, start=2):
        if not any(_clean(value) for value in values):
            continue
        raw = {field: _clean(values[index]) if index < len(values) else "" for field, index in indexes.items()}
        for field in BRAND_PROGRAM_FIELDS:
            raw.setdefault(field, "")
        if not raw["greige_code"]:
            errors.append({"row": row_number, "error": "Thiếu giá trị bắt buộc: Greige code."})
            row_details.append({"row_number": row_number, "status": "invalid", "error": "Thiếu Greige code.", "data": raw})
            continue
        parsed.append(raw)
        row_details.append({"row_number": row_number, "status": "valid", "error": None, "data": raw})
    return {"rows": parsed, "errors": errors, "total_records": len(parsed) + len(errors), "row_details": row_details}


def preview_brand_program_file(file_bytes: bytes) -> dict[str, Any]:
    """Preview 5 dòng đầu — cùng shape JSON mà `import_handler.js` (Modal Import chung
    của Dyeing Hub) đã hiểu cho Availability/Performance/Batch (`columns`/`preview`/
    `valid_rows`/`total_rows`/`errors`), không cần sửa JS."""
    parsed = parse_brand_program_file(file_bytes)
    columns = list(BRAND_PROGRAM_FIELDS)
    return {
        "status": "preview",
        "file_type": "BRAND_PROGRAM",
        "columns": columns,
        "preview": parsed["rows"][:5],
        "valid_rows": len(parsed["rows"]),
        "total_rows": parsed["total_records"],
        "errors": parsed["errors"],
    }


def ensure_brand_program_table(conn: Any) -> None:
    """CHỈ chạy DDL ở SQLite — ở Postgres bảng này PHẢI được tạo thủ công qua
    `supabase/migrate_brand_program_mapping.sql` (bảng MỚI, chưa có sẵn qua `supabase/schema.sql`
    lúc cài đặt lần đầu như các bảng khác — xem quy ước DDL ở systemPatterns.md mục 5.1)."""
    if get_dialect() == "sqlite":
        conn.execute("""
            CREATE TABLE IF NOT EXISTS brand_program_mapping (
                greige_code TEXT PRIMARY KEY,
                item_code TEXT,
                fabric_type TEXT,
                brand_program TEXT,
                brand TEXT,
                import_log_id INTEGER,
                updated_at TEXT NOT NULL DEFAULT (datetime('now'))
            )
        """)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_brand_program_mapping_greige_norm ON brand_program_mapping (LOWER(TRIM(greige_code)))")


def sync_brand_program_mapping(file_bytes: bytes, imported_by: str | None = None, filename: str = "brand_program.xlsx") -> dict[str, Any]:
    """Parse file mapping Greige Code -> Brand Program và UPSERT vào `brand_program_mapping`.

    Cùng pattern audit trail với `core/batch_importer.py::sync_batch_details()`: ghi 1 dòng
    `import_logs` trạng thái 'processing' trước, ghi dữ liệu, rồi cập nhật kết quả cuối.
    """
    result = parse_brand_program_file(file_bytes)
    total_rows = len(result["rows"]) + len(result["errors"])

    conn = get_db()
    ensure_brand_program_table(conn)
    conn.commit()

    log_id = insert_returning_id(
        conn,
        "INSERT INTO import_logs (file_name, import_type, file_type, imported_by, total_rows, error_rows, status) VALUES (?, ?, ?, ?, ?, ?, 'processing')",
        (filename, "brand_program", "BRAND_PROGRAM", imported_by or "unknown", total_rows, len(result["errors"])),
    )
    conn.commit()

    imported_rows = 0
    try:
        if result["rows"]:
            # 1 Greige Code có thể lặp lại nhiều dòng (mỗi dòng 1 màu/Item Code con) nhưng luôn
            # map về CÙNG 1 (Program, Brand) — khử trùng theo greige_code trước executemany() để
            # tránh Postgres từ chối UPSERT trùng khoá trong CÙNG 1 statement (CardinalityViolation,
            # đã gặp bug tương tự ở sync_batch_details()).
            deduped = {row["greige_code"]: row for row in result["rows"]}
            columns_with_log = BRAND_PROGRAM_FIELDS + ("import_log_id",)
            placeholders = ",".join("?" for _ in columns_with_log)
            updates = ",".join(f"{field}=excluded.{field}" for field in columns_with_log if field != "greige_code")
            sql = f"INSERT INTO brand_program_mapping ({','.join(columns_with_log)}) VALUES ({placeholders}) ON CONFLICT(greige_code) DO UPDATE SET {updates}"
            conn.execute("BEGIN")
            conn.executemany(sql, [tuple(row[field] for field in BRAND_PROGRAM_FIELDS) + (log_id,) for row in deduped.values()])
            conn.commit()
            imported_rows = len(deduped)
        record_import_rows(conn, log_id, result["row_details"])
    except Exception as exc:
        conn.rollback()
        conn.execute("UPDATE import_logs SET status='failed', error_detail=? WHERE id=?", (str(exc), log_id))
        conn.commit()
        raise

    status = "completed" if imported_rows and not result["errors"] else "partial" if imported_rows else "failed"
    error_detail = "; ".join(f"Dòng {error['row']}: {error['error']}" for error in result["errors"][:20]) or None
    conn.execute(
        "UPDATE import_logs SET imported_rows=?, success_rows=?, status=?, error_detail=? WHERE id=?",
        (imported_rows, imported_rows, status, error_detail, log_id),
    )
    conn.commit()

    affected_dates: set[Any] = set()
    if result["rows"]:
        from core.production_time import normalize_production_date, production_date_sql_expr
        from core.rollup import trigger_recompute

        greige_codes = sorted({row["greige_code"].lower() for row in result["rows"] if row["greige_code"]})
        placeholders = ",".join("?" for _ in greige_codes)
        shifted_date = production_date_sql_expr("a.end_time")
        affected_rows = conn.execute(
            f"""
            SELECT DISTINCT {shifted_date} AS production_date
            FROM availability_logs a
            JOIN batch_details b ON lower(trim(a.batch)) = lower(trim(b.dyelot)) OR lower(trim(a.batch_ref_no)) = lower(trim(b.dyelot))
            WHERE lower(trim(b.greige_code)) IN ({placeholders}) AND a.end_time IS NOT NULL AND a.end_time != ''
            """,
            greige_codes,
        ).fetchall()
        affected_dates = {
            datetime.strptime(normalize_production_date(row["production_date"]), "%Y-%m-%d").date()
            for row in affected_rows if row["production_date"]
        }
        trigger_recompute(affected_dates)

    return {
        "status": status,
        "file_type": "BRAND_PROGRAM",
        "rows_imported": imported_rows,
        "total_rows": total_rows,
        "errors": result["errors"],
        "import_log_id": log_id,
        "affected_production_dates": len(affected_dates),
    }

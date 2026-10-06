"""
tests/test_batch_day_trend_drilldown.py
-----------------------------------------
Verify drill-down danh sách mẻ + Excel export của tab Batch/Day Trend (2026-10-06):

- Bấm 1 ô (loại vải x kỳ) / ô Total -> liệt kê TẤT CẢ đoạn (mẻ x ngày) góp giờ, cờ Counted.
- Tổng giờ / số mẻ Counted / Batch/Day tính lại khớp ĐÚNG giá trị ô trên bảng Trend.
- Mẻ qua đêm: hiện ở cả 2 ngày, chỉ Counted ở ngày kết thúc.
- Machine trống ở raw (fill-down) vẫn tra được Dyelot; FabricType trống đánh dấu "from next".
- 2 mẻ trùng (machine, start_time) -> tách lại giờ/Counted riêng từng mẻ.
- Excel: cả báo cáo (Trend + Batches + Filters) và 1 ô (Batches + Filters), có khối Check.

Chạy: python tests/test_batch_day_trend_drilldown.py
"""
from __future__ import annotations

import io
import os
import shutil
import sqlite3
import sys
import tempfile
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from flask import Flask  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from core.database import close_db  # noqa: E402
from modules.dyeing.engines.batch_matrix.batch_day_trend import (  # noqa: E402
    export_batch_day_trend_excel,
    get_batch_day_trend,
    get_trend_batches,
    recompute_all,
)


def _make_temp_app(db_path: str) -> Flask:
    app = Flask(__name__)
    app.config["DATABASE_PATH"] = db_path
    app.config["SQLITE_PRAGMAS"] = {}
    return app


def _init_db(db_path: str) -> None:
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    conn.execute(
        "CREATE TABLE batch_details (id INTEGER PRIMARY KEY AUTOINCREMENT, dyelot TEXT NOT NULL, sap_lot TEXT, "
        "redye REAL DEFAULT 0, greige_code TEXT, fabric_type TEXT, machine TEXT, start_time TEXT, end_time TEXT, "
        "run_time REAL)"
    )
    conn.execute("CREATE TABLE machines (machine_code TEXT, machine_id TEXT, capacity_kg REAL, tank_type TEXT)")
    conn.execute("INSERT INTO machines (machine_code, capacity_kg, tank_type) VALUES ('M1', 500, 'J tank'), ('M2', 300, NULL)")
    # (dyelot, sap_lot, redye, fabric_type, machine, start, end) — thứ tự INSERT = thứ tự import.
    rows = [
        ("CL001", "", 0, "", "M1", "2026-06-01 07:00:00", "2026-06-01 09:00:00"),       # CM, Fabric từ mẻ sau (CVC), 2h day1
        ("D0010", "1000", 0, "CVC", "M1", "2026-06-01 09:00:00", "2026-06-01 13:00:00"),  # Normal 4h day1
        ("D0020", "3000", 0, "CVC", "M1", "2026-06-01 22:00:00", "2026-06-02 10:00:00"),  # Normal 9h day1 + 3h day2, đếm day2
        ("D0021", "1000", 0, "CVC", "", "2026-06-02 10:00:00", "2026-06-02 12:00:00"),    # Machine trống -> M1, Rework 2h
        ("D0030", "1000", 1, "CVC", "M1", "2026-06-02 12:00:00", "2026-06-03 07:00:00"),  # ReDye 1, 19h day2
        ("D0050", "1000", 0, "Cotton", "M2", "2026-06-03 07:00:00", "2026-06-03 12:00:00"),  # Normal 5h day3
        ("D0051", "1000", 0, "Cotton", "M2", "2026-06-03 07:00:00", "2026-06-03 09:00:00"),  # TRÙNG start M2, Rework 2h
        ("D0060", "4000", 0, "Cotton", "M2", "2026-06-03 12:00:00", "2026-06-03 13:00:00"),  # Sample 1h
    ]
    conn.executemany(
        "INSERT INTO batch_details (dyelot, sap_lot, redye, fabric_type, machine, start_time, end_time, run_time) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, 0)",
        rows,
    )
    conn.commit()
    conn.close()


def _check(label: str, actual, expected, failures: list[str]) -> None:
    ok = actual == expected
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}: actual={actual!r} expected={expected!r}")
    if not ok:
        failures.append(label)


def _by_dyelot(data: dict) -> dict[str, dict]:
    return {row["dyelot"]: row for row in data["batches"]}


def main() -> int:
    failures: list[str] = []
    tmp_dir = tempfile.mkdtemp()
    db_path = os.path.join(tmp_dir, "drill.db")
    _init_db(db_path)
    try:
        with _make_temp_app(db_path).app_context():
            conn = sqlite3.connect(db_path)
            conn.row_factory = sqlite3.Row
            recompute_all({date(2026, 6, 1), date(2026, 6, 2), date(2026, 6, 3)}, conn)
            conn.commit()
            conn.close()

            print("=== Ô CVC ngày 01/06 ===")
            day1 = get_trend_batches("CVC", "2026-06-01", group_by="date")
            rows = _by_dyelot(day1)
            _check("3 đoạn (CM + Normal + đoạn đầu mẻ qua đêm)", sorted(rows), ["CL001", "D0010", "D0020"], failures)
            _check("CL001 = CM, Fabric lấy từ mẻ sau", (rows["CL001"]["classification"], rows["CL001"]["fabric_filled"]), ("CM", True), failures)
            _check("D0020 ngày đầu: Normal nhưng chưa Counted, 9h", (rows["D0020"]["classification"], rows["D0020"]["counted"], rows["D0020"]["hours"]), ("Normal", False, 9.0), failures)
            _check("summary 15h / 1 mẻ / 1.6", (day1["summary"]["occupied_hours"], day1["summary"]["normal_batches"], day1["summary"]["batch_per_day"]), (15.0, 1, 1.6), failures)

            print("\n=== Ô CVC ngày 02/06 ===")
            day2 = get_trend_batches("CVC", "2026-06-02", group_by="date")
            rows = _by_dyelot(day2)
            _check("D0020 Counted ở ngày kết thúc, 3h", (rows["D0020"]["counted"], rows["D0020"]["hours"]), (True, 3.0), failures)
            _check("D0021 Machine trống vẫn tra được, = M1, Rework", (rows["D0021"]["machine"], rows["D0021"]["classification"]), ("M1", "Rework"), failures)
            _check("D0030 ReDye 1 -> Rework", rows["D0030"]["classification"], "Rework", failures)
            day2_any = _by_dyelot(get_trend_batches("CVC", "2026-06-02", group_by="date", require_redye_zero=False))
            _check("Tắt ReDye = 0 -> D0030 Normal + Counted", (day2_any["D0030"]["classification"], day2_any["D0030"]["counted"]), ("Normal", True), failures)

            print("\n=== 2 mẻ trùng (M2, 03/06 07:00) ===")
            day3 = _by_dyelot(get_trend_batches("Cotton", "2026-06-03", group_by="date"))
            _check("D0050 5h Counted / D0051 2h Rework", ((day3["D0050"]["hours"], day3["D0050"]["counted"]), (day3["D0051"]["hours"], day3["D0051"]["classification"])), ((5.0, True), (2.0, "Rework")), failures)
            _check("D0060 Sample", day3["D0060"]["classification"], "Sample", failures)

            print("\n=== Mọi ô + Total khớp bảng Trend ===")
            for group_by in ("date", "week", "month"):
                trend = get_batch_day_trend(group_by=group_by)
                for row in trend["rows"]:
                    for index, key in enumerate(trend["period_keys"]):
                        drill = get_trend_batches(row["fabric_type"], key, group_by=group_by)["summary"]
                        _check(f"{group_by} {row['fabric_type']} {key}", (drill["batch_per_day"], drill["normal_batches"]), (row["values"][index], row["counts"][index]), failures)
                    total = get_trend_batches(row["fabric_type"], None, group_by=group_by)["summary"]
                    _check(f"{group_by} {row['fabric_type']} Total", total["batch_per_day"], row["total"], failures)

            print("\n=== Bộ lọc áp dụng cho drill-down ===")
            _check("Tank Unclassified: CVC (M1 = J tank) không còn mẻ", get_trend_batches("CVC", None, tank_types="Unclassified")["batches"], [], failures)
            _check("Capacity 300: Cotton còn 3 đoạn", len(get_trend_batches("Cotton", None, capacities="300")["batches"]), 3, failures)
            try:
                get_trend_batches("Nylon")
                _check("Fabric type sai -> ValueError", "no error", "ValueError", failures)
            except ValueError:
                _check("Fabric type sai -> ValueError", "ValueError", "ValueError", failures)

            print("\n=== Excel export ===")
            full = load_workbook(io.BytesIO(export_batch_day_trend_excel(group_by="date")))
            _check("Cả báo cáo: 3 sheet", full.sheetnames, ["Trend", "Batches", "Filters"], failures)
            batches = full["Batches"]
            _check("Header Batches", [c.value for c in batches[1]][:3], ["Production Date", "Machine", "Capacity (Kg)"], failures)
            data_rows = [r for r in batches.iter_rows(min_row=2, values_only=True) if r[0] and str(r[0]).startswith("2026-")]
            _check("Batches: 6 đoạn CVC + 3 đoạn Cotton", len(data_rows), 9, failures)
            all_row = next(r for r in batches.iter_rows(values_only=True) if r[0] == "All")
            _check("Check All: 3 mẻ Normal", all_row[2], 3, failures)
            trend_sheet = full["Trend"]
            titles = [r[0] for r in trend_sheet.iter_rows(values_only=True) if r[0] in ("Batch/Day", "Normal batches (counted)", "Occupied hours")]
            _check("Trend có 3 khối", titles, ["Batch/Day", "Normal batches (counted)", "Occupied hours"], failures)

            cell = load_workbook(io.BytesIO(export_batch_day_trend_excel(group_by="date", fabric_type="CVC", period_key="2026-06-01")))
            _check("1 ô: 2 sheet", cell.sheetnames, ["Batches", "Filters"], failures)
            counted = [r[14] for r in cell["Batches"].iter_rows(min_row=2, max_row=4, values_only=True)]
            _check("Cột Counted", counted, ["No", "Yes", "No"], failures)
            check = next(r for r in cell["Batches"].iter_rows(values_only=True) if r[0] == "CVC")
            _check("Check CVC: 15h / 1 / 1.6", check[1:4], (15.0, 1, 1.6), failures)
            filters = {r[0]: r[1] for r in cell["Filters"].iter_rows(min_row=2, values_only=True)}
            _check("Filters ghi kỳ đã bấm", (filters["Fabric Type"], filters["Period"]), ("CVC", "01 Jun 2026"), failures)
            close_db()
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    if failures:
        print(f"\n{len(failures)} FAILURE(S): {failures}")
        return 1
    print("\nAll checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())

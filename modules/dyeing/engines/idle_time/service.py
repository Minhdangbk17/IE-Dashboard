"""Báo cáo Idle Time — thời gian máy KHÔNG chạy mẻ nào, theo máy x Production Date — và nhập
tay lần dừng máy (Idle Entry) để gán nguyên nhân.

NGUỒN DỮ LIỆU: `batch_day_trend_daily_summary` (rollup của Batch/Day Trend, Engine
`batch_matrix`) — mỗi dòng là 1 ĐOẠN của 1 mẻ trong 1 production_date, đã chia theo khung
07:00 -> 07:00 (`core/batch_source.py::split_production_days()`). Đọc lại ĐÚNG bảng này nên
giờ hoạt động luôn khớp Batch/Day (cùng tập mẻ, gồm cả CM/Sample/Rework/FabricType Unknown),
KHÔNG có rollup riêng, KHÔNG cần `flask rebuild-summaries`.

CÁCH TÍNH IDLE (2026-10-05, đã chốt với người dùng):
1. Khoảng thời gian của 1 đoạn: [max(StartTime mẻ, 07:00 của ngày), + hours]. Đúng vì
   `split_production_days()` tính hours = min(End, 07:00 hôm sau) - max(Start, 07:00).
2. Giờ CÓ DỮ LIỆU của 1 ngày (mẫu số, tính trên TẤT CẢ máy — không phụ thuộc filter): bắt đầu
   07:00 nếu có mẻ chạy NỐI từ hôm trước sang, ngược lại = đoạn bắt đầu sớm nhất trong ngày;
   kết thúc 07:00 hôm sau nếu có mẻ chạy nối sang hôm sau, ngược lại = đoạn kết thúc muộn
   nhất. Ngày đầu/cuối file import (VD 02/09 bắt đầu 18:42) chỉ tính phần có dữ liệu thay vì
   24h; ngày KHÔNG có đoạn nào bị bỏ hẳn (khoảng hổng giữa 2 lần import, không phải idle).
3. Danh sách máy = máy có ít nhất 1 đoạn trong kỳ (người dùng: "máy phải hoạt động mới có
   data") — KHÔNG lấy từ Machine Master. Ngày máy không chạy mẻ nào = idle trọn giờ có dữ liệu.
4. Idle của 1 ô = tổng các KHOẢNG TRỐNG giữa các đoạn (đã gộp đoạn chồng nhau) trong giờ có
   dữ liệu -> luôn = giờ có dữ liệu - giờ hoạt động.
5. % Idle = idle / giờ có dữ liệu. Target là 1 ngưỡng % Idle CHUNG, CHỈ admin sửa.

GÁN NGUYÊN NHÂN (2026-10-05, bản 2 — thay note theo khoảng idle `idle_time_notes`):
Mọi nguyên nhân là 1 LẦN DỪNG dạng khoảng thời gian trong `idle_time_stops` — nhập tay ở
trang Idle Entry (ca đêm, trước khi upload Batch; giờ làm tròn) hoặc bấm Save trong drawer báo
cáo (giờ = đúng khoảng idle). Ghép với khoảng idle ở READ TIME nên upload/upload lại Batch tự
khớp lại, không bao giờ "mất note":
a. Khoảng idle được ghép trên khoảng LIÊN TỤC (khoảng vắt qua 07:00 là 1 khoảng — các phần ở 2
   ngày đều nhận nguyên nhân), rồi mới chia lại theo ngày để hiển thị.
b. Lần dừng chồng giờ trên CÙNG máy: lần cập nhật SAU CÙNG (`updated_at`, rồi `id`) thắng phần
   chồng; lần cũ còn lại phần không bị chồng (hoặc "overridden" nếu bị chồng hết).
c. Giờ nhập chỉ là CƠ SỞ: phần còn lại của lần dừng giao với khoảng idle nào thì gán cho khoảng
   đó; không giao khoảng nào -> gán cho khoảng idle GẦN NHẤT >= 5 phút của máy (khoảng chuyển
   mẻ ngắn hơn không làm ứng viên; không giới hạn độ lệch,
   trong phạm vi dữ liệu đã nạp = kỳ đang xem ±1 ngày), ghi lại độ lệch (phút).
d. Nhiều lần dừng cùng gán vào 1 khoảng idle -> chia khoảng tại mốc bắt đầu của lần dừng sau;
   đoạn đầu kéo về đầu khoảng idle, đoạn cuối kéo tới cuối khoảng idle. Thời gian hiển thị/tính
   LUÔN là giờ chính xác từ Batch.
e. Lần dừng chưa nằm trọn trong giờ có dữ liệu Batch -> "pending" (chờ upload), chưa gán.

Filter Tank (J tank / O tank / Unclassified) + Capacity tra `machines` tại read time, chuẩn hoá
Tank bằng `_normalize_tank_label()` giống Batch Per Day by Machine + Batch/Day Trend.
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date, datetime, timedelta
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from core.batch_source import parse_batch_datetime
from core.database import execute_query, get_db, get_dialect, insert_returning_id
from core.production_time import PRODUCTION_SHIFT_START_HOUR
from modules.dyeing.engines.reports.cleaning_matrix import (
    CANONICAL_TANK_ORDER,
    UNCLASSIFIED_TANK_LABEL,
    _normalize_tank_label,
)

# Danh sách Reason cố định (người dùng chốt 2026-10-05) — bám theo các cột chờ/dừng của file
# Availability để đối chiếu được với báo cáo Downtime, + "Waiting for previous process".
IDLE_REASONS = (
    "Waiting for fabric",
    "Waiting for chemicals/dyes",
    "Waiting for water",
    "Waiting for steam",
    "Maintenance",
    "Machine breakdown",
    "Machine cleaning",
    "No order",
    "Waiting for previous process",
)
SOURCE_ENTRY = "entry"    # nhập tay ở trang Idle Entry
SOURCE_REPORT = "report"  # bấm Save trong drawer báo cáo (giờ = đúng khoảng idle)
TARGET_SETTING_KEY = "target_idle_pct"
_NOTES_MIGRATED_KEY = "idle_time_notes_migrated"
_TANK_ORDER: tuple[str, ...] = (*CANONICAL_TANK_ORDER, UNCLASSIFIED_TANK_LABEL)
_TIME_FMT = "%Y-%m-%d %H:%M:%S"
# Lần dừng chưa có giờ kết thúc (máy vẫn đang dừng lúc hết ca) được coi như 1 điểm thời gian.
_OPEN_STOP_SPAN = timedelta(seconds=1)
_MACHINE_LOOKAROUND_DAYS = 7
_NEAREST_MIN_GAP = timedelta(minutes=5)

# Màu ô vượt Target trong Excel — style "Bad" chuẩn của Excel (nền hồng/chữ đỏ đậm), tương
# đương `.ratio-warning` trên web (xem design/UI_GUIDELINES.md mục 6.3).
_EXCEL_OVER_TARGET_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
_EXCEL_OVER_TARGET_FONT = Font(color="9C0006")


# ---------------------------------------------------------------------------
# Bảng của Engine — SQLite tự tạo; Postgres tạo qua `supabase/migrate_idle_time*.sql`.
# ---------------------------------------------------------------------------


def _ensure_tables(conn: Any) -> None:
    if get_dialect() == "sqlite":
        conn.execute("""
            CREATE TABLE IF NOT EXISTS idle_time_stops (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                machine TEXT NOT NULL,
                stop_start TEXT NOT NULL,
                stop_end TEXT,
                reason TEXT,
                detail TEXT,
                source TEXT NOT NULL DEFAULT 'entry',
                created_by INTEGER NOT NULL REFERENCES users (id),
                created_at TEXT NOT NULL DEFAULT (datetime('now')),
                updated_by INTEGER NOT NULL REFERENCES users (id),
                updated_at TEXT NOT NULL DEFAULT (datetime('now'))
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS idle_time_settings (
                key TEXT PRIMARY KEY,
                value REAL NOT NULL,
                updated_at TEXT NOT NULL DEFAULT (datetime('now'))
            )
        """)
        _migrate_legacy_notes(conn)
    conn.execute("CREATE INDEX IF NOT EXISTS idx_idle_time_stops_machine_start ON idle_time_stops (machine, stop_start)")
    conn.commit()


def _migrate_legacy_notes(conn: Any) -> None:
    """Chuyển note bản 1 (`idle_time_notes`, khoá machine + gap_start) thành lần dừng
    source='report' có giờ = đúng khoảng idle lúc lưu — 1 LẦN (cờ trong `idle_time_settings`).
    Không xoá bảng cũ. Postgres làm việc này trong `supabase/migrate_idle_time_stops.sql`."""
    exists = conn.execute("SELECT name FROM sqlite_master WHERE type = 'table' AND name = 'idle_time_notes'").fetchone()
    done = conn.execute("SELECT 1 FROM idle_time_settings WHERE key = ?", (_NOTES_MIGRATED_KEY,)).fetchone()
    if exists is None or done is not None:
        return
    conn.execute("""
        INSERT INTO idle_time_stops (machine, stop_start, stop_end, reason, detail, source, created_by, created_at, updated_by, updated_at)
        SELECT machine, gap_start, gap_end, reason, detail, 'report', updated_by, updated_at, updated_by, updated_at FROM idle_time_notes
    """)
    conn.execute("INSERT INTO idle_time_settings (key, value) VALUES (?, 1)", (_NOTES_MIGRATED_KEY,))


# ---------------------------------------------------------------------------
# Tiện ích thời gian / khoảng
# ---------------------------------------------------------------------------


def _fmt(value: datetime) -> str:
    return value.strftime(_TIME_FMT)


def _day_window(day_str: str) -> tuple[datetime, datetime]:
    day = date.fromisoformat(day_str)
    start = datetime(day.year, day.month, day.day, PRODUCTION_SHIFT_START_HOUR)
    return start, start + timedelta(days=1)


def _shift_day(day_str: str, days: int) -> str:
    return (date.fromisoformat(day_str) + timedelta(days=days)).isoformat()


def _hours(start: datetime, end: datetime) -> float:
    return (end - start).total_seconds() / 3600.0


def _subtract(interval: tuple[datetime, datetime], covered: list[tuple[datetime, datetime]]) -> list[tuple[datetime, datetime]]:
    """Phần của `interval` KHÔNG bị các khoảng `covered` che."""
    pieces = [interval]
    for c_start, c_end in covered:
        next_pieces = []
        for p_start, p_end in pieces:
            if c_end <= p_start or c_start >= p_end:
                next_pieces.append((p_start, p_end))
                continue
            if c_start > p_start:
                next_pieces.append((p_start, c_start))
            if c_end < p_end:
                next_pieces.append((c_end, p_end))
        pieces = next_pieces
    return pieces


def _merge_intervals(intervals: Iterable[tuple[datetime, datetime]]) -> list[tuple[datetime, datetime]]:
    merged: list[tuple[datetime, datetime]] = []
    for start, end in sorted(intervals):
        if merged and start <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], end))
        else:
            merged.append((start, end))
    return merged


def _distance_minutes(start: datetime, end: datetime, gap: dict[str, Any]) -> float:
    if end <= gap["start"]:
        return (gap["start"] - end).total_seconds() / 60.0
    if start >= gap["end"]:
        return (start - gap["end"]).total_seconds() / 60.0
    return 0.0


def _parse_input_datetime(value: Any) -> datetime | None:
    """Nhận 'YYYY-MM-DDTHH:MM' của <input type=datetime-local> hoặc 'YYYY-MM-DD HH:MM[:SS]'."""
    text = str(value or "").strip().replace("T", " ")
    return parse_batch_datetime(text) if text else None


# ---------------------------------------------------------------------------
# Dựng đoạn mẻ / giờ có dữ liệu / khoảng idle
# ---------------------------------------------------------------------------


def _load_day_segments(from_date: str | None, to_date: str | None) -> dict[str, dict[str, list[tuple[datetime, datetime, str]]]]:
    """{production_date: {machine: [(đoạn bắt đầu, đoạn kết thúc, StartTime gốc của mẻ)]}} — TẤT
    CẢ máy (chưa lọc), đoạn đã sắp theo thời gian. Kết thúc làm tròn tới giây để 2 mẻ nối
    liền nhau không sinh khoảng idle ảo do sai số float."""
    sql = "SELECT production_date, machine, start_time, hours FROM batch_day_trend_daily_summary WHERE hours > 0"
    params: list[Any] = []
    if from_date:
        sql += " AND production_date >= ?"
        params.append(from_date)
    if to_date:
        sql += " AND production_date <= ?"
        params.append(to_date)
    result: dict[str, dict[str, list[tuple[datetime, datetime, str]]]] = defaultdict(lambda: defaultdict(list))
    for row in execute_query(sql, params):
        batch_start = parse_batch_datetime(row["start_time"])
        machine = str(row["machine"] or "").strip()
        if batch_start is None or not machine:
            continue
        window_start, _ = _day_window(row["production_date"])
        seg_start = max(batch_start, window_start)
        seg_end = seg_start + timedelta(seconds=round(float(row["hours"]) * 3600))
        result[row["production_date"]][machine].append((seg_start, seg_end, str(row["start_time"])))
    for machines in result.values():
        for segments in machines.values():
            segments.sort()
    return result


def _available_window(day_str: str, machines: dict[str, list[tuple[datetime, datetime, str]]]) -> tuple[datetime, datetime]:
    """Giờ có dữ liệu của ngày — xem bước 2 docstring đầu file."""
    window_start, window_end = _day_window(day_str)
    all_segments = [segment for segments in machines.values() for segment in segments]
    carried_in = any(seg_start <= window_start for seg_start, _, _ in all_segments)
    carried_out = any(seg_end >= window_end for _, seg_end, _ in all_segments)
    start = window_start if carried_in else min(seg_start for seg_start, _, _ in all_segments)
    end = window_end if carried_out else max(seg_end for _, seg_end, _ in all_segments)
    return start, end


def _day_gaps(segments: list[tuple[datetime, datetime, str]], avail_start: datetime, avail_end: datetime) -> tuple[float, list[dict[str, Any]]]:
    """(giờ hoạt động = độ dài HỢP các đoạn, [khoảng idle]) trong [avail_start, avail_end].
    Mỗi khoảng idle kèm StartTime gốc của mẻ ngay trước/ngay sau (None nếu ở biên ngày)."""
    operating_seconds = 0.0
    gaps: list[dict[str, Any]] = []
    cursor = avail_start
    prev_batch: str | None = None
    for seg_start, seg_end, batch_start in segments:
        seg_start, seg_end = max(seg_start, avail_start), min(seg_end, avail_end)
        if seg_end <= seg_start:
            continue
        if seg_start > cursor:
            gaps.append({"start": cursor, "end": seg_start, "prev_batch_start": prev_batch, "next_batch_start": batch_start})
        if seg_end > cursor:
            operating_seconds += (seg_end - max(seg_start, cursor)).total_seconds()
            cursor = seg_end
            prev_batch = batch_start
    if avail_end > cursor:
        gaps.append({"start": cursor, "end": avail_end, "prev_batch_start": prev_batch, "next_batch_start": None})
    return operating_seconds / 3600.0, gaps


def _machine_master(conn: Any) -> dict[str, dict[str, Any]]:
    """{mã máy (lower/trim): {"code", "tank", "capacity"}} từ Machine Master."""
    master: dict[str, dict[str, Any]] = {}
    for row in conn.execute("SELECT machine_code, machine_id, tank_type, capacity_kg FROM machines").fetchall():
        code = str(row["machine_code"] or row["machine_id"] or "").strip()
        if code and code.lower() not in master:
            capacity = round(float(row["capacity_kg"]), 2) if row["capacity_kg"] not in (None, "") else None
            master[code.lower()] = {"code": code, "tank": _normalize_tank_label(row["tank_type"]), "capacity": capacity}
    return master


def _pct(part: float, whole: float) -> float | None:
    return round(part / whole * 100, 2) if whole > 0 else None


# ---------------------------------------------------------------------------
# Lần dừng (idle_time_stops) + ghép vào khoảng idle
# ---------------------------------------------------------------------------


def _load_stops(range_start: datetime | None, range_end: datetime | None, stop_id: int | None = None) -> list[Any]:
    sql = """
        SELECT s.id, s.machine, s.stop_start, s.stop_end, s.reason, s.detail, s.source, s.created_at, s.updated_at,
               cu.username AS created_by, uu.username AS updated_by
        FROM idle_time_stops s
        LEFT JOIN users cu ON cu.id = s.created_by
        LEFT JOIN users uu ON uu.id = s.updated_by
        WHERE 1=1
    """
    params: list[Any] = []
    if stop_id is not None:
        sql += " AND s.id = ?"
        params.append(stop_id)
    if range_end is not None:
        sql += " AND s.stop_start < ?"
        params.append(_fmt(range_end))
    if range_start is not None:
        sql += " AND COALESCE(s.stop_end, s.stop_start) >= ?"
        params.append(_fmt(range_start))
    return execute_query(sql, params)


def _stop_payload(stop: Any) -> dict[str, Any]:
    return {
        "id": stop["id"], "machine": stop["machine"], "stop_start": stop["stop_start"], "stop_end": stop["stop_end"],
        "reason": stop["reason"], "detail": stop["detail"], "source": stop["source"],
        "created_by": stop["created_by"], "created_at": stop["created_at"],
        "updated_by": stop["updated_by"], "updated_at": stop["updated_at"],
    }


def _machines_near(from_date: str | None, to_date: str | None) -> set[str]:
    """Máy có mẻ trong kỳ ±`_MACHINE_LOOKAROUND_DAYS` — để máy không chạy mẻ nào quanh 1 ngày
    đang xem lẻ (drawer, danh sách lần dừng) vẫn có khoảng idle trọn ngày, khớp với ma trận cả
    kỳ (máy có mẻ ở ngày khác trong kỳ vẫn hiện, ngày không chạy = idle trọn)."""
    sql = "SELECT DISTINCT machine FROM batch_day_trend_daily_summary WHERE hours > 0"
    params: list[Any] = []
    if from_date:
        sql += " AND production_date >= ?"
        params.append(_shift_day(from_date, -_MACHINE_LOOKAROUND_DAYS))
    if to_date:
        sql += " AND production_date <= ?"
        params.append(_shift_day(to_date, _MACHINE_LOOKAROUND_DAYS))
    return {str(row["machine"]).strip() for row in execute_query(sql, params) if str(row["machine"] or "").strip()}


def _context(from_date: str | None, to_date: str | None) -> dict[str, Any]:
    """Dựng TOÀN BỘ dữ liệu cho kỳ [from_date, to_date] (nạp thêm ±1 ngày để khoảng idle vắt
    qua 07:00 ở biên kỳ và lần dừng gần biên được ghép đúng): ô máy x ngày, khoảng idle liên
    tục đã chia đoạn theo lần dừng, trạng thái từng lần dừng."""
    conn = get_db()
    _ensure_tables(conn)
    load_from = _shift_day(from_date, -1) if from_date else None
    load_to = _shift_day(to_date, 1) if to_date else None
    day_segments = _load_day_segments(load_from, load_to)
    days = sorted(day_segments)
    windows = {day: _available_window(day, day_segments[day]) for day in days}
    machines = sorted({m for day in days for m in day_segments[day]} | _machines_near(from_date, to_date))

    # 1. Ô máy x ngày + khoảng idle LIÊN TỤC (khoảng cuối ngày D nối khoảng đầu ngày D+1 tại 07:00).
    cells: dict[tuple[str, str], dict[str, Any]] = {}
    gaps_by_machine: dict[str, list[dict[str, Any]]] = defaultdict(list)
    machine_names: dict[str, str] = {}
    for machine in machines:
        key = machine.lower()
        machine_names.setdefault(key, machine)
        for day in days:
            avail_start, avail_end = windows[day]
            operating, day_gaps = _day_gaps(day_segments[day].get(machine, []), avail_start, avail_end)
            cells[(machine, day)] = {"operating": operating, "available": _hours(avail_start, avail_end)}
            for gap in day_gaps:
                physical = gaps_by_machine[key]
                if physical and physical[-1]["end"] == gap["start"]:
                    physical[-1]["end"] = gap["end"]
                    physical[-1]["next_batch_start"] = gap["next_batch_start"]
                    physical[-1]["parts"].append((day, gap["start"], gap["end"]))
                else:
                    physical.append({**gap, "machine": machine, "parts": [(day, gap["start"], gap["end"])]})
    coverage = _merge_intervals(windows.values())

    # 2. Lần dừng: phân giải chồng giờ (sau cùng thắng) rồi gán vào khoảng idle.
    range_start = windows[days[0]][0] if days else None
    range_end = windows[days[-1]][1] if days else None
    stops = _load_stops(range_start, range_end) if days else []
    stop_state: dict[int, dict[str, Any]] = {}
    assignments: dict[tuple[str, int], list[tuple[datetime, int]]] = defaultdict(list)
    stops_by_machine: dict[str, list[Any]] = defaultdict(list)
    for stop in stops:
        stops_by_machine[str(stop["machine"]).strip().lower()].append(stop)
    for key, machine_stops in stops_by_machine.items():
        covered: list[tuple[datetime, datetime]] = []
        for stop in sorted(machine_stops, key=lambda s: (str(s["updated_at"]), s["id"]), reverse=True):
            start = _parse_input_datetime(stop["stop_start"])
            end = _parse_input_datetime(stop["stop_end"]) or (start + _OPEN_STOP_SPAN if start else None)
            state = {"stop": stop, "status": "unmatched", "partially_overridden": False, "deviation_minutes": None}
            stop_state[stop["id"]] = state
            if start is None or end is None or end <= start:
                continue
            pieces = _subtract((start, end), covered)
            covered.append((start, end))
            if not pieces:
                state["status"] = "overridden"
                continue
            state["partially_overridden"] = sum(_hours(s, e) for s, e in pieces) < _hours(start, end) - 1e-9
            if not any(c_start <= start and end <= c_end for c_start, c_end in coverage):
                state["status"] = "pending"
                continue
            physical = gaps_by_machine.get(key, [])
            # Ứng viên "gần nhất" chỉ là khoảng idle >= 5 phút (người dùng chốt 2026-10-06) —
            # tránh gán lần dừng nhập lúc máy đang chạy vào 1 khoảng chuyển mẻ vài chục giây ở
            # xa. Lần dừng GIAO trực tiếp khoảng ngắn thì vẫn gán bình thường.
            nearest_candidates = [i for i, gap in enumerate(physical) if gap["end"] - gap["start"] >= _NEAREST_MIN_GAP]
            deviation = 0.0
            assigned = False
            for piece_start, piece_end in pieces:
                hits = [i for i, gap in enumerate(physical) if gap["start"] < piece_end and piece_start < gap["end"]]
                if not hits and nearest_candidates:
                    nearest = min(nearest_candidates, key=lambda i: _distance_minutes(piece_start, piece_end, physical[i]))
                    deviation = max(deviation, _distance_minutes(piece_start, piece_end, physical[nearest]))
                    hits = [nearest]
                for index in hits:
                    assignments[(key, index)].append((piece_start, stop["id"]))
                    assigned = True
            if assigned:
                state["status"] = "matched"
                state["deviation_minutes"] = round(deviation, 1)

    # 3. Chia từng khoảng idle liên tục theo lần dừng được gán, rồi chiếu về từng ngày.
    day_gaps: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    stop_segments: dict[int, list[tuple[datetime, datetime]]] = defaultdict(list)
    for key, physical in gaps_by_machine.items():
        for index, gap in enumerate(physical):
            sequence: list[tuple[datetime, int]] = []
            for piece_start, stop_id in sorted(assignments.get((key, index), [])):
                if not sequence or sequence[-1][1] != stop_id:
                    sequence.append((piece_start, stop_id))
            if not sequence:
                segments: list[tuple[datetime, datetime, int | None]] = [(gap["start"], gap["end"], None)]
            else:
                bounds = [gap["start"], *[min(max(s, gap["start"]), gap["end"]) for s, _ in sequence[1:]], gap["end"]]
                segments = [(bounds[i], bounds[i + 1], sequence[i][1]) for i in range(len(sequence)) if bounds[i + 1] > bounds[i]]
            for seg_start, seg_end, stop_id in segments:
                if stop_id is not None:
                    stop_segments[stop_id].append((seg_start, seg_end))
            for day, part_start, part_end in gap["parts"]:
                part_segments = [
                    {"start": max(s, part_start), "end": min(e, part_end), "stop_id": sid}
                    for s, e, sid in segments if min(e, part_end) > max(s, part_start)
                ]
                day_gaps[(gap["machine"], day)].append({
                    "start": part_start, "end": part_end, "gap_start": gap["start"], "gap_end": gap["end"],
                    "prev_batch_start": gap["prev_batch_start"], "next_batch_start": gap["next_batch_start"],
                    "segments": part_segments,
                })
    for stop_id, state in stop_state.items():
        state["segments"] = sorted(stop_segments.get(stop_id, []))
        if state["status"] == "matched" and not state["segments"]:
            state["status"] = "no_time"  # bị lần dừng sau chiếm hết khoảng idle được gán

    for (machine, day), gaps in day_gaps.items():
        cell = cells[(machine, day)]
        cell["idle"] = sum(_hours(g["start"], g["end"]) for g in gaps)
        cell["explained"] = sum(_hours(s["start"], s["end"]) for g in gaps for s in g["segments"] if s["stop_id"] is not None)
    return {
        "days": days, "windows": windows, "machines": machines, "cells": cells, "day_gaps": day_gaps,
        "stop_state": stop_state, "conn": conn,
    }


# ---------------------------------------------------------------------------
# Target (% Idle chung cho mọi máy — chỉ admin sửa)
# ---------------------------------------------------------------------------


def get_target() -> float | None:
    conn = get_db()
    _ensure_tables(conn)
    rows = execute_query("SELECT value FROM idle_time_settings WHERE key = ?", [TARGET_SETTING_KEY])
    return float(rows[0]["value"]) if rows else None


def set_target(target_idle_pct: float) -> dict[str, Any]:
    if not 0 <= target_idle_pct <= 100:
        raise ValueError("Target % Idle phải trong khoảng 0-100.")
    conn = get_db()
    _ensure_tables(conn)
    now_str = datetime.now().strftime(_TIME_FMT)
    conn.execute(
        "INSERT INTO idle_time_settings (key, value, updated_at) VALUES (?, ?, ?) "
        "ON CONFLICT(key) DO UPDATE SET value = excluded.value, updated_at = excluded.updated_at",
        (TARGET_SETTING_KEY, target_idle_pct, now_str),
    )
    conn.commit()
    return {"target_idle_pct": target_idle_pct, "updated_at": now_str}


# ---------------------------------------------------------------------------
# Ma trận máy x ngày
# ---------------------------------------------------------------------------


def _parse_capacities(capacities: Iterable[Any] | None) -> set[float] | None:
    values: set[float] = set()
    for raw in capacities or []:
        try:
            values.add(round(float(raw), 2))
        except (TypeError, ValueError):
            continue
    return values or None


def _build(from_date: str | None, to_date: str | None, tank_types: list[str] | None, capacities: Iterable[Any] | None) -> dict[str, Any]:
    """Tính toàn bộ ô + khoảng idle cho kỳ đang lọc — dùng chung cho API ma trận, Excel."""
    ctx = _context(from_date, to_date)
    master = _machine_master(ctx["conn"])
    days = [day for day in ctx["days"] if (not from_date or day >= from_date) and (not to_date or day <= to_date)]
    day_set = set(days)
    active = sorted({m for (m, day), cell in ctx["cells"].items() if day in day_set and cell["operating"] > 0})

    present_tanks = {master.get(m.lower(), {}).get("tank", UNCLASSIFIED_TANK_LABEL) for m in active}
    available_tank_types = [label for label in _TANK_ORDER if label in present_tanks]
    available_capacities = sorted({master[m.lower()]["capacity"] for m in active if m.lower() in master and master[m.lower()]["capacity"] is not None})
    tank_filter = set(tank_types) if tank_types else None
    capacity_filter = _parse_capacities(capacities)

    def keep(machine: str) -> bool:
        info = master.get(machine.lower(), {})
        if tank_filter is not None and info.get("tank", UNCLASSIFIED_TANK_LABEL) not in tank_filter:
            return False
        return capacity_filter is None or info.get("capacity") in capacity_filter

    stop_state = ctx["stop_state"]
    target = get_target()
    rows: list[dict[str, Any]] = []
    all_segments: list[dict[str, Any]] = []
    totals_keys = ("operating", "idle", "available", "explained")
    day_totals = {day: dict.fromkeys(totals_keys, 0.0) for day in days}
    for machine in (m for m in active if keep(m)):
        info = master.get(machine.lower(), {})
        row = {"machine": machine, "tank": info.get("tank", UNCLASSIFIED_TANK_LABEL), "capacity": info.get("capacity"), "cells": {}}
        row_totals = dict.fromkeys(totals_keys, 0.0)
        for day in days:
            cell = ctx["cells"][(machine, day)]
            gaps = ctx["day_gaps"].get((machine, day), [])
            idle, explained = cell.get("idle", 0.0), cell.get("explained", 0.0)
            for gap in gaps:
                for segment in gap["segments"]:
                    stop = stop_state[segment["stop_id"]]["stop"] if segment["stop_id"] is not None else None
                    all_segments.append({"machine": machine, "production_date": day, **segment, "stop": stop})
            row["cells"][day] = {
                "operating": round(cell["operating"], 2), "idle": round(idle, 2), "available": round(cell["available"], 2),
                "explained": round(explained, 2), "idle_pct": _pct(idle, cell["available"]), "gap_count": len(gaps),
            }
            values = {"operating": cell["operating"], "idle": idle, "available": cell["available"], "explained": explained}
            for totals in (row_totals, day_totals[day]):
                for key in totals_keys:
                    totals[key] += values[key]
        row.update({key: round(value, 2) for key, value in row_totals.items()})
        row["idle_pct"] = _pct(row_totals["idle"], row_totals["available"])
        rows.append(row)

    grand = {key: sum(totals[key] for totals in day_totals.values()) for key in totals_keys}
    windows = ctx["windows"]
    return {
        "days": days,
        "day_available_hours": {day: round(_hours(*windows[day]), 2) for day in days},
        "partial_days": [day for day in days if windows[day] != _day_window(day)],
        "rows": rows,
        "day_totals": {
            day: {**{key: round(value, 2) for key, value in totals.items()}, "idle_pct": _pct(totals["idle"], totals["available"])}
            for day, totals in day_totals.items()
        },
        "grand_total": {
            **{key: round(value, 2) for key, value in grand.items()},
            "idle_pct": _pct(grand["idle"], grand["available"]),
            "explained_pct": _pct(grand["explained"], grand["idle"]),
        },
        "target_idle_pct": target,
        "available_tank_types": available_tank_types,
        "available_capacities": available_capacities,
        "reasons": list(IDLE_REASONS),
        "_segments": all_segments,
    }


def get_idle_matrix(
    from_date: str | None = None, to_date: str | None = None,
    tank_types: list[str] | None = None, capacities: Iterable[Any] | None = None,
) -> dict[str, Any]:
    data = _build(from_date, to_date, tank_types, capacities)
    data.pop("_segments")
    return data


# ---------------------------------------------------------------------------
# Drill-down 1 ô
# ---------------------------------------------------------------------------


def _batch_dyelots(machine: str, start_times: set[str]) -> dict[str, str]:
    """{StartTime gốc: Dyelot} — chỉ để hiển thị mẻ trước/sau khoảng idle. Machine trong
    batch_details có thể trống (fill-down khi chuẩn hoá) nên nhận cả dòng machine rỗng."""
    if not start_times:
        return {}
    placeholders = ", ".join("?" for _ in start_times)
    rows = execute_query(f"SELECT dyelot, machine, start_time FROM batch_details WHERE start_time IN ({placeholders})", list(start_times))
    result: dict[str, str] = {}
    for row in sorted(rows, key=lambda r: str(r["machine"] or "").strip().lower() != machine.lower()):
        raw_machine = str(row["machine"] or "").strip().lower()
        if raw_machine in ("", machine.lower()):
            result.setdefault(str(row["start_time"]), row["dyelot"])
    return result


def get_idle_gaps(machine: str, production_date: str) -> dict[str, Any]:
    """Khoảng idle của 1 ô (máy x ngày), mỗi khoảng chia thành đoạn theo lần dừng đã gán."""
    ctx = _context(production_date, production_date)
    window = ctx["windows"].get(production_date)
    cell = ctx["cells"].get((machine, production_date))
    gaps = ctx["day_gaps"].get((machine, production_date), [])
    if window is None or cell is None:
        return {"machine": machine, "production_date": production_date, "gaps": [], "reasons": list(IDLE_REASONS)}
    dyelots = _batch_dyelots(machine, {s for gap in gaps for s in (gap["prev_batch_start"], gap["next_batch_start"]) if s})
    stop_state = ctx["stop_state"]

    def segment_payload(segment: dict[str, Any]) -> dict[str, Any]:
        state = stop_state.get(segment["stop_id"]) if segment["stop_id"] is not None else None
        return {
            "start": _fmt(segment["start"]), "end": _fmt(segment["end"]), "hours": round(_hours(segment["start"], segment["end"]), 2),
            "stop": {**_stop_payload(state["stop"]), "deviation_minutes": state["deviation_minutes"]} if state else None,
        }

    return {
        "machine": machine,
        "production_date": production_date,
        "available_from": _fmt(window[0]),
        "available_to": _fmt(window[1]),
        "operating": round(cell["operating"], 2),
        "idle": round(cell.get("idle", 0.0), 2),
        "gaps": [
            {
                "start": _fmt(gap["start"]), "end": _fmt(gap["end"]), "hours": round(_hours(gap["start"], gap["end"]), 2),
                "gap_start": _fmt(gap["gap_start"]), "gap_end": _fmt(gap["gap_end"]),
                "prev_batch": dyelots.get(gap["prev_batch_start"] or "") or gap["prev_batch_start"],
                "next_batch": dyelots.get(gap["next_batch_start"] or "") or gap["next_batch_start"],
                "segments": [segment_payload(segment) for segment in gap["segments"]],
            }
            for gap in gaps
        ],
        "reasons": list(IDLE_REASONS),
    }


# ---------------------------------------------------------------------------
# Lần dừng — danh sách / tạo / sửa / xoá
# ---------------------------------------------------------------------------


def get_master_machines() -> list[str]:
    """Danh sách máy để chọn ở trang Idle Entry — Machine Master (domain dyeing)."""
    conn = get_db()
    rows = conn.execute("SELECT machine_code, machine_id FROM machines WHERE domain = 'dyeing'").fetchall()
    return sorted({str(r["machine_code"] or r["machine_id"]).strip() for r in rows if (r["machine_code"] or r["machine_id"])})


def list_stops(from_date: str, to_date: str, machine: str | None = None) -> list[dict[str, Any]]:
    """Lần dừng có giờ bắt đầu trong [from_date 07:00, to_date+1 07:00) + trạng thái ghép:
    pending (chờ Batch) / matched (kèm giờ chính xác đã gán) / unmatched (máy không có khoảng
    idle nào trong dữ liệu) / overridden / no_time (bị lần dừng sau chiếm hết khoảng idle)."""
    ctx = _context(from_date, to_date)
    range_start, range_end = _day_window(from_date)[0], _day_window(to_date)[1]
    rows = _load_stops(range_start, range_end)
    result = []
    for stop in rows:
        start = _parse_input_datetime(stop["stop_start"])
        if start is None or not range_start <= start < range_end:
            continue
        if machine and str(stop["machine"]).strip().lower() != machine.strip().lower():
            continue
        state = ctx["stop_state"].get(stop["id"], {})
        status = state.get("status", "pending")
        if status == "unmatched" and not any(m.lower() == str(stop["machine"]).strip().lower() for m in ctx["machines"]):
            status = "pending"  # máy chưa có mẻ nào trong dữ liệu đã nạp
        segments = state.get("segments", [])
        result.append({
            **_stop_payload(stop),
            "status": status,
            "partially_overridden": state.get("partially_overridden", False),
            "deviation_minutes": state.get("deviation_minutes"),
            "matched": [{"start": _fmt(s), "end": _fmt(e), "hours": round(_hours(s, e), 2)} for s, e in segments],
            "matched_hours": round(sum(_hours(s, e) for s, e in segments), 2),
        })
    result.sort(key=lambda item: (item["stop_start"], item["machine"]), reverse=True)
    return result


def _validate_stop(machine: str, stop_start: Any, stop_end: Any, reason: str | None, source: str) -> tuple[str, str, str | None]:
    machine = (machine or "").strip()
    if not machine:
        raise ValueError("Chưa chọn máy.")
    if source == SOURCE_ENTRY:
        master = {code.lower(): code for code in get_master_machines()}
        if machine.lower() not in master:
            raise ValueError(f"Máy {machine} không có trong Machine Master.")
        machine = master[machine.lower()]
    start = _parse_input_datetime(stop_start)
    if start is None:
        raise ValueError("Thời điểm bắt đầu dừng không hợp lệ.")
    end = _parse_input_datetime(stop_end)
    if stop_end and end is None:
        raise ValueError("Thời điểm kết thúc dừng không hợp lệ.")
    if end is not None and end <= start:
        raise ValueError("Thời điểm kết thúc phải sau thời điểm bắt đầu.")
    if reason not in IDLE_REASONS:
        raise ValueError("Chưa chọn Reason hợp lệ.")
    return machine, _fmt(start), _fmt(end) if end else None


def create_stop(machine: str, stop_start: Any, stop_end: Any, reason: str | None, detail: str | None, user_id: int, source: str = SOURCE_ENTRY) -> dict[str, Any]:
    machine, start, end = _validate_stop(machine, stop_start, stop_end, reason, source)
    conn = get_db()
    _ensure_tables(conn)
    now_str = datetime.now().strftime(_TIME_FMT)
    stop_id = insert_returning_id(
        conn,
        "INSERT INTO idle_time_stops (machine, stop_start, stop_end, reason, detail, source, created_by, created_at, updated_by, updated_at) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (machine, start, end, reason, detail, source, user_id, now_str, user_id, now_str),
    )
    conn.commit()
    return _stop_payload(_load_stops(None, None, stop_id)[0])


def update_stop(stop_id: int, machine: str, stop_start: Any, stop_end: Any, reason: str | None, detail: str | None, user_id: int) -> dict[str, Any]:
    conn = get_db()
    _ensure_tables(conn)
    existing = _load_stops(None, None, stop_id)
    if not existing:
        raise ValueError(f"Không tìm thấy lần dừng id={stop_id}.")
    machine, start, end = _validate_stop(machine, stop_start, stop_end, reason, existing[0]["source"])
    conn.execute(
        "UPDATE idle_time_stops SET machine = ?, stop_start = ?, stop_end = ?, reason = ?, detail = ?, updated_by = ?, updated_at = ? WHERE id = ?",
        (machine, start, end, reason, detail, user_id, datetime.now().strftime(_TIME_FMT), stop_id),
    )
    conn.commit()
    return _stop_payload(_load_stops(None, None, stop_id)[0])


def delete_stop(stop_id: int) -> None:
    conn = get_db()
    _ensure_tables(conn)
    if not _load_stops(None, None, stop_id):
        raise ValueError(f"Không tìm thấy lần dừng id={stop_id}.")
    conn.execute("DELETE FROM idle_time_stops WHERE id = ?", (stop_id,))
    conn.commit()


def save_from_report(machine: str, segment_start: str, segment_end: str, reason: str | None, detail: str | None, user_id: int) -> dict[str, Any]:
    """Save trong drawer báo cáo = lần dừng source='report' có giờ ĐÚNG bằng đoạn idle đang sửa.
    Nếu đoạn đó đã do 1 lần dừng 'report' cùng giờ tạo ra -> sửa chính lần đó (không nhân bản)."""
    conn = get_db()
    _ensure_tables(conn)
    same = execute_query(
        "SELECT id FROM idle_time_stops WHERE machine = ? AND source = ? AND stop_start = ? AND stop_end = ?",
        [machine, SOURCE_REPORT, segment_start, segment_end],
    )
    if same:
        return update_stop(same[0]["id"], machine, segment_start, segment_end, reason, detail, user_id)
    return create_stop(machine, segment_start, segment_end, reason, detail, user_id, source=SOURCE_REPORT)


# ---------------------------------------------------------------------------
# Excel export — y hệt bảng trên UI + sheet chi tiết đoạn idle kèm nguyên nhân
# ---------------------------------------------------------------------------


def export_idle_excel(
    from_date: str | None = None, to_date: str | None = None,
    tank_types: list[str] | None = None, capacities: Iterable[Any] | None = None, mode: str = "idle",
) -> bytes:
    """Sheet "Idle Time": máy x ngày theo `mode` ("idle" = giờ idle, "operating" = giờ hoạt
    động) + Total/% Idle/Target, dòng Total/% Idle/Available (h) theo ngày; ô có % Idle vượt
    Target tô màu cảnh báo. Sheet "Idle Gaps": từng đoạn idle (giờ chính xác từ Batch) + nguyên
    nhân + giờ người dùng đã nhập."""
    data = _build(from_date, to_date, tank_types, capacities)
    value_key = "operating" if mode == "operating" else "idle"
    target = data["target_idle_pct"]
    days = data["days"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="24292F", end_color="24292F", fill_type="solid")
    bold = Font(bold=True)

    def over_target(pct: float | None) -> bool:
        return target is not None and pct is not None and pct > target

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Idle Time"
    headers = ["Machine", "Tank", "Capacity (Kg)", *days, "Total (h)", "% Idle", "Target % Idle"]
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        sheet.column_dimensions[cell.column_letter].width = 12
    first_day_col = 4
    total_col = first_day_col + len(days)
    for row_idx, row in enumerate(data["rows"], start=2):
        sheet.cell(row=row_idx, column=1, value=row["machine"])
        sheet.cell(row=row_idx, column=2, value=row["tank"])
        sheet.cell(row=row_idx, column=3, value=row["capacity"])
        for offset, day in enumerate(days):
            cell_data = row["cells"][day]
            cell = sheet.cell(row=row_idx, column=first_day_col + offset, value=round(cell_data[value_key], 1))
            if over_target(cell_data["idle_pct"]):
                cell.fill, cell.font = _EXCEL_OVER_TARGET_FILL, _EXCEL_OVER_TARGET_FONT
        sheet.cell(row=row_idx, column=total_col, value=round(row[value_key], 1))
        pct_cell = sheet.cell(row=row_idx, column=total_col + 1, value=row["idle_pct"])
        if over_target(row["idle_pct"]):
            pct_cell.fill, pct_cell.font = _EXCEL_OVER_TARGET_FILL, _EXCEL_OVER_TARGET_FONT
        sheet.cell(row=row_idx, column=total_col + 2, value=target)

    footer_row = len(data["rows"]) + 2
    for label_offset, (label, key) in enumerate((("Total", value_key), ("% Idle", "idle_pct"), ("Available (h/machine)", None))):
        row_idx = footer_row + label_offset
        sheet.cell(row=row_idx, column=1, value=label).font = bold
        for offset, day in enumerate(days):
            value = data["day_available_hours"][day] if key is None else data["day_totals"][day][key]
            cell = sheet.cell(row=row_idx, column=first_day_col + offset, value=round(value, 1) if key == value_key else value)
            cell.font = bold
            if key == "idle_pct" and over_target(value):
                cell.fill, cell.font = _EXCEL_OVER_TARGET_FILL, _EXCEL_OVER_TARGET_FONT
        if key == value_key:
            sheet.cell(row=row_idx, column=total_col, value=round(data["grand_total"][value_key], 1)).font = bold
        elif key == "idle_pct":
            sheet.cell(row=row_idx, column=total_col + 1, value=data["grand_total"]["idle_pct"]).font = bold
    sheet.freeze_panes = sheet.cell(row=2, column=first_day_col)

    gap_sheet = workbook.create_sheet("Idle Gaps")
    gap_headers = ["Machine", "Production Date", "From", "To", "Idle (h)", "Reason", "Detail", "Source", "Input From", "Input To", "Updated by", "Updated at"]
    for col, header in enumerate(gap_headers, start=1):
        cell = gap_sheet.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        gap_sheet.column_dimensions[cell.column_letter].width = 18
    for row_idx, segment in enumerate(data["_segments"], start=2):
        stop = segment["stop"]
        values = [
            segment["machine"], segment["production_date"], _fmt(segment["start"]), _fmt(segment["end"]),
            round(_hours(segment["start"], segment["end"]), 2),
            stop["reason"] if stop else None, stop["detail"] if stop else None,
            ("Idle Entry" if stop["source"] == SOURCE_ENTRY else "Report") if stop else None,
            stop["stop_start"] if stop else None, stop["stop_end"] if stop else None,
            stop["updated_by"] if stop else None, stop["updated_at"] if stop else None,
        ]
        for col, value in enumerate(values, start=1):
            gap_sheet.cell(row=row_idx, column=col, value=value)
    gap_sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()

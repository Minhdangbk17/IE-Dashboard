"""Batch/Day Trend — tab thứ 2 của Engine `batch_matrix` (theo yêu cầu người dùng).

CÔNG THỨC (mỗi kỳ Day/Week/Month, tính riêng từng loại vải Cotton/CVC/Polyester):
    Batch/Day = [Số mẻ Normal dyeing] * 24 / [Tổng Occupied Hours của TẤT CẢ mẻ trong kỳ]

2026-09-29 — VIẾT LẠI HOÀN TOÀN theo Power Query người dùng cung cấp (thay thế bản cũ dùng
`availability_logs.planned_prd_time_hour` + carry-forward giờ của mẻ FabricType rỗng):

1. **Nguồn dữ liệu DUY NHẤT: `batch_details`** (sheet "Batch") — BỎ HẲN `availability_logs`.
   Bước 2-3 + ngày kết thúc nằm ở `core/batch_source.py::load_resolved_batches()` (2026-10-02),
   dùng CHUNG với "Batch Per Day by Machine" để 2 báo cáo luôn cùng 1 tập mẻ.
2. **Machine trống** -> lấy Machine của dòng NGAY PHÍA TRÊN (fill-down theo thứ tự `id`,
   tức thứ tự dòng lúc import — khớp `Table.FillDown` trên sheet gốc).
3. **FabricType trống/"Unknown"** -> lấy FabricType của mẻ KẾ TIẾP trên CÙNG máy (sắp theo
   StartTime, khớp `Table.FillUp`); mẻ cuối cùng của máy không có mẻ sau -> giữ "Unknown"
   (bị loại khỏi báo cáo vì không thuộc 3 loại vải chính). Dòng được điền này VẪN là 1 mẻ
   riêng (giờ + đếm của CHÍNH nó), KHÔNG còn cộng dồn giờ sang mẻ sau như bản cũ.
4. **Occupied Hours**: mỗi mẻ được TÁCH theo khung Production Date 07:00 -> 07:00 (không giới
   hạn số ngày); giờ mỗi đoạn = thời gian THỰC (EndTime - StartTime) nằm trong khung đó, KHÔNG
   dùng cột RunTime. Đoạn 0h (EndTime đúng 07:00:00) bị bỏ. Mẻ thiếu/ngược Start-End bị bỏ.
5. **Đếm mẻ (tử số)**: mẻ Normal dyeing theo ĐÚNG quy tắc Dyelot/SapLot/ReDye của báo cáo
   "Batch Per Day by Machine" (`reports/cleaning_matrix.py::classify_batch_badge()`; điều kiện
   ReDye = 0 bật/tắt qua checkbox "ReDye = 0" trên UI, mặc định BẬT) — KHÔNG phải CM ("WA"/"CL*"),
   KHÔNG phải Sample ("DU"/"KN"/SapLot 4*), KHÔNG phải Rework. Mẻ chạy qua nhiều ngày CHỈ đếm
   1 lần vào NGÀY KẾT THÚC (= Production Date của đoạn cuối cùng có giờ > 0).
6. **Mẫu số**: Occupied Hours của TẤT CẢ mẻ (kể cả CM/Sample/Rework) cộng theo từng đoạn ngày.

DAILY ROLLUP PATTERN (`batch_day_trend_daily_summary`, grain = 1 dòng/1 ĐOẠN của 1 mẻ trong 1
production_date): `hours` = Occupied Hours của đoạn đó, `is_valid` = số mẻ Normal được ĐẾM
tại đoạn đó (1 ở đoạn ngày kết thúc của mẻ Normal, 0 ở mọi đoạn khác), `is_valid_any_redye` =
như `is_valid` nhưng BỎ điều kiện ReDye = 0 (checkbox TẮT). Cột `is_valid_any_redye` cần
migration Postgres `supabase/migrate_batch_day_trend_redye.sql`; sau khi deploy PHẢI chạy
`flask rebuild-summaries` để tính lại toàn bộ lịch sử.

Fill-down Machine/fill-up FabricType phụ thuộc THỨ TỰ xuyên suốt lịch sử (không tách rời theo
từng ngày) nên `recompute_all()` đọc TOÀN BỘ `batch_details` 1 LẦN (1 round-trip, chỉ các cột
cần dùng), chạy thuật toán ở Python rồi CHỈ lưu các đoạn thuộc tập ngày được yêu cầu.
`recompute_daily(D)` = `recompute_all([D])`.

Bộ lọc: Capacity (Kg) — `machines.capacity_kg` đã cấu hình qua UI Batch Per Day by Machine
(máy CHƯA cấu hình bị ẩn khi lọc Capacity cụ thể). Brand - Program — suy từ `greige_code`.
Tank Type (2026-09-29) — tra `machines.tank_type` tại read time (xem `_machine_tank_labels()`).
"""
from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from core.batch_source import INVALID_FABRIC_TYPES, load_resolved_batches, parse_batch_datetime, split_production_days
from core.brand_program_importer import ensure_brand_program_table
from core.database import DatabaseError, execute_query, get_db, get_dialect
from modules.dyeing.engines.reports.cleaning_matrix import (
    CANONICAL_TANK_ORDER,
    UNCLASSIFIED_TANK_LABEL,
    _normalize_tank_label,
    classify_batch_badge,
)

# Báo cáo "Batch/Day Trend" LUÔN thể hiện ĐÚNG 3 loại vải chính này (theo yêu cầu người
# dùng — 1 đường/1 dòng cố định mỗi loại, không phụ thuộc filter hay dữ liệu thực tế đang
# có). Mọi `fabric_type` KHÁC 3 loại này (VD Nylon, Spandex...) bị loại khỏi báo cáo Trend
# hoàn toàn — khác các báo cáo khác trong `batch_matrix`/`downtime`/`rft` vốn hiển thị ĐỘNG
# theo `available_fabric_types` tìm thấy trong dữ liệu. So khớp không phân biệt hoa/thường
# (xem `_normalize_main_fabric_type()`) — CHƯA xác nhận dữ liệu thật có biến thể viết khác
# (VD viết tắt "PES" cho Polyester) hay không, nếu có sẽ cần bổ sung alias sau.
MAIN_FABRIC_TYPES: tuple[str, ...] = ("Cotton", "CVC", "Polyester")
_MAIN_FABRIC_TYPE_BY_NORM: dict[str, str] = {name.lower(): name for name in MAIN_FABRIC_TYPES}


def _normalize_main_fabric_type(value: str | None) -> str | None:
    """Khớp `fabric_type` thô (bất kỳ hoa/thường) với ĐÚNG 1 trong `MAIN_FABRIC_TYPES` —
    trả `None` nếu không khớp loại nào (loại đó bị loại khỏi báo cáo Trend)."""
    return _MAIN_FABRIC_TYPE_BY_NORM.get(str(value or "").strip().lower())
def _parse_brand_programs(brand_programs: str | list[str] | None) -> list[str]:
    if not brand_programs:
        return []
    values = brand_programs if isinstance(brand_programs, list) else str(brand_programs).split(",")
    return [text for value in values if (text := str(value).strip())]


def _parse_capacities(capacities: str | list[str] | None) -> list[float]:
    if not capacities:
        return []
    values = capacities if isinstance(capacities, list) else str(capacities).split(",")
    parsed: list[float] = []
    for value in values:
        text = str(value).strip()
        if not text:
            continue
        try:
            number = float(text)
        except ValueError:
            continue
        if number not in parsed:
            parsed.append(number)
    return parsed


def _period(day: date, group_by: str) -> tuple[str, str]:
    if group_by == "month":
        return day.strftime("%Y-%m"), day.strftime("%b-%y")
    if group_by == "week":
        year, week, _ = day.isocalendar()
        return f"{year}-W{week:02d}", f"W{week:02d} {year}"
    return day.isoformat(), day.strftime("%d %b %Y")


# ---------------------------------------------------------------------------
# Daily Rollup: recompute_daily()
# ---------------------------------------------------------------------------


def _ensure_summary_table(conn: Any) -> None:
    """CHỈ chạy CREATE TABLE ở SQLite — ở Postgres bảng đã có sẵn qua `supabase/schema.sql`
    (DB production đã có dữ liệu thì áp `supabase/migrate_batch_day_trend_brand_program.sql`
    + `supabase/migrate_batch_day_trend_redye.sql` rồi chạy lại `flask rebuild-summaries`)."""
    if get_dialect() == "sqlite":
        existing_columns = {row["name"] for row in conn.execute("PRAGMA table_info(batch_day_trend_daily_summary)")}
        if existing_columns and not {"brand_program", "is_valid_any_redye"} <= existing_columns:
            # An toàn xoá/tạo lại (giống `batch_matrix/service.py::_ensure_summary_table()`)
            # vì bảng luôn tái tạo được 100% từ raw data qua `recompute_daily()`.
            conn.execute("DROP TABLE batch_day_trend_daily_summary")
        conn.execute("""
            CREATE TABLE IF NOT EXISTS batch_day_trend_daily_summary (
                production_date TEXT NOT NULL,
                machine TEXT NOT NULL,
                start_time TEXT NOT NULL,
                fabric_type TEXT NOT NULL,
                capacity_kg REAL,
                brand_program TEXT NOT NULL DEFAULT '',
                hours REAL NOT NULL DEFAULT 0,
                is_valid INTEGER NOT NULL DEFAULT 0,
                is_valid_any_redye INTEGER NOT NULL DEFAULT 0,
                updated_at TEXT NOT NULL DEFAULT (datetime('now')),
                PRIMARY KEY (production_date, machine, start_time)
            )
        """)
    conn.execute("CREATE INDEX IF NOT EXISTS idx_batch_day_trend_daily_summary_date ON batch_day_trend_daily_summary (production_date)")


def _machine_capacities(conn: Any) -> dict[str, float]:
    """{mã máy (lower/trim): capacity_kg} — capacity cấu hình qua UI Batch Per Day by Machine."""
    capacities: dict[str, float] = {}
    for row in conn.execute("SELECT machine_code, machine_id, capacity_kg FROM machines").fetchall():
        code = str(row["machine_code"] or row["machine_id"] or "").strip().lower()
        if code and row["capacity_kg"] not in (None, "") and code not in capacities:
            capacities[code] = round(float(row["capacity_kg"]), 2)
    return capacities


def _is_normal_batch(row: dict[str, Any], require_redye_zero: bool) -> bool:
    """Normal dyeing = KHÔNG phải CM/Sample/Rework theo `classify_batch_badge()` của báo cáo
    "Batch Per Day by Machine" — `require_redye_zero` khớp checkbox "ReDye = 0" trên UI."""
    badge = classify_batch_badge(
        {"dyelot": row["dyelot"], "sap_lot": row["sap_lot"], "redye": row["redye"]},
        require_redye_zero=require_redye_zero,
    )
    return badge not in ("CM", "S") and not badge.endswith("R")


def _brand_program_label(batch: dict[str, Any]) -> str:
    """Nhãn "Brand - Program" (VD "UQ - Ht Fleece") — rỗng nếu greige_code chưa map được.
    AN TOÀN lưu thẳng vào grain summary: mỗi dòng LÀ 1 mẻ cụ thể, lọc/gộp không cộng trùng giờ."""
    brand = str(batch["brand"] or "").strip()
    program = str(batch["brand_program"] or "").strip()
    return f"{brand} - {program}" if brand and program else ""


def _build_segments(conn: Any) -> list[dict[str, Any]]:
    """Tách từng mẻ ĐÃ CHUẨN HOÁ (`core/batch_source.py::load_resolved_batches()` — nguồn sự
    thật chung với Batch Per Day by Machine: fill-down Machine, fill-up FabricType) thành các
    đoạn theo Production Date, đếm mẻ Normal ở đoạn cuối (ngày kết thúc). Trả danh sách đoạn
    (1 đoạn = 1 mẻ x 1 production_date) CHƯA gộp — xem docstring đầu file bước 4-6."""
    capacities = _machine_capacities(conn)
    segments: list[dict[str, Any]] = []
    for batch in load_resolved_batches(conn):
        parts = batch["_segments"]
        if not parts:
            continue
        is_normal = _is_normal_batch(batch, require_redye_zero=True)
        is_normal_any_redye = _is_normal_batch(batch, require_redye_zero=False)
        for index, (day, hours) in enumerate(parts):
            is_end_day = index == len(parts) - 1
            segments.append({
                "production_date": day.isoformat(),
                "machine": batch["machine"],
                "start_time": str(batch["start_time"]),
                "fabric_type": batch["fabric_type"],
                "capacity_kg": capacities.get(batch["machine"].lower()),
                "brand_program": _brand_program_label(batch),
                "hours": hours,
                "is_valid": int(is_normal and is_end_day),
                "is_valid_any_redye": int(is_normal_any_redye and is_end_day),
            })
    return segments


def recompute_daily(production_date: date, conn: Any) -> None:
    """Tính lại `batch_day_trend_daily_summary` cho ĐÚNG 1 production_date — idempotent."""
    recompute_all([production_date], conn)


def recompute_all(dates: Iterable[date], conn: Any) -> None:
    """Tính lại `batch_day_trend_daily_summary` cho tập `dates` — idempotent (DELETE rồi
    INSERT lại đúng các ngày này). Đọc `batch_details` ĐÚNG 1 LẦN cho cả tập ngày (tránh bug
    2026-09-18: gọi lặp từng ngày trên Postgres remote chậm tới mức bị coi là "treo")."""
    _ensure_summary_table(conn)
    ensure_brand_program_table(conn)
    date_strs = {d.isoformat() for d in dates}
    if not date_strs:
        return

    # Gộp theo PRIMARY KEY (production_date, machine, start_time) phòng trường hợp 2 dòng
    # batch_details trùng máy + StartTime (khác Dyelot) — cộng giờ/số mẻ, không lỗi trùng khoá.
    merged: dict[tuple[str, str, str], dict[str, Any]] = {}
    for segment in _build_segments(conn):
        if segment["production_date"] not in date_strs:
            continue
        key = (segment["production_date"], segment["machine"], segment["start_time"])
        if key in merged:
            merged[key]["hours"] += segment["hours"]
            merged[key]["is_valid"] += segment["is_valid"]
            merged[key]["is_valid_any_redye"] += segment["is_valid_any_redye"]
        else:
            merged[key] = segment

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for day_str in date_strs:
        conn.execute("DELETE FROM batch_day_trend_daily_summary WHERE production_date = ?", (day_str,))
    if merged:
        conn.executemany(
            """
            INSERT INTO batch_day_trend_daily_summary
                (production_date, machine, start_time, fabric_type, capacity_kg, brand_program, hours,
                 is_valid, is_valid_any_redye, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (s["production_date"], s["machine"], s["start_time"], s["fabric_type"], s["capacity_kg"],
                 s["brand_program"], s["hours"], s["is_valid"], s["is_valid_any_redye"], now_str)
                for s in merged.values()
            ],
        )


# ---------------------------------------------------------------------------
# Target — bảng cấu hình `batch_day_trend_targets`, khoá theo `fabric_type` (CHỈ 3 giá trị
# `MAIN_FABRIC_TYPES`). Cùng pattern `downtime_targets`/`batch_matrix_targets` đã có sẵn
# trong dự án — mỗi báo cáo 1 bảng riêng, khoá theo đúng "đơn vị hàng" của báo cáo đó.
# KHÔNG tái dùng `batch_matrix_targets` (khoá `(fabric_type, color_group)`) — 2 bảng trả lời
# 2 câu hỏi khác nhau ("target cho 1 ô fabric+color của Matrix" vs "target cho tổng
# Batch/Day của 1 loại vải ở Trend"), đơn vị đo cũng khác quy mô nhau dù cùng công thức gốc.
# ---------------------------------------------------------------------------


def _ensure_trend_targets_table(conn: Any) -> None:
    """CHỈ chạy CREATE TABLE ở SQLite — Postgres tạo qua `supabase/schema.sql`."""
    if get_dialect() == "sqlite":
        conn.execute("""
            CREATE TABLE IF NOT EXISTS batch_day_trend_targets (
                fabric_type TEXT PRIMARY KEY,
                target_value REAL NOT NULL DEFAULT 0,
                updated_at TEXT NOT NULL DEFAULT (datetime('now'))
            )
        """)
    conn.commit()


def get_trend_targets() -> dict[str, float | None]:
    """{fabric_type: target_value} cho ĐÚNG 3 loại `MAIN_FABRIC_TYPES` — loại CHƯA từng
    được cấu hình trả `None` (hiển thị "-" trên UI, phân biệt với target THẬT SỰ = 0)."""
    conn = get_db()
    _ensure_trend_targets_table(conn)
    rows = execute_query("SELECT fabric_type, target_value FROM batch_day_trend_targets")
    result: dict[str, float | None] = {row["fabric_type"]: float(row["target_value"]) for row in rows}
    for name in MAIN_FABRIC_TYPES:
        result.setdefault(name, None)
    return result


def set_trend_target(fabric_type: str, target_value: float) -> dict[str, Any]:
    if fabric_type not in MAIN_FABRIC_TYPES:
        raise ValueError(f"Fabric type không hợp lệ: {fabric_type}")
    conn = get_db()
    _ensure_trend_targets_table(conn)
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn.execute(
        "INSERT INTO batch_day_trend_targets (fabric_type, target_value, updated_at) VALUES (?, ?, ?) "
        "ON CONFLICT(fabric_type) DO UPDATE SET target_value = excluded.target_value, updated_at = excluded.updated_at",
        (fabric_type, target_value, now_str),
    )
    conn.commit()
    return {"fabric_type": fabric_type, "target_value": target_value, "updated_at": now_str}


# ---------------------------------------------------------------------------
# Đọc báo cáo — CHỈ SELECT từ `batch_day_trend_daily_summary`, lọc/gộp ở Python.
# ---------------------------------------------------------------------------


# Tank Type — tra `machines.tank_type` (Machine Master) NGAY TẠI READ TIME (không lưu vào
# bảng summary: không cần migration/rebuild, đổi Tank Type trên Machine Master có hiệu lực
# ngay). Chuẩn hoá giống tab Summary của Batch Per Day by Machine: "J tank"/"O tank", mọi giá
# trị khác (rỗng/lạ/máy chưa khai báo) -> "Unclassified".
_TANK_ORDER: tuple[str, ...] = (*CANONICAL_TANK_ORDER, UNCLASSIFIED_TANK_LABEL)


def _machine_tank_labels(conn: Any) -> dict[str, str]:
    """{mã máy (lower/trim): nhãn Tank đã chuẩn hoá}."""
    labels: dict[str, str] = {}
    for row in conn.execute("SELECT machine_code, machine_id, tank_type FROM machines").fetchall():
        code = str(row["machine_code"] or row["machine_id"] or "").strip().lower()
        if code and code not in labels:
            labels[code] = _normalize_tank_label(row["tank_type"])
    return labels


def _row_tank(row: Any, tank_by_machine: dict[str, str]) -> str:
    return tank_by_machine.get(str(row["machine"] or "").strip().lower(), UNCLASSIFIED_TANK_LABEL)


def _empty_trend(group_by: str, targets: dict[str, float | None] | None = None) -> dict[str, Any]:
    targets = targets or {}
    return {
        "filters": {"capacities": [], "brand_programs": [], "tank_types": [], "from_date": None, "to_date": None, "group_by": group_by},
        "periods": [], "period_keys": [],
        "rows": [
            {"fabric_type": name, "target": targets.get(name), "values": [], "total": 0.0,
             "counts": [], "hours": [], "total_count": 0, "total_hours": 0.0}
            for name in MAIN_FABRIC_TYPES
        ],
        "chart": {"categories": [], "series": []},
        "kpis": {"batch_per_day": 0.0, "valid_batches": 0, "total_planned_hours": 0.0},
        "available_capacities": [],
        "available_brand_programs": [],
        "available_tank_types": [],
    }


def _load_filtered_rows(
    capacities: str | list[str] | None,
    brand_programs: str | list[str] | None,
    tank_types: str | list[str] | None,
    from_date: str | None, to_date: str | None,
    group_by: str,
    require_redye_zero: bool,
) -> dict[str, Any] | None:
    """Đọc `batch_day_trend_daily_summary` + áp bộ lọc — dùng CHUNG cho bảng Trend, drill-down
    danh sách mẻ và Excel export để cả 3 LUÔN cùng 1 tập đoạn (mẻ x ngày). Trả `None` khi bảng
    chưa có/không có dữ liệu. Mỗi phần tử `rows` = 1 đoạn đã qua filter, kèm loại vải chính,
    kỳ (`period_key`/`period_label`) và nhãn Tank."""
    conn = get_db()
    # Checkbox "ReDye = 0": BẬT -> `is_valid` (Normal có ReDye = 0), TẮT -> `is_valid_any_redye`
    # (bỏ điều kiện ReDye). Cả 2 cột đã tính sẵn lúc rollup — không cần recompute khi đổi.
    count_column = "is_valid" if require_redye_zero else "is_valid_any_redye"
    sql = (
        "SELECT production_date, machine, start_time, fabric_type, capacity_kg, brand_program, hours, "
        f"{count_column} AS batch_count FROM batch_day_trend_daily_summary WHERE 1=1"
    )
    params: list[Any] = []
    if from_date:
        sql += " AND production_date >= ?"
        params.append(from_date)
    if to_date:
        sql += " AND production_date <= ?"
        params.append(to_date)
    try:
        _ensure_summary_table(conn)
        rows = execute_query(sql, params)
    except DatabaseError:
        # Bảng `batch_day_trend_daily_summary` chưa tồn tại trên Postgres (chưa áp
        # `supabase/migrate_batch_day_trend.sql`) hoặc lỗi tương tự — trả kết quả rỗng
        # thay vì crash cả trang, giống cách `rft`/`tank_loading` xử lý.
        return None
    if not rows:
        return None

    tank_by_machine = _machine_tank_labels(conn)
    available_capacities = sorted({round(float(row["capacity_kg"]), 2) for row in rows if row["capacity_kg"] not in (None, "")})
    available_brand_programs = sorted({row["brand_program"] for row in rows if row["brand_program"]})
    present_tanks = {_row_tank(row, tank_by_machine) for row in rows}
    available_tank_types = [label for label in _TANK_ORDER if label in present_tanks]

    selected_capacities = _parse_capacities(capacities)
    capacity_filter = {round(value, 2) for value in selected_capacities} if selected_capacities else None
    selected_brand_programs = _parse_brand_programs(brand_programs)
    brand_program_filter = set(selected_brand_programs) if selected_brand_programs else None
    selected_tank_types = [value for value in _parse_brand_programs(tank_types) if value in _TANK_ORDER]
    tank_filter = set(selected_tank_types) if selected_tank_types else None

    # Loại KHÁC 3 loại chính (`_normalize_main_fabric_type()` trả None) bị loại bỏ hoàn toàn
    # khỏi báo cáo Trend.
    filtered: list[dict[str, Any]] = []
    for row in rows:
        fabric_type = _normalize_main_fabric_type(row["fabric_type"])
        if fabric_type is None:
            continue
        if capacity_filter is not None:
            row_capacity = round(float(row["capacity_kg"]), 2) if row["capacity_kg"] not in (None, "") else None
            if row_capacity not in capacity_filter:
                continue
        if brand_program_filter is not None and (row["brand_program"] or "") not in brand_program_filter:
            continue
        tank = _row_tank(row, tank_by_machine)
        if tank_filter is not None and tank not in tank_filter:
            continue
        day = datetime.strptime(row["production_date"], "%Y-%m-%d").date()
        key, label = _period(day, group_by)
        filtered.append({"row": row, "fabric_type": fabric_type, "period_key": key, "period_label": label, "tank": tank})

    return {
        "rows": filtered,
        "available_capacities": available_capacities,
        "available_brand_programs": available_brand_programs,
        "available_tank_types": available_tank_types,
        "selected_capacities": selected_capacities,
        "selected_brand_programs": selected_brand_programs,
        "selected_tank_types": selected_tank_types,
    }


def get_batch_day_trend(
    capacities: str | list[str] | None = None,
    brand_programs: str | list[str] | None = None,
    tank_types: str | list[str] | None = None,
    from_date: str | None = None, to_date: str | None = None,
    group_by: str = "date",
    require_redye_zero: bool = True,
) -> dict[str, Any]:
    group_by = group_by if group_by in {"date", "week", "month"} else "date"
    trend_targets = get_trend_targets()
    loaded = _load_filtered_rows(capacities, brand_programs, tank_types, from_date, to_date, group_by, require_redye_zero)
    if loaded is None:
        return _empty_trend(group_by, trend_targets)
    available_capacities = loaded["available_capacities"]
    available_brand_programs = loaded["available_brand_programs"]
    available_tank_types = loaded["available_tank_types"]
    selected_capacities = loaded["selected_capacities"]
    selected_brand_programs = loaded["selected_brand_programs"]
    selected_tank_types = loaded["selected_tank_types"]

    # Mỗi loại vải chính (Cotton/CVC/Polyester) có tử số/mẫu số RIÊNG — tính ĐỘC LẬP theo
    # ĐÚNG công thức gốc (đếm mẻ Normal * 24 / tổng Occupied Hours TẤT CẢ mẻ), không dùng chung mẫu
    # số như bản 1-đường-gộp cũ.
    period_counts: dict[str, dict[str, int]] = {name: {} for name in MAIN_FABRIC_TYPES}
    period_hours: dict[str, dict[str, float]] = {name: {} for name in MAIN_FABRIC_TYPES}
    period_labels: dict[str, str] = {}
    for item in loaded["rows"]:
        row, fabric_type, key = item["row"], item["fabric_type"], item["period_key"]
        period_labels[key] = item["period_label"]
        period_hours[fabric_type][key] = period_hours[fabric_type].get(key, 0.0) + float(row["hours"] or 0)
        if row["batch_count"]:
            period_counts[fabric_type][key] = period_counts[fabric_type].get(key, 0) + int(row["batch_count"])

    if not period_labels:
        result = _empty_trend(group_by, trend_targets)
        result["available_capacities"] = available_capacities
        result["available_brand_programs"] = available_brand_programs
        result["available_tank_types"] = available_tank_types
        return result

    ordered_keys = sorted(period_labels.keys())
    labels = [period_labels[key] for key in ordered_keys]

    rows_out: list[dict[str, Any]] = []
    total_count_all = 0
    total_hours_all = 0.0
    for name in MAIN_FABRIC_TYPES:
        counts = period_counts[name]
        hours = period_hours[name]
        values = [round(counts.get(key, 0) * 24 / hours[key], 2) if hours.get(key) else 0.0 for key in ordered_keys]
        total_count = sum(counts.values())
        total_hours = sum(hours.values())
        total_value = round(total_count * 24 / total_hours, 2) if total_hours else 0.0
        rows_out.append({
            "fabric_type": name, "target": trend_targets.get(name), "values": values, "total": total_value,
            # Tử số/mẫu số từng kỳ — Excel export dùng để đối chiếu lại từng ô.
            "counts": [counts.get(key, 0) for key in ordered_keys],
            "hours": [round(hours.get(key, 0.0), 2) for key in ordered_keys],
            "total_count": total_count, "total_hours": round(total_hours, 2),
        })
        total_count_all += total_count
        total_hours_all += total_hours

    overall_value = round(total_count_all * 24 / total_hours_all, 2) if total_hours_all else 0.0

    return {
        "filters": {
            "capacities": selected_capacities or "all",
            "brand_programs": selected_brand_programs or "all",
            "tank_types": selected_tank_types or "all",
            "require_redye_zero": require_redye_zero,
            "from_date": from_date, "to_date": to_date, "group_by": group_by,
        },
        "periods": labels, "period_keys": ordered_keys,
        "rows": rows_out,
        "chart": {"categories": labels, "series": [{"name": row["fabric_type"], "data": row["values"]} for row in rows_out]},
        "kpis": {
            "batch_per_day": overall_value,
            "valid_batches": total_count_all,
            "total_planned_hours": round(total_hours_all, 1),
        },
        "available_capacities": available_capacities,
        "available_brand_programs": available_brand_programs,
        "available_tank_types": available_tank_types,
    }


# ---------------------------------------------------------------------------
# Drill-down danh sách mẻ + Excel export (2026-10-06, yêu cầu người dùng): bấm 1 ô (loại vải x
# kỳ) hoặc ô Total -> liệt kê TẤT CẢ đoạn (mẻ x ngày) góp giờ vào mẫu số, cờ "Counted" = mẻ được
# đếm vào tử số. Grain = 1 dòng/1 đoạn của summary (mẻ chạy qua nhiều ngày hiện nhiều dòng, chỉ
# dòng ngày kết thúc có Counted = Yes) nên tổng giờ/số mẻ Counted khớp ĐÚNG con số trên web.
# Dyelot/SapLot/ReDye/End không lưu trong summary -> tra ngược `batch_details` theo StartTime
# (+ Machine; dòng Machine trống ở raw được fill-down nên khớp theo StartTime).
# ---------------------------------------------------------------------------

_DETAIL_QUERY_CHUNK = 500
_EXCEL_HEADER_FONT = Font(bold=True, color="FFFFFF")
_EXCEL_HEADER_FILL = PatternFill(start_color="24292F", end_color="24292F", fill_type="solid")
_EXCEL_BOLD = Font(bold=True)
BATCH_COLUMNS: tuple[tuple[str, str], ...] = (
    ("production_date", "Production Date"), ("machine", "Machine"), ("capacity_kg", "Capacity (Kg)"),
    ("tank", "Tank"), ("dyelot", "Dyelot"), ("sap_lot", "SapLot"), ("redye", "ReDye"),
    ("brand_program", "Brand Program"), ("fabric_type", "Fabric Type"),
    ("fabric_filled", "Fabric Type from next batch"), ("start_time", "Start"), ("end_time", "End"),
    ("hours", "Hours in period"), ("classification", "Classification"), ("counted", "Counted"),
)


def _classification_label(batch: dict[str, Any], require_redye_zero: bool) -> str:
    """Nhãn phân loại theo `classify_batch_badge()` (cùng quy tắc với cột Counted)."""
    badge = classify_batch_badge(
        {"dyelot": batch["dyelot"], "sap_lot": batch["sap_lot"], "redye": batch["redye"]},
        require_redye_zero=require_redye_zero,
    )
    if badge == "CM":
        return "CM"
    if badge == "S":
        return "Sample"
    return "Rework" if badge.endswith("R") else "Normal"


def _load_batch_details_by_start(start_times: set[str]) -> dict[str, list[dict[str, Any]]]:
    """{start_time: [dòng batch_details]} — truy vấn theo lô để không vượt giới hạn tham số."""
    result: dict[str, list[dict[str, Any]]] = {}
    values = sorted(start_times)
    for index in range(0, len(values), _DETAIL_QUERY_CHUNK):
        chunk = values[index:index + _DETAIL_QUERY_CHUNK]
        placeholders = ", ".join("?" for _ in chunk)
        rows = execute_query(
            "SELECT id, dyelot, sap_lot, redye, machine, fabric_type, start_time, end_time "
            f"FROM batch_details WHERE start_time IN ({placeholders}) ORDER BY id",
            chunk,
        )
        for row in rows:
            result.setdefault(str(row["start_time"]), []).append(dict(row))
    return result


def _match_details(machine: str, candidates: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Dòng raw cùng Machine; không có thì lấy dòng Machine trống (đã được fill-down lúc rollup)."""
    key = str(machine or "").strip().lower()
    exact = [batch for batch in candidates if str(batch["machine"] or "").strip().lower() == key]
    return exact or [batch for batch in candidates if not str(batch["machine"] or "").strip()]


def _own_segment(batch: dict[str, Any], production_date: str, is_normal: bool) -> tuple[float, bool]:
    """Giờ + cờ Counted của RIÊNG 1 mẻ trong `production_date` — chỉ dùng khi 2 mẻ trùng
    (machine, start_time) bị gộp chung 1 dòng summary."""
    start, end = parse_batch_datetime(batch["start_time"]), parse_batch_datetime(batch["end_time"])
    segments = split_production_days(start, end) if start and end else []
    hours = sum(value for day, value in segments if day.isoformat() == production_date)
    counted = is_normal and bool(segments) and segments[-1][0].isoformat() == production_date
    return hours, counted


def _collect_records(
    loaded: dict[str, Any], require_redye_zero: bool,
    fabric_type: str | None = None, period_key: str | None = None,
) -> list[dict[str, Any]]:
    items = [
        item for item in loaded["rows"]
        if (fabric_type is None or item["fabric_type"] == fabric_type)
        and (period_key is None or item["period_key"] == period_key)
    ]
    details = _load_batch_details_by_start({str(item["row"]["start_time"]) for item in items})
    records: list[dict[str, Any]] = []
    for item in items:
        row = item["row"]
        base = {
            "production_date": row["production_date"], "machine": row["machine"],
            "capacity_kg": row["capacity_kg"], "tank": item["tank"],
            "brand_program": row["brand_program"] or "", "fabric_type": item["fabric_type"],
            "start_time": str(row["start_time"]),
        }
        matched = _match_details(row["machine"], details.get(str(row["start_time"]), []))
        if not matched:
            # Summary cũ hơn raw (chưa rebuild) — vẫn hiện để tổng giờ/số mẻ khớp web.
            records.append({
                **base, "dyelot": None, "sap_lot": None, "redye": None, "fabric_filled": False,
                "end_time": None, "hours": float(row["hours"] or 0), "classification": "Not found in Batch",
                "counted": bool(row["batch_count"]),
            })
            continue
        for batch in matched:
            classification = _classification_label(batch, require_redye_zero)
            if len(matched) == 1:
                hours, counted = float(row["hours"] or 0), bool(row["batch_count"])
            else:
                hours, counted = _own_segment(batch, row["production_date"], classification == "Normal")
            records.append({
                **base, "dyelot": batch["dyelot"], "sap_lot": batch["sap_lot"], "redye": batch["redye"],
                "fabric_filled": str(batch["fabric_type"] or "").strip().lower() in INVALID_FABRIC_TYPES,
                "end_time": batch["end_time"], "hours": hours, "classification": classification,
                "counted": counted,
            })
    records.sort(key=lambda r: (r["production_date"], str(r["machine"]).lower(), r["start_time"]))
    return records


def _summarize_records(records: list[dict[str, Any]]) -> dict[str, Any]:
    """Tổng giờ, số mẻ Normal được đếm, Batch/Day tính lại — để đối chiếu với ô trên web."""
    hours = sum(record["hours"] for record in records)
    counted = sum(1 for record in records if record["counted"])
    return {
        "segments": len(records),
        "occupied_hours": round(hours, 2),
        "normal_batches": counted,
        "batch_per_day": round(counted * 24 / hours, 2) if hours else 0.0,
    }


def get_trend_batches(
    fabric_type: str,
    period_key: str | None = None,
    capacities: str | list[str] | None = None,
    brand_programs: str | list[str] | None = None,
    tank_types: str | list[str] | None = None,
    from_date: str | None = None, to_date: str | None = None,
    group_by: str = "date",
    require_redye_zero: bool = True,
) -> dict[str, Any]:
    """Danh sách đoạn (mẻ x ngày) của 1 ô Trend (`period_key`) hoặc ô Total (`period_key`
    = None), theo ĐÚNG bộ lọc đang chọn trên UI."""
    fabric = _normalize_main_fabric_type(fabric_type)
    if fabric is None:
        raise ValueError(f"Fabric type không hợp lệ: {fabric_type}")
    group_by = group_by if group_by in {"date", "week", "month"} else "date"
    loaded = _load_filtered_rows(capacities, brand_programs, tank_types, from_date, to_date, group_by, require_redye_zero)
    records = _collect_records(loaded, require_redye_zero, fabric, period_key) if loaded else []
    return {
        "fabric_type": fabric,
        "period_key": period_key,
        "batches": [{**record, "hours": round(record["hours"], 2)} for record in records],
        "summary": _summarize_records(records),
    }


def _write_header(sheet: Any, row_idx: int, headers: list[str]) -> None:
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row_idx, column=col, value=header)
        cell.font = _EXCEL_HEADER_FONT
        cell.fill = _EXCEL_HEADER_FILL


def _write_batches_sheet(sheet: Any, records: list[dict[str, Any]]) -> None:
    """Sheet "Batches": 1 dòng/1 đoạn + khối "Check" (tổng giờ, số mẻ Normal, Batch/Day tính
    lại theo từng loại vải) ở cuối."""
    _write_header(sheet, 1, [label for _, label in BATCH_COLUMNS])
    for col in range(1, len(BATCH_COLUMNS) + 1):
        sheet.column_dimensions[sheet.cell(row=1, column=col).column_letter].width = 16
    for row_idx, record in enumerate(records, start=2):
        for col, (key, _) in enumerate(BATCH_COLUMNS, start=1):
            value = record[key]
            if key == "hours":
                value = round(value, 2)
            elif key == "counted":
                value = "Yes" if value else "No"
            elif key == "fabric_filled":
                value = "Yes" if value else None
            sheet.cell(row=row_idx, column=col, value=value)
    last_row = len(records) + 1
    sheet.auto_filter.ref = f"A1:{sheet.cell(row=1, column=len(BATCH_COLUMNS)).column_letter}{last_row}"
    sheet.freeze_panes = "A2"

    check_row = last_row + 2
    sheet.cell(row=check_row, column=1, value="Check (compare with web)").font = _EXCEL_BOLD
    _write_header(sheet, check_row + 1, ["Fabric Type", "Occupied hours", "Normal batches", "Batch/Day"])
    groups = [(name, [r for r in records if r["fabric_type"] == name]) for name in MAIN_FABRIC_TYPES]
    groups = [(name, group) for name, group in groups if group]
    if len(groups) != 1:
        groups.append(("All", records))
    for row_idx, (name, group) in enumerate(groups, start=check_row + 2):
        summary = _summarize_records(group)
        values = (name, summary["occupied_hours"], summary["normal_batches"], summary["batch_per_day"])
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_idx, column=col, value=value)
            if name == "All":
                cell.font = _EXCEL_BOLD


def export_batch_day_trend_excel(
    capacities: str | list[str] | None = None,
    brand_programs: str | list[str] | None = None,
    tank_types: str | list[str] | None = None,
    from_date: str | None = None, to_date: str | None = None,
    group_by: str = "date",
    require_redye_zero: bool = True,
    fabric_type: str | None = None,
    period_key: str | None = None,
) -> bytes:
    """`fabric_type` rỗng -> cả báo cáo: sheet "Trend" (Batch/Day + số mẻ Normal + giờ theo kỳ)
    + "Batches" (mọi đoạn trong khoảng lọc). Có `fabric_type` -> chỉ danh sách mẻ của ô đang
    mở trên web (`period_key` rỗng = ô Total). Luôn kèm sheet "Filters" ghi bộ lọc đã dùng."""
    group_by = group_by if group_by in {"date", "week", "month"} else "date"
    fabric = None
    if fabric_type:
        fabric = _normalize_main_fabric_type(fabric_type)
        if fabric is None:
            raise ValueError(f"Fabric type không hợp lệ: {fabric_type}")
    loaded = _load_filtered_rows(capacities, brand_programs, tank_types, from_date, to_date, group_by, require_redye_zero)
    records = _collect_records(loaded, require_redye_zero, fabric, period_key) if loaded else []

    workbook = Workbook()
    if fabric is None:
        trend = get_batch_day_trend(capacities, brand_programs, tank_types, from_date, to_date, group_by, require_redye_zero)
        sheet = workbook.active
        sheet.title = "Trend"
        periods = trend["periods"]
        row_idx = 1
        for title, value_key, total_key in (
            ("Batch/Day", "values", "total"),
            ("Normal batches (counted)", "counts", "total_count"),
            ("Occupied hours", "hours", "total_hours"),
        ):
            sheet.cell(row=row_idx, column=1, value=title).font = _EXCEL_BOLD
            _write_header(sheet, row_idx + 1, ["Fabric Type", "Target", *periods, "Total"])
            for offset, row in enumerate(trend["rows"], start=row_idx + 2):
                sheet.cell(row=offset, column=1, value=row["fabric_type"])
                sheet.cell(row=offset, column=2, value=row["target"] if value_key == "values" else None)
                for col, value in enumerate(row[value_key], start=3):
                    sheet.cell(row=offset, column=col, value=value)
                sheet.cell(row=offset, column=3 + len(periods), value=row[total_key]).font = _EXCEL_BOLD
            row_idx += len(trend["rows"]) + 3
        for col in range(1, len(periods) + 4):
            sheet.column_dimensions[sheet.cell(row=1, column=col).column_letter].width = 13
        batches_sheet = workbook.create_sheet("Batches")
    else:
        batches_sheet = workbook.active
        batches_sheet.title = "Batches"
    _write_batches_sheet(batches_sheet, records)

    period_label = "Total" if period_key is None else next(
        (item["period_label"] for item in (loaded or {}).get("rows", []) if item["period_key"] == period_key),
        period_key,
    )
    filters_sheet = workbook.create_sheet("Filters")
    _write_header(filters_sheet, 1, ["Filter", "Value"])
    filters_sheet.column_dimensions["A"].width = 20
    filters_sheet.column_dimensions["B"].width = 40
    filters = [
        ("From Date", from_date or "All"), ("To Date", to_date or "All"),
        ("Group By", {"date": "Day", "week": "Week", "month": "Month"}[group_by]),
        ("Capacity (Kg)", ", ".join(str(value) for value in _parse_capacities(capacities)) or "All"),
        ("Brand Program", ", ".join(_parse_brand_programs(brand_programs)) or "All"),
        ("Tank Type", ", ".join(_parse_brand_programs(tank_types)) or "All"),
        ("ReDye = 0", "Yes" if require_redye_zero else "No"),
        ("Fabric Type", fabric or "All"),
        ("Period", period_label if fabric else "All"),
    ]
    for row_idx, (label, value) in enumerate(filters, start=2):
        filters_sheet.cell(row=row_idx, column=1, value=label)
        filters_sheet.cell(row=row_idx, column=2, value=value)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()

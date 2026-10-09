"""
modules/knitting/engines/excel_import/program_importer.py
--------------------------------------------------------
Parse 2 nguồn dùng cho bộ lọc Program của báo cáo Downtime Dệt (người dùng chốt 2026-10-09:
"cột Greige ID sẽ quyết định program"):

1. "CET-Piece Produced report - <từ> - <đến>.csv" — 1 dòng = 1 CUỘN vải (Roll No duy nhất), có
   M/c Code, Job ID, Greige ID, Available/Running/Stopped (phút) và Record Start/End (CHỈ có ngày
   dd/mm/yyyy, không có giờ). Header 2 dòng (dòng nhóm + dòng tên cột, tìm dòng có "Roll No"); có 2
   cột trùng tên "Time" (Aut. Stops / Decl.Stop) — không dùng. File thường phủ 1 tháng, import
   UPSERT theo Roll No (cuộn vắt qua 2 tháng xuất hiện ở cả 2 file).
   `production_date` của cuộn = Record End KẸP vào khoảng sản xuất trong TÊN FILE ("01-07-2026
   07;00;00 - 01-08-2026 07;00;00" -> 01/07..31/07): cuộn có Record End = ngày kết thúc file chắc
   chắn xong TRƯỚC 07:00 nên thuộc ngày sản xuất hôm trước (bug thật 2026-10-09: 87 cuộn sáng 01/08
   bị tính sang tháng 8). Tổng theo THÁNG đúng tuyệt đối; theo NGÀY vẫn gần đúng (file không có giờ).
   Tên file không có khoảng -> production_date = Record End.
2. "Knitting program.xlsx" — sheet đầu: cột A "Greige code SAP", cột B "Program"; bảng Core program
   không tiêu đề ở cột F:G (F = "Core program", G = tên program). Cột H/I người dùng xác nhận KHÔNG
   liên quan. Import = THAY THẾ toàn bộ danh mục.
"""
from __future__ import annotations

import csv
import io
import re
from collections import Counter
from datetime import datetime, timedelta
from typing import Any

from openpyxl import load_workbook

from .importer import _decode, _header_key, _to_number, parse_period_from_filename

PIECE_FILE_TYPE = "KNITTING_PIECE_PRODUCED"
PROGRAM_FILE_TYPE = "KNITTING_PROGRAM"

PIECE_HEADER_MAP: dict[str, str] = {
    "machine group": "machine_group",
    "machine": "machine",
    "m/c code": "machine_code",
    "job id": "job_id",
    "sap lot": "sap_lot",
    "sale order": "sale_order",
    "roll no": "roll_no",
    "material type": "material_type",
    "kniting structure": "knitting_structure",
    "knitting structure": "knitting_structure",
    "available": "available",
    "running": "running",
    "stopped": "stopped",
    "total qty": "total_qty",
    "good qty": "good_qty",
    "greige id": "greige_id",
    "record start": "record_start",
    "record end": "record_end",
    # Báo cáo Incentive (2026-10-09): %Achieve = Σ(KNT N.W × Std.PTM) / Σ Available.
    "std.ptm": "std_ptm",
    "knt n.w(kg)": "knt_nw_kg",
    "final n.w(kg)": "final_nw_kg",
    "operator code": "operator_code",
}
PIECE_FIELDS: tuple[str, ...] = (
    "roll_no", "machine_group", "machine", "machine_code", "job_id", "sap_lot", "sale_order", "material_type",
    "knitting_structure", "greige_id", "available", "running", "stopped", "total_qty", "good_qty",
    "record_start", "record_end", "std_ptm", "knt_nw_kg", "final_nw_kg", "operator_code", "production_date",
)
_PIECE_REQUIRED = {"machine_code", "roll_no", "greige_id", "record_start", "record_end", "available"}
_PIECE_NUMERIC = {"available", "running", "stopped", "total_qty", "good_qty", "std_ptm", "knt_nw_kg", "final_nw_kg"}


def normalize_program(name: Any) -> str:
    """Khoá so khớp tên Program: bỏ khoảng trắng đặc biệt (NBSP trong "Airism\\xa0 (50S/1 165G)"),
    gộp khoảng trắng, viết thường — "GEL"/"Gel", "Cooling tee"/"Cooling Tee" là 1 program."""
    return " ".join(str(name or "").replace("\xa0", " ").lower().split())


def _clean_name(name: Any) -> str:
    return " ".join(str(name or "").replace("\xa0", " ").split())


def _parse_day(value: str) -> str:
    return datetime.strptime(value.strip(), "%d/%m/%Y").date().isoformat()


def _production_window(filename: str | None) -> tuple[str, str] | None:
    """(ngày sản xuất đầu, ngày sản xuất cuối) từ tên file, None nếu tên file không có khoảng."""
    try:
        start, end = parse_period_from_filename(filename or "")
    except ValueError:
        return None
    last = (end - timedelta(seconds=1))
    first_day = start.date() if start.hour >= 7 else start.date() - timedelta(days=1)
    last_day = last.date() if last.hour >= 7 else last.date() - timedelta(days=1)
    return first_day.isoformat(), last_day.isoformat()


def parse_piece_produced_file(file_bytes: bytes, filename: str | None = None) -> dict[str, Any]:
    window = _production_window(filename)
    rows = list(csv.reader(io.StringIO(_decode(file_bytes))))
    header_index = next((i for i, row in enumerate(rows) if any(_header_key(c) == "roll no" for c in row)), None)
    if header_index is None:
        raise ValueError("Không tìm thấy dòng header có cột 'Roll No' — không phải file Piece Produced report.")
    columns: dict[str, int] = {}
    for idx, cell in enumerate(rows[header_index]):
        field = PIECE_HEADER_MAP.get(_header_key(cell))
        if field and field not in columns:  # cột trùng tên: giữ cột đầu tiên
            columns[field] = idx
    missing = _PIECE_REQUIRED - columns.keys()
    if missing:
        raise ValueError(f"File Piece Produced thiếu cột bắt buộc: {', '.join(sorted(missing))}.")

    rolls: dict[str, dict[str, Any]] = {}
    errors: list[dict[str, Any]] = []
    for offset, raw in enumerate(rows[header_index + 1:], start=header_index + 2):
        if not any(cell.strip() for cell in raw):
            continue
        data = {field: (raw[idx].strip() if idx < len(raw) else "") for field, idx in columns.items()}
        try:
            for field in ("machine_code", "roll_no", "greige_id"):
                if not data[field]:
                    raise ValueError(f"thiếu {field}")
            record: dict[str, Any] = {field: data.get(field) or None for field in PIECE_FIELDS}
            for field in _PIECE_NUMERIC & data.keys():
                record[field] = _to_number(data[field], field)
            try:
                record["record_start"] = _parse_day(data["record_start"])
                record["record_end"] = _parse_day(data["record_end"])
            except ValueError as exc:
                raise ValueError("Record Start/End phải dạng dd/mm/yyyy") from exc
            if record["record_end"] < record["record_start"]:
                raise ValueError("Record End trước Record Start")
            production_date = record["record_end"]
            if window:
                production_date = min(max(production_date, window[0]), window[1])
            record["production_date"] = production_date
            rolls[record["roll_no"]] = record  # Roll No trùng trong file: giữ dòng cuối
        except ValueError as exc:
            errors.append({"row": offset, "error": str(exc)})

    values = list(rolls.values())
    return {
        "rolls": values,
        "errors": errors,
        "total_rows": len(values) + len(errors),
        "date_from": min((r["production_date"] for r in values), default=None),
        "date_to": max((r["production_date"] for r in values), default=None),
        "production_window": window,
        "machines": len({r["machine_code"] for r in values}),
        "greige_ids": sorted({r["greige_id"] for r in values}),
    }


def parse_program_file(file_bytes: bytes) -> dict[str, Any]:
    """Greige -> Program (Greige trùng: dòng CUỐI thắng) + danh sách Core program.
    Tên hiển thị của 1 program = cách viết gặp nhiều nhất trong nhóm cùng khoá chuẩn hoá."""
    try:
        sheet = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True).worksheets[0]
    except Exception as exc:  # noqa: BLE001 — openpyxl ném nhiều loại lỗi khác nhau cho file hỏng
        raise ValueError(f"Không đọc được file Excel: {exc}") from exc
    rows = [tuple(row) + (None,) * 7 for row in sheet.iter_rows(values_only=True)]
    if not rows or _header_key(rows[0][0]) != "greige code sap" or _header_key(rows[0][1]) != "program":
        raise ValueError("Không phải file Knitting program — ô A1/B1 phải là 'Greige code SAP' / 'Program'.")

    spellings: dict[str, Counter[str]] = {}
    greige_to_key: dict[str, str] = {}
    conflicts: dict[str, set[str]] = {}
    core_keys: list[str] = []
    for row in rows:
        if _header_key(row[5]) == "core program" and _clean_name(row[6]):
            key = normalize_program(row[6])
            spellings.setdefault(key, Counter())[_clean_name(row[6])] += 0  # tên core cũng là 1 cách viết hợp lệ
            if key not in core_keys:
                core_keys.append(key)
    for row in rows[1:]:
        greige, program = str(row[0] or "").strip(), _clean_name(row[1])
        if not greige or not program:
            continue
        key = normalize_program(program)
        spellings.setdefault(key, Counter())[program] += 1
        if greige in greige_to_key and greige_to_key[greige] != key:
            conflicts.setdefault(greige, {greige_to_key[greige]}).add(key)
        greige_to_key[greige] = key

    def display(key: str) -> str:
        counter = spellings[key]
        return max(counter, key=lambda name: (counter[name], -list(counter).index(name)))

    return {
        "greige_programs": [{"greige_code": g, "program_key": k, "program": display(k)} for g, k in greige_to_key.items()],
        "core_programs": [{"program_key": k, "program": display(k)} for k in core_keys],
        "conflicts": [{"greige_code": g, "programs": sorted(display(k) for k in keys), "used": display(greige_to_key[g])}
                      for g, keys in sorted(conflicts.items())],
        "programs": len({k for k in greige_to_key.values()}),
    }


def detect_file_type(filename: str, file_bytes: bytes) -> str:
    """Nhận loại file knitting theo đuôi + header (không dựa vào tên file)."""
    from .importer import FILE_TYPE as STOP_FILE_TYPE

    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        return PROGRAM_FILE_TYPE
    if name.endswith(".csv"):
        head = file_bytes[:20000].decode("utf-8", errors="ignore").lower()
        if re.search(r"\broll no\b", head) and "greige id" in head:
            return PIECE_FILE_TYPE
        if "stop code" in head and "m/c code" in head:
            return STOP_FILE_TYPE
        raise ValueError("Không nhận ra file CSV — cần Stop Reason Analysis by Machine hoặc Piece Produced report.")
    raise ValueError("Chỉ nhận .csv (Stop Reason / Piece Produced) hoặc .xlsx (Knitting program).")

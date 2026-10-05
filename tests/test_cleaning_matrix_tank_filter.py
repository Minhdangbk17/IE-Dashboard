"""
tests/test_cleaning_matrix_tank_filter.py
-------------------------------------------
2026-10-05 — filter Tank Type (J tank / O tank) cho Batch Per Day by Machine
(`get_cleaning_matrix(tank_types=...)` + Excel export). Kiểm tra:

  1. `available_tank_types` theo thứ tự J tank -> O tank -> Unclassified; Tank trong Machine
     Master được chuẩn hoá không phân biệt hoa/thường + khoảng trắng ("o TANK " -> "O tank").
  2. Lọc J tank -> chỉ còn máy J tank; KPI + bảng Colour chỉ đếm mẻ của máy đó.
  3. Máy Tank trống + máy unmapped (không có trong Machine Master) -> "Unclassified": bị loại
     khi lọc J/O tank, hiện khi chọn Unclassified.
  4. Excel export theo đúng filter Tank.

Chạy: python tests/test_cleaning_matrix_tank_filter.py
"""
from __future__ import annotations

import io
import os
import shutil
import sqlite3
import sys
import tempfile
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from flask import Flask  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from core.database import close_db  # noqa: E402
from modules.dyeing.engines.reports import cleaning_matrix  # noqa: E402

# (machine_code, tank_type)
MACHINES = [("M1", "J tank"), ("M2", "o TANK "), ("M3", None)]
# (dyelot, machine, fabric_type, sap_lot, shade, start, end) — mọi mẻ đều Normal.
BATCHES = [
    ("C100000", "M1", "Cotton", "1000000001", "Dark", "2026-09-01 08:00:00", "2026-09-01 10:00:00"),
    ("C100010", "M1", "Cotton", "1000000002", "Light", "2026-09-01 11:00:00", "2026-09-01 13:00:00"),
    ("C100020", "M2", "CVC", "1000000003", "Dark", "2026-09-01 08:00:00", "2026-09-01 10:00:00"),
    ("C100030", "M3", "Polyester", "1000000004", "Dark", "2026-09-01 08:00:00", "2026-09-01 10:00:00"),
    ("C100040", "M9", "Polyester", "1000000005", "Dark", "2026-09-01 08:00:00", "2026-09-01 10:00:00"),
]
DAY = "2026-09-01"


def _check(label: str, actual, expected, failures: list[str]) -> None:
    ok = actual == expected
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}: actual={actual!r} expected={expected!r}")
    if not ok:
        failures.append(label)


def _init_db(db_path: str) -> None:
    conn = sqlite3.connect(db_path)
    conn.execute("""
        CREATE TABLE batch_details (
            id INTEGER PRIMARY KEY AUTOINCREMENT, dyelot TEXT NOT NULL, sap_lot TEXT, redye REAL DEFAULT 0,
            shade TEXT, colour_no TEXT, recipe_no TEXT, customer_color TEXT, batch_type TEXT,
            is_rework INTEGER DEFAULT 0, run_time REAL DEFAULT 0, greige_code TEXT,
            machine TEXT, fabric_type TEXT, start_time TEXT, end_time TEXT
        )
    """)
    conn.execute("CREATE TABLE machines (machine_id TEXT, machine_code TEXT, mc_brand TEXT, tank_type TEXT, mc_quantity INTEGER, tube_no INTEGER, capacity_kg REAL, domain TEXT)")
    conn.executemany(
        "INSERT INTO machines (machine_id, machine_code, tank_type, capacity_kg, domain) VALUES (?, ?, ?, 600, 'dyeing')",
        [(code, code, tank) for code, tank in MACHINES],
    )
    conn.executemany(
        "INSERT INTO batch_details (dyelot, machine, fabric_type, sap_lot, shade, start_time, end_time) VALUES (?, ?, ?, ?, ?, ?, ?)",
        BATCHES,
    )
    conn.commit()
    conn.close()


def _machines(data: dict) -> list[str]:
    return [item["machine"] for item in data["matrix"]]


def _colour_total(data: dict) -> int:
    return sum(sum(counts) for counts in data["color_summary"]["by_day"].values())


def _scenarios(failures: list[str]) -> None:
    print("\n=== Kịch bản 1: không lọc ===")
    data = cleaning_matrix.get_cleaning_matrix(DAY, DAY)
    _check("available_tank_types", data["available_tank_types"], ["J tank", "O tank", "Unclassified"], failures)
    _check("đủ máy (gồm M9 unmapped)", _machines(data), ["M1", "M2", "M3", "M9"], failures)
    _check("KPI Normal", data["kpis"]["normal_batches"], 5, failures)

    print("\n=== Kịch bản 2: lọc J tank ===")
    data = cleaning_matrix.get_cleaning_matrix(DAY, DAY, tank_types=["J tank"])
    _check("chỉ còn M1", _machines(data), ["M1"], failures)
    _check("KPI Normal", data["kpis"]["normal_batches"], 2, failures)
    _check("bảng Colour = KPI", _colour_total(data), 2, failures)
    _check("không cảnh báo unmapped", data["unmapped_machines"], [], failures)
    _check("available_tank_types không đổi khi đang lọc", data["available_tank_types"], ["J tank", "O tank", "Unclassified"], failures)

    print("\n=== Kịch bản 3: lọc O tank (Tank 'o TANK ' được chuẩn hoá) ===")
    data = cleaning_matrix.get_cleaning_matrix(DAY, DAY, tank_types=["O tank"])
    _check("chỉ còn M2", _machines(data), ["M2"], failures)
    _check("KPI Normal", data["kpis"]["normal_batches"], 1, failures)

    print("\n=== Kịch bản 4: J tank + O tank ===")
    data = cleaning_matrix.get_cleaning_matrix(DAY, DAY, tank_types=["J tank", "O tank"])
    _check("M1 + M2", _machines(data), ["M1", "M2"], failures)
    _check("KPI Normal", data["kpis"]["normal_batches"], 3, failures)

    print("\n=== Kịch bản 5: Unclassified (Tank trống + máy unmapped) ===")
    data = cleaning_matrix.get_cleaning_matrix(DAY, DAY, tank_types=["Unclassified"])
    _check("M3 + M9", _machines(data), ["M3", "M9"], failures)
    _check("KPI Normal", data["kpis"]["normal_batches"], 2, failures)

    print("\n=== Kịch bản 6: Excel export lọc J tank ===")
    sheets = load_workbook(io.BytesIO(cleaning_matrix.export_cleaning_matrix_excel(DAY, DAY, tank_types=["J tank"])))
    detail_machines = [row[0] for row in sheets["Detail"].iter_rows(min_row=2, values_only=True)]
    _check("sheet Detail chỉ có M1", detail_machines, ["M1"], failures)
    colour = sheets["Color Summary"]
    _check("Color Summary dòng Total = 2", colour.cell(colour.max_row, colour.max_column).value, 2, failures)


def main() -> int:
    failures: list[str] = []
    tmp_dir = tempfile.mkdtemp()
    try:
        db_path = os.path.join(tmp_dir, "test.db")
        _init_db(db_path)
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        cleaning_matrix.recompute_all([date(2026, 9, 1)], conn)
        conn.commit()
        conn.close()

        app = Flask(__name__)
        app.config["DATABASE_PATH"] = db_path
        app.config["SQLITE_PRAGMAS"] = {}
        with app.app_context():
            _scenarios(failures)
            close_db()
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    print(f"\n{'=' * 60}\nKẾT QUẢ: {'TẤT CẢ KHỚP' if not failures else f'{len(failures)} CASE LỆCH'}\n{'=' * 60}")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())

"""
modules/knitting/engines/incentive/service.py
------------------------------------------------
Báo cáo Incentive xưởng Dệt (người dùng chốt 2026-10-09):

    %Achieve = Σ(KNT N.W(kg) × Std.PTM) / Σ Available        (nguồn: Piece Produced report)
    Incentive = Đơn giá (VND/kg) của bậc chứa %Achieve × Σ KNT N.W(kg)

- Tính cho CẢ XƯỞNG và từng NHÓM (cột Machine Group của file, VD "CETVN Block D-1F"); mỗi dòng có
  %Achieve RIÊNG -> bậc riêng (dòng Xưởng không phải tổng tiền các nhóm).
- Kỳ: THÁNG, cộng dồn từ ngày 1 tới ngày "As of" (mặc định ngày Record End mới nhất trong tháng).
- Cuộn thuộc ngày `production_date` = Record End kẹp vào khoảng sản xuất trong tên file Piece Produced
  (xem `excel_import/program_importer.py`) -> tổng THÁNG đúng; theo ngày gần đúng (file không có giờ).
- Bậc đơn giá "From (>) To (<=)" lưu ở `knitting_incentive_bands` (seed bảng người dùng, admin sửa).
  %Achieve <= From của bậc đầu -> bậc đầu; > To của bậc cuối -> bậc cuối (bảng gốc dừng ở 100%,
  vượt 100% vẫn hưởng mức cao nhất — giả định, chờ người dùng xác nhận).
- Cuộn chưa có KNT N.W (đang dệt / chưa cân, = 0) vẫn cộng Available theo đúng công thức; cuộn import
  TRƯỚC khi có cột KNT N.W / Std.PTM (NULL) được đếm riêng để cảnh báo import lại.
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date, datetime, timedelta
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from core.database import get_dialect
from modules.knitting.engines.excel_import.service import ensure_tables as ensure_data_tables

WORKSHOP = "Workshop"

# (From >, To <=, VND/kg) — bảng người dùng gửi 2026-10-09.
DEFAULT_BANDS: tuple[tuple[float, float, float], ...] = (
    (0.0, 82.0, 0), (82.0, 84.5, 15), (84.5, 87.0, 45), (87.0, 89.5, 60), (89.5, 92.0, 80),
    (92.0, 94.0, 104), (94.0, 96.0, 130), (96.0, 98.0, 162), (98.0, 100.0, 192),
)


# ---------------------------------------------------------------------------
# Bậc đơn giá
# ---------------------------------------------------------------------------


def ensure_tables(conn: Any) -> None:
    ensure_data_tables(conn)
    if get_dialect() == "sqlite":
        conn.execute("""
            CREATE TABLE IF NOT EXISTS knitting_incentive_bands (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                from_pct REAL NOT NULL,
                to_pct REAL NOT NULL,
                unit_vnd_per_kg REAL NOT NULL,
                updated_by TEXT,
                updated_at TEXT NOT NULL DEFAULT (datetime('now'))
            )
        """)
    if conn.execute("SELECT COUNT(*) AS n FROM knitting_incentive_bands").fetchone()["n"] == 0:
        conn.executemany(
            "INSERT INTO knitting_incentive_bands (from_pct, to_pct, unit_vnd_per_kg, updated_by) VALUES (?, ?, ?, 'seed')",
            DEFAULT_BANDS,
        )
    conn.commit()


def load_bands(conn: Any) -> list[dict[str, float]]:
    ensure_tables(conn)
    rows = conn.execute("SELECT from_pct, to_pct, unit_vnd_per_kg FROM knitting_incentive_bands ORDER BY from_pct").fetchall()
    return [{"from_pct": r["from_pct"], "to_pct": r["to_pct"], "unit": r["unit_vnd_per_kg"]} for r in rows]


def save_bands(conn: Any, bands: list[dict[str, Any]], user: str | None) -> None:
    """Thay toàn bộ bảng bậc. Bậc phải liên tục (To bậc trước = From bậc sau), From < To, đơn giá >= 0."""
    try:
        parsed = sorted(((float(b["from_pct"]), float(b["to_pct"]), float(b["unit"])) for b in bands), key=lambda b: b[0])
    except (KeyError, TypeError, ValueError) as exc:
        raise ValueError("Mỗi bậc cần From, To, Unit là số.") from exc
    if not parsed:
        raise ValueError("Cần ít nhất 1 bậc.")
    for index, (low, high, unit) in enumerate(parsed):
        if low >= high:
            raise ValueError(f"Bậc {index + 1}: From phải nhỏ hơn To.")
        if unit < 0:
            raise ValueError(f"Bậc {index + 1}: đơn giá không được âm.")
        if index and abs(parsed[index - 1][1] - low) > 1e-9:
            raise ValueError(f"Bậc {index + 1}: From ({low}) phải bằng To của bậc trước ({parsed[index - 1][1]}).")
    ensure_tables(conn)
    conn.execute("BEGIN")
    conn.execute("DELETE FROM knitting_incentive_bands")
    conn.executemany(
        "INSERT INTO knitting_incentive_bands (from_pct, to_pct, unit_vnd_per_kg, updated_by) VALUES (?, ?, ?, ?)",
        [(low, high, unit, user or "unknown") for low, high, unit in parsed],
    )
    conn.commit()


def band_for(achievement_pct: float | None, bands: list[dict[str, float]]) -> dict[str, Any] | None:
    """Bậc chứa %Achieve theo quy tắc From (>) To (<=), kèm khoảng cách tới bậc kế tiếp."""
    if achievement_pct is None or not bands:
        return None
    index = len(bands) - 1  # > To bậc cuối -> bậc cuối
    if achievement_pct <= bands[0]["from_pct"] + 1e-9:
        index = 0
    else:
        for i, band in enumerate(bands):
            if band["from_pct"] + 1e-9 < achievement_pct <= band["to_pct"] + 1e-9:
                index = i
                break
    nxt = bands[index + 1] if index + 1 < len(bands) else None
    return {
        "index": index, **bands[index],
        "next_from_pct": nxt["from_pct"] if nxt else None,
        "next_unit": nxt["unit"] if nxt else None,
        "gap_to_next_pct": max(nxt["from_pct"] - achievement_pct, 0.0) if nxt else None,
    }


# ---------------------------------------------------------------------------
# Kỳ + bộ lọc
# ---------------------------------------------------------------------------


def _split(value: str | Iterable[str] | None) -> list[str]:
    if value is None:
        return []
    items = value.split("|") if isinstance(value, str) else list(value)
    return sorted({str(item).strip() for item in items if str(item).strip()})


def list_options(conn: Any) -> dict[str, Any]:
    ensure_tables(conn)
    months = [r["m"] for r in conn.execute("SELECT DISTINCT substr(COALESCE(production_date, record_end), 1, 7) AS m FROM knitting_piece_rolls ORDER BY m DESC")]
    groups = [r["g"] for r in conn.execute(
        "SELECT DISTINCT machine_group AS g FROM knitting_piece_rolls WHERE machine_group IS NOT NULL AND machine_group <> '' ORDER BY g"
    )]
    return {"months": months, "groups": groups, "latest_date": conn.execute("SELECT MAX(COALESCE(production_date, record_end)) AS d FROM knitting_piece_rolls").fetchone()["d"]}


def normalize_filters(conn: Any, month: str | None, as_of: str | None, groups: str | Iterable[str] | None = None) -> dict[str, Any]:
    """Tháng "YYYY-MM" (mặc định tháng có dữ liệu mới nhất) + As of trong tháng (mặc định ngày
    Record End mới nhất của tháng, không có dữ liệu thì ngày cuối tháng)."""
    ensure_tables(conn)
    if not month:
        latest = conn.execute("SELECT MAX(COALESCE(production_date, record_end)) AS d FROM knitting_piece_rolls").fetchone()["d"]
        month = (latest or date.today().isoformat())[:7]
    try:
        start = date.fromisoformat(f"{month}-01")
    except ValueError as exc:
        raise ValueError(f"Tháng không hợp lệ: {month!r} (cần YYYY-MM).") from exc
    month_end = (start.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
    if as_of:
        try:
            as_of_day = date.fromisoformat(as_of)
        except ValueError as exc:
            raise ValueError(f"As of không hợp lệ: {as_of!r}") from exc
        if not start <= as_of_day <= month_end:
            raise ValueError("As of phải nằm trong tháng đã chọn.")
    else:
        latest = conn.execute(
            "SELECT MAX(COALESCE(production_date, record_end)) AS d FROM knitting_piece_rolls WHERE COALESCE(production_date, record_end) BETWEEN ? AND ?",
            (start.isoformat(), month_end.isoformat()),
        ).fetchone()["d"]
        as_of_day = date.fromisoformat(latest) if latest else month_end
    return {"month": month, "from_date": start.isoformat(), "to_date": as_of_day.isoformat(), "groups": _split(groups)}


def _where(filters: dict[str, Any], from_date: str, to_date: str) -> tuple[str, list[Any]]:
    clauses = ["COALESCE(production_date, record_end) BETWEEN ? AND ?"]
    params: list[Any] = [from_date, to_date]
    if filters["groups"]:
        clauses.append(f"machine_group IN ({','.join('?' for _ in filters['groups'])})")
        params += filters["groups"]
    return " AND ".join(clauses), params


# ---------------------------------------------------------------------------
# Báo cáo
# ---------------------------------------------------------------------------


def _new_bucket() -> dict[str, float]:
    return {"kg": 0.0, "std_minutes": 0.0, "available": 0.0, "rolls": 0, "missing": 0}


def _add(bucket: dict[str, float], row: Any) -> None:
    bucket["kg"] += row["kg"] or 0
    bucket["std_minutes"] += row["std_minutes"] or 0
    bucket["available"] += row["available"] or 0
    bucket["rolls"] += row["rolls"]
    bucket["missing"] += row["missing"]


def _pct(bucket: dict[str, float]) -> float | None:
    return bucket["std_minutes"] / bucket["available"] * 100 if bucket["available"] else None


def _with_incentive(name: str, bucket: dict[str, float], bands: list[dict[str, float]]) -> dict[str, Any]:
    pct = _pct(bucket)
    band = band_for(pct, bands)
    unit = band["unit"] if band else 0.0
    return {
        "name": name, "kg": bucket["kg"], "std_minutes": bucket["std_minutes"], "available": bucket["available"],
        "rolls": bucket["rolls"], "achievement_pct": pct, "band": band, "unit": unit, "incentive": unit * bucket["kg"],
    }


def _load(conn: Any, filters: dict[str, Any], from_date: str, to_date: str) -> list[Any]:
    where, params = _where(filters, from_date, to_date)
    return conn.execute(
        f"""
        SELECT COALESCE(production_date, record_end) AS d, COALESCE(machine_group, '') AS grp, machine_code,
               SUM(knt_nw_kg * std_ptm) AS std_minutes, SUM(available) AS available, SUM(knt_nw_kg) AS kg,
               COUNT(*) AS rolls,
               SUM(CASE WHEN knt_nw_kg IS NULL OR std_ptm IS NULL THEN 1 ELSE 0 END) AS missing
        FROM knitting_piece_rolls WHERE {where}
        GROUP BY COALESCE(production_date, record_end), COALESCE(machine_group, ''), machine_code
        """,
        params,
    ).fetchall()


def build_report(conn: Any, filters: dict[str, Any]) -> dict[str, Any]:
    ensure_tables(conn)
    bands = load_bands(conn)
    rows = _load(conn, filters, filters["from_date"], filters["to_date"])

    workshop = _new_bucket()
    by_group: dict[str, dict[str, float]] = defaultdict(_new_bucket)
    by_machine: dict[tuple[str, str], dict[str, float]] = defaultdict(_new_bucket)
    by_day: dict[str, dict[str, float]] = defaultdict(_new_bucket)
    by_day_group: dict[tuple[str, str], dict[str, float]] = defaultdict(_new_bucket)
    for row in rows:
        _add(workshop, row)
        _add(by_group[row["grp"]], row)
        _add(by_machine[(row["machine_code"], row["grp"])], row)
        _add(by_day[row["d"]], row)
        _add(by_day_group[(row["d"], row["grp"])], row)

    groups = sorted(by_group)
    # Cộng dồn từ ngày 1 (MTD) cho từng ngày có dữ liệu — Xưởng + từng nhóm.
    daily = []
    mtd = _new_bucket()
    mtd_group = {g: _new_bucket() for g in groups}
    for day in sorted(by_day):
        for key in ("kg", "std_minutes", "available", "rolls", "missing"):
            mtd[key] += by_day[day][key]
            for g in groups:
                mtd_group[g][key] += by_day_group.get((day, g), _new_bucket())[key]
        mtd_row = _with_incentive(WORKSHOP, mtd, bands)
        daily.append({
            "date": day, "kg": by_day[day]["kg"], "std_minutes": by_day[day]["std_minutes"], "available": by_day[day]["available"],
            "rolls": by_day[day]["rolls"], "day_pct": _pct(by_day[day]),
            "mtd_kg": mtd["kg"], "mtd_pct": mtd_row["achievement_pct"], "mtd_unit": mtd_row["unit"], "mtd_incentive": mtd_row["incentive"],
            "mtd_pct_by_group": {(g or "(No group)"): _pct(mtd_group[g]) for g in groups},
        })

    machines = [
        {**_with_incentive(machine, bucket, bands), "group": grp}
        for (machine, grp), bucket in by_machine.items()
    ]
    machines.sort(key=lambda m: (m["achievement_pct"] is None, m["achievement_pct"] or 0))
    return {
        "filters": filters,
        "bands": bands,
        "threshold_pct": next((b["from_pct"] for b in bands if b["unit"] > 0), None),
        "workshop": _with_incentive(WORKSHOP, workshop, bands),
        "groups": [_with_incentive(g or "(No group)", by_group[g], bands) for g in groups],
        "group_names": [g or "(No group)" for g in groups],
        "daily": daily,
        "machines": machines,
        "missing_rolls": workshop["missing"],
    }


def get_day_details(conn: Any, filters: dict[str, Any], day: str) -> dict[str, Any]:
    """Chi tiết 1 ngày: từng máy (kg, phút chuẩn, Available, %Achieve) + từng cuộn."""
    if not filters["from_date"] <= day <= filters["to_date"]:
        raise ValueError("Ngày nằm ngoài kỳ đang xem.")
    ensure_tables(conn)
    bands = load_bands(conn)
    machines = [
        {**_with_incentive(row["machine_code"], {"kg": row["kg"] or 0, "std_minutes": row["std_minutes"] or 0, "available": row["available"] or 0, "rolls": row["rolls"], "missing": row["missing"]}, bands), "group": row["grp"]}
        for row in _load(conn, filters, day, day)
    ]
    machines.sort(key=lambda m: (m["achievement_pct"] is None, m["achievement_pct"] or 0))
    where, params = _where(filters, day, day)
    rolls = [dict(r) for r in conn.execute(
        f"""
        SELECT roll_no, machine_group, machine_code, job_id, greige_id, knitting_structure, operator_code,
               knt_nw_kg, std_ptm, knt_nw_kg * std_ptm AS std_minutes, available, record_start, record_end, production_date
        FROM knitting_piece_rolls WHERE {where} ORDER BY machine_code, roll_no
        """,
        params,
    )]
    total = _new_bucket()
    for m in machines:
        total["kg"] += m["kg"]
        total["std_minutes"] += m["std_minutes"]
        total["available"] += m["available"]
    return {"date": day, "achievement_pct": _pct(total), "kg": total["kg"], "machines": machines, "rolls": rolls}


# ---------------------------------------------------------------------------
# Export Excel
# ---------------------------------------------------------------------------

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="24292F", end_color="24292F", fill_type="solid")
_BOLD = Font(bold=True)


def _sheet(workbook: Workbook, title: str, headers: list[str], widths: list[int], first: bool = False) -> Any:
    sheet = workbook.active if first else workbook.create_sheet(title)
    sheet.title = title
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
    for index, width in enumerate(widths):
        sheet.column_dimensions[chr(65 + index)].width = width
    sheet.freeze_panes = "A2"
    return sheet


def _formats(sheet: Any, formats: dict[int, str]) -> None:
    for row in sheet.iter_rows(min_row=2):
        for col, fmt in formats.items():
            row[col - 1].number_format = fmt


def export_excel(conn: Any, filters: dict[str, Any]) -> bytes:
    report = build_report(conn, filters)
    workbook = Workbook()
    summary = _sheet(workbook, "Summary", ["Scope", "KNT N.W (kg)", "Std minutes", "Available", "%Achieve", "Unit (VND/kg)", "Incentive (VND)", "Rolls"],
                     [26, 14, 14, 14, 11, 14, 16, 8], first=True)
    for row in [report["workshop"], *report["groups"]]:
        summary.append([row["name"], row["kg"], row["std_minutes"], row["available"],
                        None if row["achievement_pct"] is None else row["achievement_pct"] / 100, row["unit"], row["incentive"], row["rolls"]])
    summary.cell(row=2, column=1).font = _BOLD
    _formats(summary, {2: "#,##0.00", 3: "#,##0.00", 4: "#,##0.00", 5: "0.00%", 6: "#,##0", 7: "#,##0"})

    daily = _sheet(workbook, "Daily", ["Date", "KNT N.W (kg)", "Std minutes", "Available", "Day %Achieve", "MTD %Achieve", "MTD Unit", "MTD Incentive (VND)", "Rolls"],
                   [12, 14, 14, 14, 13, 13, 10, 18, 8])
    for d in report["daily"]:
        daily.append([d["date"], d["kg"], d["std_minutes"], d["available"], None if d["day_pct"] is None else d["day_pct"] / 100,
                      None if d["mtd_pct"] is None else d["mtd_pct"] / 100, d["mtd_unit"], d["mtd_incentive"], d["rolls"]])
    _formats(daily, {2: "#,##0.00", 3: "#,##0.00", 4: "#,##0.00", 5: "0.00%", 6: "0.00%", 7: "#,##0", 8: "#,##0"})

    machines = _sheet(workbook, "Machines", ["M/c Code", "Machine Group", "KNT N.W (kg)", "Std minutes", "Available", "%Achieve", "Rolls"], [11, 20, 14, 14, 14, 11, 8])
    for m in report["machines"]:
        machines.append([m["name"], m["group"], m["kg"], m["std_minutes"], m["available"], None if m["achievement_pct"] is None else m["achievement_pct"] / 100, m["rolls"]])
    _formats(machines, {3: "#,##0.00", 4: "#,##0.00", 5: "#,##0.00", 6: "0.00%"})

    band_sheet = _sheet(workbook, "Bands", ["From (>)", "To (<=)", "Unit Incentive (VND/Kg)"], [10, 10, 22])
    for b in report["bands"]:
        band_sheet.append([b["from_pct"] / 100, b["to_pct"] / 100, b["unit"]])
    _formats(band_sheet, {1: "0.0%", 2: "0.0%", 3: "#,##0"})

    info = _sheet(workbook, "Filters", ["Filter", "Value"], [20, 100])
    for label, value in (
        ("Month", filters["month"]), ("From", filters["from_date"]), ("As of", filters["to_date"]),
        ("Machine Group", " | ".join(filters["groups"]) or "All"),
        ("%Achieve", "Σ(KNT N.W(kg) × Std.PTM) / Σ Available (Piece Produced report; roll date = Record End clamped to the file's production period)"),
        ("Incentive", "Unit of the band containing %Achieve (From > , To <=) × Σ KNT N.W(kg), month to date"),
        ("Exported at", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ):
        info.append([label, value])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()

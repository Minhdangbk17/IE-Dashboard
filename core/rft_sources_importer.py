"""Importer 2 nguồn tra cứu của báo cáo Right First Time (RFT) — Production Report Dye và
NC Report, nạp HẰNG NGÀY và TÍCH LUỸ (UPSERT, không ghi đè theo file).

Cả 2 file đều import nguyên file THÔ do hệ thống xuất (không lọc tay). Mọi dòng được lưu lại;
quy tắc lọc nghiệp vụ (chỉ công đoạn `DG*`; NC "khác màu"/Closed/không phải thẻ MA) áp ở READ
TIME trong `modules/dyeing/engines/rft/service.py` — đổi quy tắc sau này không phải import lại.

- Production Report Dye (`production_report_dye_*.xlsx`, sheet thô `DYE`): header 3 dòng
  (Trung / Việt / Anh trong cùng 1 ô) — so khớp theo DÒNG CUỐI (tiếng Anh). Khoá UPSERT
  `(batch_no, operation, op_start_time)` -> bảng `dye_production_ops`.
- NC Report (`nc_report_*.xlsx`, sheet thô đầu tiên): khoá UPSERT `nc_no` -> bảng
  `dye_nc_reports` (1 NC có thể đổi trạng thái giữa các ngày — bản import sau ghi đè bản trước).

File có thể chứa thêm sheet người dùng tự lọc (VD `DG`, `NC`) — luôn đọc sheet ĐẦU TIÊN có đủ
cột nhận diện (sheet thô đứng trước trong file hệ thống xuất).
Cùng pattern ghi `import_logs` với `core/batch_importer.py::sync_batch_details()`.
"""
from __future__ import annotations

import io
import math
from datetime import datetime, timedelta
from typing import Any

from openpyxl import load_workbook

from core.database import get_db, get_dialect, insert_returning_id
from core.excel_importer import record_import_rows

DYE_PRODUCTION_FILE_TYPE = "DYE_PRODUCTION"
NC_REPORT_FILE_TYPE = "NC_REPORT"

# Header (đã chuẩn hoá: dòng cuối của ô, viết thường, gộp khoảng trắng) -> field DB.
DYE_PRODUCTION_HEADER_MAP: dict[str, str] = {
    "report op end time": "report_time",
    "department": "department",
    "operation#": "operation",
    "plant": "plant",
    "machine#": "machine",
    "batch#": "batch_no",
    "batch status": "batch_status",
    "sap lot": "sap_lot",
    "so#": "so_no",
    "brand": "brand",
    "customer": "customer",
    "grey code": "greige_code",
    "fabric code": "fabric_code",
    "recipe": "recipe",
    "crystal color code": "color_code",
    "operation output qty(kg)": "output_qty",
    "shift": "shift",
    "op start time": "op_start_time",
    "op end time": "op_end_time",
    "batch type": "batch_type",
}
DYE_PRODUCTION_FIELDS: tuple[str, ...] = tuple(dict.fromkeys(DYE_PRODUCTION_HEADER_MAP.values()))
DYE_PRODUCTION_REQUIRED = ("batch_no", "operation")
DYE_PRODUCTION_KEY = ("batch_no", "operation", "op_start_time")

NC_REPORT_HEADER_MAP: dict[str, str] = {
    "nc#": "nc_no",
    "defect": "defect",
    "defect qty (kg)": "defect_qty",
    "operation route": "operation_route",
    "status": "status",
    "dept report": "dept_report",
    "mp#": "mp_no",
    "plant": "plant",
    "delivery date": "delivery_date",
    "customer": "customer",
    "sale#": "sale_no",
    "batch ref": "batch_ref",
    "batch qty (kg)": "batch_qty",
    "colorist": "colorist",
    "sap lot": "sap_lot",
    "color": "color",
    "recipe": "recipe",
    "fabric code": "fabric_code",
    "greige code": "greige_code",
    "new batch": "new_batch",
    "corrective": "corrective",
    "reason": "reason",
    "dept in charge": "dept_in_charge",
    "remark": "remark",
    "created nc (1)": "created_nc_at",
    "confirm nc (2)": "confirm_nc_at",
    "created nb (3)": "created_nb_at",
}
NC_REPORT_FIELDS: tuple[str, ...] = tuple(dict.fromkeys(NC_REPORT_HEADER_MAP.values()))
NC_REPORT_REQUIRED = ("nc_no", "batch_ref")

_NUMERIC_FIELDS = {"output_qty", "defect_qty", "batch_qty"}
_DATETIME_FIELDS = {"report_time", "op_start_time", "op_end_time", "created_nc_at", "confirm_nc_at", "created_nb_at"}
_EXCEL_SERIAL_EPOCH = datetime(1899, 12, 30)

# Cột nhận diện loại file (header đã chuẩn hoá) — dùng cho cả chọn sheet lẫn auto-detect.
_SIGNATURES: dict[str, set[str]] = {
    DYE_PRODUCTION_FILE_TYPE: {"batch#", "operation#"},
    NC_REPORT_FILE_TYPE: {"nc#", "batch ref"},
}


def _header_key(value: Any) -> str:
    """Dòng cuối không rỗng của ô header, viết thường, gộp khoảng trắng — header Production
    Report có dạng "工卡号\\nMã công lệnh\\nBatch#", chỉ dòng tiếng Anh là ổn định."""
    lines = [line.strip() for line in str(value or "").splitlines() if line.strip()]
    key = " ".join(lines[-1].lower().split()) if lines else ""
    # Ô "Batch Status" bị dính tiếng Việt cùng dòng ("Trạng thái mã công lệnh Batch Status").
    return "batch status" if key.endswith("batch status") else key


def _clean(value: Any) -> Any:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    text = str(value).strip()
    return None if not text or text.lower() in {"nan", "none", "nat"} else value


def _to_text(value: Any) -> str:
    value = _clean(value)
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _to_datetime_text(value: Any) -> str:
    """'YYYY-MM-DD HH:MM:SS' — nhận `datetime`, số serial Excel hoặc chuỗi ngày giờ."""
    value = _clean(value)
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, (int, float)) and 1.0 <= float(value) <= 100000.0:
        return (_EXCEL_SERIAL_EPOCH + timedelta(days=float(value))).strftime("%Y-%m-%d %H:%M:%S")
    text = str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%m/%d/%Y %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            continue
    return text


def _to_number(value: Any) -> float | None:
    value = _clean(value)
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _find_sheet_rows(file_bytes: bytes, file_type: str) -> tuple[list[str], list[tuple[Any, ...]]] | None:
    """(header đã chuẩn hoá, các dòng dữ liệu) của sheet ĐẦU TIÊN có đủ cột nhận diện."""
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header_row = next(rows, None)
            if header_row is None:
                continue
            headers = [_header_key(value) for value in header_row]
            if _SIGNATURES[file_type] <= set(headers):
                return headers, list(rows)
    finally:
        workbook.close()
    return None


def detect_rft_source_type(file_bytes: bytes) -> str | None:
    """`DYE_PRODUCTION` / `NC_REPORT` nếu file là 1 trong 2 nguồn RFT, ngược lại `None`.
    Không raise với file không phải Excel hợp lệ — để luồng auto-detect chung xử lý tiếp."""
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception:
        return None
    try:
        for sheet in workbook.worksheets:
            header_row = next(sheet.iter_rows(values_only=True), None)
            headers = {_header_key(value) for value in header_row or ()}
            for file_type, signature in _SIGNATURES.items():
                if signature <= headers:
                    return file_type
    finally:
        workbook.close()
    return None


def _parse_file(
    file_bytes: bytes, file_type: str, header_map: dict[str, str], fields: tuple[str, ...], required: tuple[str, ...],
) -> dict[str, Any]:
    found = _find_sheet_rows(file_bytes, file_type)
    if found is None:
        columns = ", ".join(sorted(_SIGNATURES[file_type]))
        raise ValueError(f"Không tìm thấy sheet có đủ cột nhận diện ({columns}).")
    headers, raw_rows = found
    indexes: dict[str, int] = {}
    for index, header in enumerate(headers):
        field = header_map.get(header)
        if field is not None and field not in indexes:
            indexes[field] = index

    parsed: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    row_details: list[dict[str, Any]] = []
    for row_number, values in enumerate(raw_rows, start=2):
        if not any(_clean(value) is not None for value in values):
            continue
        record: dict[str, Any] = {}
        for field in fields:
            index = indexes.get(field)
            value = values[index] if index is not None and index < len(values) else None
            if field in _NUMERIC_FIELDS:
                record[field] = _to_number(value)
            elif field in _DATETIME_FIELDS:
                record[field] = _to_datetime_text(value)
            else:
                record[field] = _to_text(value)
        raw = {field: ("" if record[field] is None else str(record[field])) for field in fields}
        missing = [field for field in required if not record[field]]
        if missing:
            error = f"Thiếu giá trị bắt buộc: {', '.join(missing)}."
            errors.append({"row": row_number, "error": error})
            row_details.append({"row_number": row_number, "status": "invalid", "error": error, "data": raw})
            continue
        parsed.append(record)
        row_details.append({"row_number": row_number, "status": "valid", "error": None, "data": raw})
    return {"rows": parsed, "errors": errors, "total_records": len(parsed) + len(errors), "row_details": row_details}


def parse_dye_production_file(file_bytes: bytes) -> dict[str, Any]:
    return _parse_file(file_bytes, DYE_PRODUCTION_FILE_TYPE, DYE_PRODUCTION_HEADER_MAP, DYE_PRODUCTION_FIELDS, DYE_PRODUCTION_REQUIRED)


def parse_nc_report_file(file_bytes: bytes) -> dict[str, Any]:
    return _parse_file(file_bytes, NC_REPORT_FILE_TYPE, NC_REPORT_HEADER_MAP, NC_REPORT_FIELDS, NC_REPORT_REQUIRED)


def preview_rft_source_file(file_bytes: bytes, file_type: str) -> dict[str, Any]:
    parsed = parse_dye_production_file(file_bytes) if file_type == DYE_PRODUCTION_FILE_TYPE else parse_nc_report_file(file_bytes)
    return {
        "status": "preview", "file_type": file_type,
        "columns": list(parsed["rows"][0]) if parsed["rows"] else [],
        "preview": parsed["rows"][:5], "valid_rows": len(parsed["rows"]),
        "total_rows": parsed["total_records"], "errors": parsed["errors"],
    }


# ---------------------------------------------------------------------------
# Schema (CHỈ SQLite tự tạo — Postgres tạo qua `supabase/schema.sql` /
# `supabase/migrate_rft_sources.sql`).
# ---------------------------------------------------------------------------


def _ensure_import_logs_table(conn: Any) -> None:
    """Duplicate có chủ đích của `core/batch_importer.py::_ensure_import_logs_table()` — mỗi
    importer tự đảm bảo bảng chung tồn tại, không import chéo hàm private giữa module `core/`."""
    if get_dialect() == "sqlite":
        conn.execute("""
            CREATE TABLE IF NOT EXISTS import_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT, file_name TEXT NOT NULL,
                import_type TEXT NOT NULL, imported_by TEXT NOT NULL,
                total_rows INTEGER NOT NULL DEFAULT 0, success_rows INTEGER NOT NULL DEFAULT 0,
                error_rows INTEGER NOT NULL DEFAULT 0, status TEXT NOT NULL DEFAULT 'completed',
                error_detail TEXT, imported_at TEXT NOT NULL DEFAULT (datetime('now')),
                imported_rows INTEGER NOT NULL DEFAULT 0, file_type TEXT,
                created_at TEXT NOT NULL DEFAULT (datetime('now'))
            )
        """)
        existing = {row["name"] for row in conn.execute("PRAGMA table_info(import_logs)")}
        for column, definition in (("imported_rows", "INTEGER NOT NULL DEFAULT 0"), ("file_type", "TEXT"), ("created_at", "TEXT")):
            if column not in existing:
                conn.execute(f"ALTER TABLE import_logs ADD COLUMN {column} {definition}")


def ensure_rft_source_tables(conn: Any) -> None:
    if get_dialect() != "sqlite":
        return
    conn.execute("""
        CREATE TABLE IF NOT EXISTS dye_production_ops (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            report_time TEXT, department TEXT, operation TEXT NOT NULL, plant TEXT, machine TEXT,
            batch_no TEXT NOT NULL, batch_status TEXT, sap_lot TEXT, so_no TEXT, brand TEXT,
            customer TEXT, greige_code TEXT, fabric_code TEXT, recipe TEXT, color_code TEXT,
            output_qty REAL, shift TEXT, op_start_time TEXT NOT NULL DEFAULT '', op_end_time TEXT,
            batch_type TEXT, import_log_id INTEGER,
            created_at TEXT NOT NULL DEFAULT (datetime('now')),
            UNIQUE (batch_no, operation, op_start_time)
        )
    """)
    conn.execute("CREATE INDEX IF NOT EXISTS idx_dye_production_ops_batch_norm ON dye_production_ops (lower(trim(batch_no)))")
    conn.execute("""
        CREATE TABLE IF NOT EXISTS dye_nc_reports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nc_no TEXT NOT NULL UNIQUE, defect TEXT, defect_qty REAL, operation_route TEXT,
            status TEXT, dept_report TEXT, mp_no TEXT, plant TEXT, delivery_date TEXT, customer TEXT,
            sale_no TEXT, batch_ref TEXT NOT NULL, batch_qty REAL, colorist TEXT, sap_lot TEXT,
            color TEXT, recipe TEXT, fabric_code TEXT, greige_code TEXT, new_batch TEXT,
            corrective TEXT, reason TEXT, dept_in_charge TEXT, remark TEXT,
            created_nc_at TEXT, confirm_nc_at TEXT, created_nb_at TEXT, import_log_id INTEGER,
            created_at TEXT NOT NULL DEFAULT (datetime('now'))
        )
    """)
    conn.execute("CREATE INDEX IF NOT EXISTS idx_dye_nc_reports_batch_ref_norm ON dye_nc_reports (lower(trim(batch_ref)))")


def _sync(
    file_bytes: bytes, imported_by: str | None, filename: str, file_type: str,
    table: str, fields: tuple[str, ...], key: tuple[str, ...],
) -> dict[str, Any]:
    result = parse_dye_production_file(file_bytes) if file_type == DYE_PRODUCTION_FILE_TYPE else parse_nc_report_file(file_bytes)
    total_rows = len(result["rows"]) + len(result["errors"])

    conn = get_db()
    ensure_rft_source_tables(conn)
    _ensure_import_logs_table(conn)
    conn.commit()

    log_id = insert_returning_id(
        conn,
        "INSERT INTO import_logs (file_name, import_type, file_type, imported_by, total_rows, error_rows, status) VALUES (?, ?, ?, ?, ?, ?, 'processing')",
        (filename, file_type.lower(), file_type, imported_by or "unknown", total_rows, len(result["errors"])),
    )
    conn.commit()

    imported_rows = 0
    try:
        if result["rows"]:
            columns = fields + ("import_log_id",)
            placeholders = ",".join("?" for _ in columns)
            updates = ",".join(f"{field}=excluded.{field}" for field in columns if field not in key)
            sql = (
                f"INSERT INTO {table} ({','.join(columns)}) VALUES ({placeholders}) "
                f"ON CONFLICT({','.join(key)}) DO UPDATE SET {updates}"
            )
            # Khử trùng theo khoá (giữ dòng CUỐI) TRƯỚC executemany — Postgres gộp executemany
            # thành 1 INSERT, 2 dòng trùng khoá trong cùng statement bị từ chối
            # (`CardinalityViolation`), cùng lý do ở `sync_batch_details()`.
            deduped = {tuple(row[field] for field in key): row for row in result["rows"]}
            conn.execute("BEGIN")
            conn.executemany(sql, [tuple(row[field] for field in fields) + (log_id,) for row in deduped.values()])
            conn.commit()
            imported_rows = len(result["rows"])
        record_import_rows(conn, log_id, result["row_details"])
    except Exception as exc:
        conn.rollback()
        conn.execute("UPDATE import_logs SET status='failed', error_detail=? WHERE id=?", (str(exc), log_id))
        conn.commit()
        raise

    status = "completed" if imported_rows and not result["errors"] else "partial" if imported_rows else "failed"
    error_detail = "; ".join(f"Dòng {error['row']}: {error['error']}" for error in result["errors"][:20]) or None
    conn.execute(
        "UPDATE import_logs SET imported_rows=?, success_rows=?, status=?, error_detail=? WHERE id=?",
        (imported_rows, imported_rows, status, error_detail, log_id),
    )
    conn.commit()

    # KHÔNG gọi trigger_recompute(): Engine `rft` tính trực tiếp ở read time, không có rollup.
    return {
        "status": status,
        "file_type": file_type,
        "rows_imported": imported_rows,
        "total_rows": total_rows,
        "errors": result["errors"],
        "preview": result["rows"][:5],
        "columns": list(result["rows"][0].keys()) if result["rows"] else [],
        "imported_rows": imported_rows,
        "import_log_id": log_id,
    }


def sync_dye_production_ops(file_bytes: bytes, imported_by: str | None = None, filename: str = "production_report_dye.xlsx") -> dict[str, Any]:
    """UPSERT Production Report Dye vào `dye_production_ops` (khoá batch_no+operation+op_start_time)."""
    return _sync(file_bytes, imported_by, filename, DYE_PRODUCTION_FILE_TYPE, "dye_production_ops", DYE_PRODUCTION_FIELDS, DYE_PRODUCTION_KEY)


def sync_nc_reports(file_bytes: bytes, imported_by: str | None = None, filename: str = "nc_report.xlsx") -> dict[str, Any]:
    """UPSERT NC Report vào `dye_nc_reports` (khoá nc_no)."""
    return _sync(file_bytes, imported_by, filename, NC_REPORT_FILE_TYPE, "dye_nc_reports", NC_REPORT_FIELDS, ("nc_no",))

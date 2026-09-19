"""Right First Time (RFT) report Excel importer — UPSERT vào `rft_dye_results`.

Cùng pattern với `core/batch_importer.py` (module `core/` riêng, KHÔNG dùng chung
`detect_and_parse_file()`/`save_to_db()` generic của `core/excel_importer.py` — file RFT
report có cấu trúc/ngữ nghĩa hoàn toàn khác Availability/Performance). Nguồn: file
"RFT report.xlsx" do QC xuất, 13 cột (Customer, Color, OrderNo, GreigeCode, Dyelot,
MachineType, NC-DG, ResultDYE, NewBatch2, Rework Count, Stage, recipe, body/rib).
"""
from __future__ import annotations

import io
import math
from datetime import datetime
from typing import Any

from openpyxl import load_workbook

from core.database import get_db, get_dialect, insert_returning_id
from core.excel_importer import record_import_rows

RFT_HEADER_MAP: dict[str, str] = {
    "Customer": "customer",
    "Color": "color",
    "OrderNo": "order_no", "Order No": "order_no",
    "GreigeCode": "greige_code", "Greige Code": "greige_code",
    "Dyelot": "dyelot", "Dyelot No": "dyelot",
    "MachineType": "machine_type", "Machine Type": "machine_type",
    "NC-DG": "nc_dg", "NC DG": "nc_dg", "NCDG": "nc_dg",
    "ResultDYE": "result_dye", "Result DYE": "result_dye", "ResultDye": "result_dye", "Result Dye": "result_dye",
    "NewBatch2": "new_batch2", "New Batch2": "new_batch2", "New Batch 2": "new_batch2",
    "Rework Count": "rework_count", "ReworkCount": "rework_count",
    "Stage": "stage",
    "recipe": "recipe", "Recipe": "recipe",
    "body/rib": "body_rib", "Body/Rib": "body_rib", "Body/rib": "body_rib", "BodyRib": "body_rib",
}
RFT_FIELDS: tuple[str, ...] = (
    "dyelot", "customer", "color", "order_no", "greige_code", "machine_type", "nc_dg",
    "result_dye", "new_batch2", "rework_count", "stage", "recipe", "body_rib",
)
# Giá trị chuẩn hoá (bỏ khoảng trắng, viết thường) của MachineType máy lớn — dùng để so sánh
# nhất quán ở cả import (không validate enum) lẫn `rft/service.py` (giới hạn tab Rework/Adjust
# Color chỉ tính máy lớn theo yêu cầu người dùng).
LARGE_MACHINE_KEY = ">=500kg"
RESULT_DYE_VALUES = {"OK", "NG"}


def _clean(value: Any) -> Any:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    text = str(value).strip()
    return None if not text or text.lower() in {"nan", "none", "nat"} else value


def _read_xlsx(file_bytes: bytes) -> tuple[list[str], list[list[Any]]]:
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    try:
        rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return [], []
    return [str(value).strip() if value is not None else "" for value in rows[0]], [list(row) for row in rows[1:]]


def _build_rft_raw_row(indexes: dict[str, int], values: list[Any]) -> dict[str, str]:
    raw: dict[str, str] = {}
    for field in RFT_FIELDS:
        index = indexes.get(field)
        value = values[index] if index is not None and index < len(values) else None
        raw[field] = str(_clean(value) or "").strip()
    return raw


def _parse_rft_row(raw: dict[str, Any]) -> tuple[dict[str, Any] | None, str | None]:
    """Ép kiểu + validate 1 dòng RFT report ở dạng raw dict — dùng chung cho import hàng
    loạt (`parse_rft_file`) và revalidate 1 dòng qua Raw Data Viewer (nếu sau này cần)."""
    record = {field: raw.get(field, "") for field in RFT_FIELDS}
    if not record["dyelot"]:
        return None, "Thiếu giá trị bắt buộc: Dyelot."
    if not record["stage"]:
        return None, "Thiếu giá trị bắt buộc: Stage."
    result_dye = record["result_dye"].strip().upper()
    if result_dye not in RESULT_DYE_VALUES:
        return None, "ResultDYE không hợp lệ (chỉ nhận OK/NG)."
    record["result_dye"] = result_dye
    return record, None


def parse_rft_file(file_bytes: bytes) -> dict[str, Any]:
    headers, raw_rows = _read_xlsx(file_bytes)
    normalized = {header.lower(): index for index, header in enumerate(headers)}
    indexes = {field: normalized[header.lower()] for header, field in RFT_HEADER_MAP.items() if header.lower() in normalized}
    if "dyelot" not in indexes:
        raise ValueError("File RFT report thiếu cột bắt buộc Dyelot.")
    if "stage" not in indexes:
        raise ValueError("File RFT report thiếu cột bắt buộc Stage.")
    parsed: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    row_details: list[dict[str, Any]] = []
    for row_number, values in enumerate(raw_rows, start=2):
        if not any(_clean(value) is not None for value in values):
            continue
        raw = _build_rft_raw_row(indexes, values)
        record, error = _parse_rft_row(raw)
        if error is not None:
            errors.append({"row": row_number, "error": error})
            row_details.append({"row_number": row_number, "status": "invalid", "error": error, "data": raw})
            continue
        parsed.append(record)
        row_details.append({"row_number": row_number, "status": "valid", "error": None, "data": raw})
    return {"rows": parsed, "errors": errors, "total_records": len(parsed) + len(errors), "row_details": row_details}


def _ensure_import_logs_table(conn: Any) -> None:
    """Duplicate có chủ đích của `core/batch_importer.py::_ensure_import_logs_table()` — mỗi
    importer độc lập tự đảm bảo bảng chung tồn tại (cùng convention đã áp dụng ở
    `excel_import/service.py::import_raw_file()` và `batch_importer.py`, không import chéo
    hàm private giữa các module `core/`)."""
    if get_dialect() == "sqlite":
        conn.execute("""
            CREATE TABLE IF NOT EXISTS import_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT, file_name TEXT NOT NULL,
                import_type TEXT NOT NULL, imported_by TEXT NOT NULL,
                total_rows INTEGER NOT NULL DEFAULT 0, success_rows INTEGER NOT NULL DEFAULT 0,
                error_rows INTEGER NOT NULL DEFAULT 0, status TEXT NOT NULL DEFAULT 'completed',
                error_detail TEXT, imported_at TEXT NOT NULL DEFAULT (datetime('now')),
                imported_rows INTEGER NOT NULL DEFAULT 0, file_type TEXT,
                created_at TEXT NOT NULL DEFAULT (datetime('now'))
            )
        """)
    existing = {row["name"] for row in conn.execute("PRAGMA table_info(import_logs)")}
    for column, definition in (("imported_rows", "INTEGER NOT NULL DEFAULT 0"), ("file_type", "TEXT"), ("created_at", "TEXT")):
        if column not in existing:
            conn.execute(f"ALTER TABLE import_logs ADD COLUMN {column} {definition}")


def _ensure_rft_dye_results_table(conn: Any) -> None:
    if get_dialect() == "sqlite":
        conn.execute("""
            CREATE TABLE IF NOT EXISTS rft_dye_results (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                dyelot TEXT NOT NULL UNIQUE,
                customer TEXT, color TEXT, order_no TEXT, greige_code TEXT,
                machine_type TEXT, nc_dg TEXT, result_dye TEXT, new_batch2 TEXT,
                rework_count TEXT, stage TEXT NOT NULL, recipe TEXT, body_rib TEXT,
                import_log_id INTEGER,
                created_at TEXT NOT NULL DEFAULT (datetime('now'))
            )
        """)
    existing_columns = {row["name"] for row in conn.execute("PRAGMA table_info(rft_dye_results)")}
    for field in RFT_FIELDS:
        if field not in existing_columns:
            conn.execute(f"ALTER TABLE rft_dye_results ADD COLUMN {field} TEXT")
    if "import_log_id" not in existing_columns:
        conn.execute("ALTER TABLE rft_dye_results ADD COLUMN import_log_id INTEGER")


def sync_rft_results(file_bytes: bytes, imported_by: str | None = None, filename: str = "rft_report.xlsx") -> dict[str, Any]:
    """Parse file RFT report và UPSERT vào `rft_dye_results` (khoá `dyelot`) — cùng pattern
    ghi `import_logs` 'processing' trước/sau của `sync_batch_details()`."""
    result = parse_rft_file(file_bytes)
    total_rows = len(result["rows"]) + len(result["errors"])

    conn = get_db()
    _ensure_rft_dye_results_table(conn)
    _ensure_import_logs_table(conn)
    conn.commit()

    log_id = insert_returning_id(
        conn,
        "INSERT INTO import_logs (file_name, import_type, file_type, imported_by, total_rows, error_rows, status) VALUES (?, ?, ?, ?, ?, ?, 'processing')",
        (filename, "rft", "RFT", imported_by or "unknown", total_rows, len(result["errors"])),
    )
    conn.commit()

    imported_rows = 0
    try:
        if result["rows"]:
            columns_with_log = RFT_FIELDS + ("import_log_id",)
            placeholders = ",".join("?" for _ in columns_with_log)
            updates = ",".join(f"{field}=excluded.{field}" for field in columns_with_log if field != "dyelot")
            sql = f"INSERT INTO rft_dye_results ({','.join(columns_with_log)}) VALUES ({placeholders}) ON CONFLICT(dyelot) DO UPDATE SET {updates}"
            # Khử trùng theo dyelot (giữ dòng CUỐI xuất hiện trong file) TRƯỚC khi executemany —
            # cùng lý do đã áp dụng ở `sync_batch_details()`: Postgres gộp executemany() thành 1
            # câu INSERT...VALUES duy nhất, 2 dòng trùng khoá UPSERT trong CÙNG 1 statement sẽ bị
            # từ chối (`CardinalityViolation`).
            deduped_by_dyelot = {row["dyelot"]: row for row in result["rows"]}
            conn.execute("BEGIN")
            conn.executemany(sql, [tuple(row[field] for field in RFT_FIELDS) + (log_id,) for row in deduped_by_dyelot.values()])
            conn.commit()
            imported_rows = len(result["rows"])
        record_import_rows(conn, log_id, result["row_details"])
    except Exception as exc:
        conn.rollback()
        conn.execute("UPDATE import_logs SET status='failed', error_detail=? WHERE id=?", (str(exc), log_id))
        conn.commit()
        raise

    status = "completed" if imported_rows and not result["errors"] else "partial" if imported_rows else "failed"
    error_detail = "; ".join(f"Dòng {error['row']}: {error['error']}" for error in result["errors"][:20]) or None
    conn.execute(
        "UPDATE import_logs SET imported_rows=?, success_rows=?, status=?, error_detail=? WHERE id=?",
        (imported_rows, imported_rows, status, error_detail, log_id),
    )
    conn.commit()

    # KHÔNG gọi trigger_recompute(): Engine `rft` chưa có Daily Rollup/recompute_daily (query
    # trực tiếp mỗi request, xem `modules/dyeing/engines/rft/service.py`).
    return {
        "status": status,
        "file_type": "RFT",
        "rows_imported": imported_rows,
        "total_rows": total_rows,
        "errors": result["errors"],
        "preview": result["rows"][:5],
        "columns": list(result["rows"][0].keys()) if result["rows"] else [],
        "imported_rows": imported_rows,
        "import_log_id": log_id,
    }

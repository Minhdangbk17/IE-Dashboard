"""
modules/knitting/engines/downtime/report.py
----------------------------------------------
Báo cáo % Downtime xưởng Dệt (bảng người dùng chốt 2026-10-09): 13 nhóm dừng x kỳ (Day / Week /
Month), cột Before + Target, dòng Total.

Công thức (đã đối chiếu pivot tháng 10 của người dùng, khớp từng ô):
    % ô = Σ Stop Time của nhóm / Σ Available (Plan PRD) của MỌI máy trong bộ lọc, mọi ngày thuộc kỳ.
Gộp Week/Month cộng TỬ và MẪU rồi mới chia — KHÔNG trung bình các % theo ngày (Chủ nhật chỉ có
nửa Plan, trung bình % sẽ lệch). Total = tổng các nhóm (kể cả "Unmapped" nếu có, để Total luôn
khớp tổng Stop Time của file). Target Total = tổng Target các nhóm (25.6% Before / 16.3% Target
trong bảng người dùng đều là tổng cột).

Tuần = tuần ISO (Thứ Hai -> Chủ nhật), nhãn "W40-Oct" lấy tháng của NGÀY THỨ NĂM (đã đối chiếu:
W40 = 28/9..4/10 khớp 15.3%, tuần CN->T7 ra 15.1% không khớp).

Nhóm của 1 Stop Code: bản ghi tay `knitting_stop_category_map` (admin) > `DEFAULT_CODE_MAP`
(người dùng xác nhận) > từ khoá trong Stop Description (`_DESCRIPTION_RULES`, cho mã chưa gặp ở
file mẫu như No Material / MC Adjustment / Cleaning (Scheduled) / Drop stitch) > "Unmapped".

Bộ lọc Program / Core program: Program của từng máy-ngày tính ở `programs.py` (từ Piece Produced
report + danh mục Greige -> Program); lọc giữ NGUYÊN máy-ngày (cả tử Stop Time lẫn mẫu Available).
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date, datetime, timedelta
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import programs as program_lookup
from .service import ensure_tables

CATEGORIES: tuple[str, ...] = (
    "Doffing + Cleaning",
    "Yarn Broken",
    "No Material",
    "Loading/Unloading Yarn",
    "Needle Broken",
    "MC Part Broken",
    "MC Adjustment",
    "Needle Cleaning",
    "Safe Door",
    "MC Set Up",
    "Cleaning (Scheduled)",
    "Drop stitch",
    "Others",
)
UNMAPPED = "Unmapped"
TOTAL = "Total"

# (Before %, 2026 Target %) — bảng người dùng gửi 2026-10-09.
DEFAULT_TARGETS: dict[str, tuple[float, float]] = {
    "Doffing + Cleaning": (7.5, 5.5),
    "Yarn Broken": (6.5, 4.5),
    "No Material": (0.8, 0.4),
    "Loading/Unloading Yarn": (0.6, 0.5),
    "Needle Broken": (0.6, 0.4),
    "MC Part Broken": (0.4, 0.3),
    "MC Adjustment": (0.4, 0.3),
    "Needle Cleaning": (0.8, 0.5),
    "Safe Door": (5.0, 1.5),
    "MC Set Up": (2.0, 0.5),
    "Cleaning (Scheduled)": (0.0, 0.8),
    "Drop stitch": (0.0, 0.3),
    "Others": (1.0, 0.8),
}

# Stop Code -> nhóm, người dùng xác nhận 2026-10-09 (Others: 18 Others, 14 Waiting Tool,
# 25 "19 Ready to run", 17 Quality Issue, 271 Shift Change, Bad Material).
DEFAULT_CODE_MAP: dict[str, str] = {
    "10": "Doffing + Cleaning", "257": "Doffing + Cleaning", "11": "Doffing + Cleaning",
    "259": "Yarn Broken", "258": "Yarn Broken", "3": "Yarn Broken",
    "9": "Loading/Unloading Yarn",
    "1": "Needle Broken",
    "2": "MC Part Broken",
    "6": "MC Set Up",
    "12": "Needle Cleaning",
    "256": "Safe Door",
    "18": "Others", "14": "Others", "25": "Others", "17": "Others", "271": "Others",
}

# Từ khoá Stop Description (viết thường) -> nhóm, xét THEO THỨ TỰ (cụ thể trước, chung sau).
_DESCRIPTION_RULES: tuple[tuple[str, str], ...] = (
    ("scheduled", "Cleaning (Scheduled)"),
    ("drop stitch", "Drop stitch"),
    ("no material", "No Material"),
    ("bad material", "Others"),
    ("adjust", "MC Adjustment"),
    ("needle clean", "Needle Cleaning"),
    ("needle broken", "Needle Broken"),
    ("yarn broken", "Yarn Broken"),
    ("loading", "Loading/Unloading Yarn"),
    ("part broken", "MC Part Broken"),
    ("set up", "MC Set Up"),
    ("safe door", "Safe Door"),
    ("doffing", "Doffing + Cleaning"),
    ("other", "Others"),
)

GROUP_BYS = ("date", "week", "month")


# ---------------------------------------------------------------------------
# Nhóm dừng + Target
# ---------------------------------------------------------------------------


def _ensure_seeded(conn: Any) -> None:
    ensure_tables(conn)
    if conn.execute("SELECT COUNT(*) AS n FROM knitting_downtime_targets").fetchone()["n"] == 0:
        conn.executemany(
            "INSERT INTO knitting_downtime_targets (category, before_pct, target_pct, updated_by) VALUES (?, ?, ?, 'seed') "
            "ON CONFLICT (category) DO NOTHING",
            [(category, before, target) for category, (before, target) in DEFAULT_TARGETS.items()],
        )
    conn.commit()


def resolve_category(stop_code: str, description: str | None, overrides: dict[str, str]) -> tuple[str, str]:
    """(nhóm, nguồn) — nguồn: manual / default / rule / unmapped."""
    code = str(stop_code or "").strip()
    if code in overrides:
        return overrides[code], "manual"
    if code in DEFAULT_CODE_MAP:
        return DEFAULT_CODE_MAP[code], "default"
    text = str(description or "").lower()
    for keyword, category in _DESCRIPTION_RULES:
        if keyword in text:
            return category, "rule"
    return UNMAPPED, "unmapped"


def _overrides(conn: Any) -> dict[str, str]:
    return {row["stop_code"]: row["category"] for row in conn.execute("SELECT stop_code, category FROM knitting_stop_category_map")}


def load_targets(conn: Any) -> dict[str, dict[str, float | None]]:
    _ensure_seeded(conn)
    rows = conn.execute("SELECT category, before_pct, target_pct FROM knitting_downtime_targets").fetchall()
    return {row["category"]: {"before": row["before_pct"], "target": row["target_pct"]} for row in rows}


def set_target(conn: Any, category: str, before_pct: float | None, target_pct: float | None, user: str | None) -> None:
    if category not in CATEGORIES:
        raise ValueError(f"Nhóm không hợp lệ: {category}")
    for value in (before_pct, target_pct):
        if value is not None and not 0 <= value <= 100:
            raise ValueError("Before / Target phải trong khoảng 0–100 (%).")
    _ensure_seeded(conn)
    conn.execute(
        "INSERT INTO knitting_downtime_targets (category, before_pct, target_pct, updated_by, updated_at) "
        "VALUES (?, ?, ?, ?, ?) ON CONFLICT (category) DO UPDATE SET before_pct = excluded.before_pct, "
        "target_pct = excluded.target_pct, updated_by = excluded.updated_by, updated_at = excluded.updated_at",
        (category, before_pct, target_pct, user or "unknown", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    )
    conn.commit()


def list_stop_codes(conn: Any) -> list[dict[str, Any]]:
    """Mọi Stop Code đã từng import + nhóm hiện tại + nguồn — cho tab Stop Code Mapping."""
    ensure_tables(conn)
    overrides = _overrides(conn)
    rows = conn.execute(
        """
        SELECT stop_code, MAX(stop_description) AS stop_description, SUM(stop_time) AS stop_time,
               COUNT(*) AS records, MIN(production_date) AS first_date, MAX(production_date) AS last_date
        FROM knitting_stop_details GROUP BY stop_code
        """
    ).fetchall()
    result = []
    for row in rows:
        category, source = resolve_category(row["stop_code"], row["stop_description"], overrides)
        default_category, _ = resolve_category(row["stop_code"], row["stop_description"], {})
        result.append({**dict(row), "category": category, "source": source, "default_category": default_category})
    order = {name: index for index, name in enumerate(CATEGORIES + (UNMAPPED,))}
    return sorted(result, key=lambda r: (order.get(r["category"], 99), -(r["stop_time"] or 0)))


def set_stop_category(conn: Any, stop_code: str, category: str | None, user: str | None) -> None:
    """Ghi đè nhóm cho 1 Stop Code; `category` rỗng/None = bỏ ghi đè, quay về mặc định."""
    code = str(stop_code or "").strip()
    if not code:
        raise ValueError("Thiếu Stop Code.")
    ensure_tables(conn)
    if not category:
        conn.execute("DELETE FROM knitting_stop_category_map WHERE stop_code = ?", (code,))
    else:
        if category not in CATEGORIES:
            raise ValueError(f"Nhóm không hợp lệ: {category}")
        conn.execute(
            "INSERT INTO knitting_stop_category_map (stop_code, category, updated_by, updated_at) VALUES (?, ?, ?, ?) "
            "ON CONFLICT (stop_code) DO UPDATE SET category = excluded.category, updated_by = excluded.updated_by, "
            "updated_at = excluded.updated_at",
            (code, category, user or "unknown", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        )
    conn.commit()


# ---------------------------------------------------------------------------
# Bộ lọc + kỳ
# ---------------------------------------------------------------------------


def _split(value: str | Iterable[str] | None) -> list[str]:
    if value is None:
        return []
    items = value.split(",") if isinstance(value, str) else list(value)
    return sorted({str(item).strip() for item in items if str(item).strip()})


def _parse_day(value: str | None, field: str) -> date:
    try:
        return date.fromisoformat(str(value))
    except ValueError as exc:
        raise ValueError(f"{field} không hợp lệ: {value!r}") from exc


def normalize_filters(
    from_date: str | None, to_date: str | None, group_by: str | None,
    machines: str | Iterable[str] | None = None, structures: str | Iterable[str] | None = None,
    programs: str | Iterable[str] | None = None, core_only: bool | str | None = False,
) -> dict[str, Any]:
    if not from_date or not to_date:
        raise ValueError("Cần From Date và To Date.")
    start, end = _parse_day(from_date, "From Date"), _parse_day(to_date, "To Date")
    if start > end:
        raise ValueError("From Date phải nhỏ hơn hoặc bằng To Date.")
    return {
        "from_date": start.isoformat(), "to_date": end.isoformat(),
        "group_by": group_by if group_by in GROUP_BYS else "week",
        "machines": _split(machines), "structures": _split(structures),
        # Tên Program có thể chứa dấu phẩy -> nhận list, hoặc chuỗi phân tách bằng "|".
        "programs": _split(programs.split("|") if isinstance(programs, str) else programs),
        "core_only": str(core_only).lower() in {"1", "true", "yes", "on"},
    }


def _machine_where(filters: dict[str, Any], alias: str = "m") -> tuple[str, list[Any]]:
    clauses = [f"{alias}.production_date BETWEEN ? AND ?"]
    params: list[Any] = [filters["from_date"], filters["to_date"]]
    if filters["machines"]:
        clauses.append(f"{alias}.machine_code IN ({','.join('?' for _ in filters['machines'])})")
        params += filters["machines"]
    if filters["structures"]:
        clauses.append(f"{alias}.knitting_structure IN ({','.join('?' for _ in filters['structures'])})")
        params += filters["structures"]
    return " AND ".join(clauses), params


def period_key(day: date, group_by: str) -> str:
    if group_by == "month":
        return day.strftime("%Y-%m")
    if group_by == "week":
        year, week, _ = day.isocalendar()
        return f"{year}-W{week:02d}"
    return day.isoformat()


def period_range(key: str, group_by: str) -> tuple[date, date]:
    if group_by == "month":
        first = date.fromisoformat(f"{key}-01")
        following = (first.replace(day=28) + timedelta(days=4)).replace(day=1)
        return first, following - timedelta(days=1)
    if group_by == "week":
        year, week = key.split("-W")
        monday = date.fromisocalendar(int(year), int(week), 1)
        return monday, monday + timedelta(days=6)
    day = date.fromisoformat(key)
    return day, day


def _period_label(key: str, group_by: str, with_year: bool) -> str:
    """Nhãn theo bảng người dùng: "01-Oct" / "W40-Oct" (tháng của ngày thứ Năm) / "Oct"."""
    suffix_fmt = "-%y" if with_year else ""
    if group_by == "month":
        return date.fromisoformat(f"{key}-01").strftime("%b" + suffix_fmt)
    if group_by == "week":
        year, week = key.split("-W")
        thursday = date.fromisocalendar(int(year), int(week), 4)
        return f"W{int(week):02d}-" + thursday.strftime("%b" + suffix_fmt)
    return date.fromisoformat(key).strftime("%d-%b" + suffix_fmt)


def list_filter_options(conn: Any) -> dict[str, list[str]]:
    ensure_tables(conn)
    machines = [r["machine_code"] for r in conn.execute("SELECT DISTINCT machine_code FROM knitting_machine_daily ORDER BY machine_code")]
    structures = [
        r["knitting_structure"] for r in conn.execute(
            "SELECT DISTINCT knitting_structure FROM knitting_machine_daily WHERE knitting_structure IS NOT NULL ORDER BY knitting_structure"
        )
    ]
    latest = conn.execute("SELECT MAX(production_date) AS d FROM knitting_machine_daily").fetchone()["d"]
    return {"machines": machines, "structures": structures, "latest_date": latest, **program_lookup.list_program_options(conn)}


def _program_filter(conn: Any, filters: dict[str, Any]) -> tuple[set[str] | None, dict[tuple[str, str], dict[str, Any]]]:
    """(tập khoá Program được giữ hoặc None, bảng Program theo máy-ngày — chỉ tính khi cần lọc)."""
    allowed = program_lookup.allowed_program_keys(conn, filters["programs"], filters["core_only"])
    if allowed is None:
        return None, {}
    return allowed, program_lookup.machine_day_programs(conn, filters["from_date"], filters["to_date"], filters["machines"] or None)


def _keep(allowed: set[str] | None, mapping: dict[tuple[str, str], dict[str, Any]], day: str, machine: str) -> bool:
    return allowed is None or program_lookup.program_of(mapping, day, machine)["program_key"] in allowed


# ---------------------------------------------------------------------------
# Báo cáo
# ---------------------------------------------------------------------------


def _pct(stop: float, available: float) -> float | None:
    return stop / available * 100 if available else None


def build_report(conn: Any, filters: dict[str, Any]) -> dict[str, Any]:
    _ensure_seeded(conn)
    group_by = filters["group_by"]
    where, params = _machine_where(filters)
    allowed, mapping = _program_filter(conn, filters)

    available_by_period: dict[str, float] = defaultdict(float)
    for row in conn.execute(
        f"SELECT m.production_date AS d, m.machine_code, m.available_time FROM knitting_machine_daily m WHERE {where}",
        params,
    ):
        if _keep(allowed, mapping, row["d"], row["machine_code"]):
            available_by_period[period_key(date.fromisoformat(row["d"]), group_by)] += row["available_time"] or 0

    overrides = _overrides(conn)
    stop_by_cell: dict[tuple[str, str], float] = defaultdict(float)
    unmapped_codes: dict[str, str] = {}
    for row in conn.execute(
        f"""
        SELECT s.production_date AS d, s.machine_code, s.stop_code, s.stop_description, s.stop_time
        FROM knitting_stop_details s
        JOIN knitting_machine_daily m ON m.production_date = s.production_date AND m.machine_code = s.machine_code
        WHERE {where}
        """,
        params,
    ):
        if not _keep(allowed, mapping, row["d"], row["machine_code"]):
            continue
        category, _source = resolve_category(row["stop_code"], row["stop_description"], overrides)
        if category == UNMAPPED:
            unmapped_codes[row["stop_code"]] = row["stop_description"] or ""
        stop_by_cell[(category, period_key(date.fromisoformat(row["d"]), group_by))] += row["stop_time"] or 0

    keys = sorted(available_by_period)
    years = {period_range(key, group_by)[0].year for key in keys} | {period_range(key, group_by)[1].year for key in keys}
    labels = [_period_label(key, group_by, len(years) > 1) for key in keys]
    total_available = sum(available_by_period.values())
    targets = load_targets(conn)

    categories = list(CATEGORIES) + ([UNMAPPED] if unmapped_codes else [])
    rows = []
    for category in categories:
        times = [stop_by_cell.get((category, key), 0.0) for key in keys]
        rows.append({
            "category": category,
            "before": targets.get(category, {}).get("before"),
            "target": targets.get(category, {}).get("target"),
            "stop_time": times,
            "pct": [_pct(t, available_by_period[key]) for t, key in zip(times, keys)],
            "total_stop_time": sum(times),
            "total_pct": _pct(sum(times), total_available),
        })
    total_times = [sum(row["stop_time"][i] for row in rows) for i in range(len(keys))]
    total_row = {
        "category": TOTAL,
        "before": sum(t["before"] or 0 for c, t in targets.items() if c in CATEGORIES),
        "target": sum(t["target"] or 0 for c, t in targets.items() if c in CATEGORIES),
        "stop_time": total_times,
        "pct": [_pct(t, available_by_period[key]) for t, key in zip(total_times, keys)],
        "total_stop_time": sum(total_times),
        "total_pct": _pct(sum(total_times), total_available),
    }
    return {
        "filters": filters,
        "period_keys": keys,
        "periods": labels,
        "rows": rows,
        "total_row": total_row,
        "available": [available_by_period[key] for key in keys],
        "total_available": total_available,
        "unmapped_codes": [{"stop_code": code, "stop_description": desc} for code, desc in sorted(unmapped_codes.items())],
    }


def get_cell_details(conn: Any, filters: dict[str, Any], period: str, category: str) -> dict[str, Any]:
    """Các dòng (ngày x máy x mã dừng) tạo nên 1 ô — tổng Stop Time PHẢI khớp ô đó.
    `category` = "Total" -> mọi nhóm; `period` = "ALL" -> cột Total (cả khoảng lọc)."""
    if category not in CATEGORIES + (UNMAPPED, TOTAL):
        raise ValueError(f"Nhóm không hợp lệ: {category}")
    if period == "ALL":
        start, end = date.fromisoformat(filters["from_date"]), date.fromisoformat(filters["to_date"])
    else:
        try:
            start, end = period_range(period, filters["group_by"])
        except ValueError as exc:
            raise ValueError(f"Kỳ không hợp lệ: {period!r}") from exc
    cell_filters = {
        **filters,
        "from_date": max(start.isoformat(), filters["from_date"]),
        "to_date": min(end.isoformat(), filters["to_date"]),
    }
    ensure_tables(conn)
    where, params = _machine_where(cell_filters)
    allowed = program_lookup.allowed_program_keys(conn, filters["programs"], filters["core_only"])
    mapping = program_lookup.machine_day_programs(conn, cell_filters["from_date"], cell_filters["to_date"], filters["machines"] or None)
    machine_days = [
        row for row in conn.execute(f"SELECT m.production_date AS d, m.machine_code, m.available_time FROM knitting_machine_daily m WHERE {where}", params)
        if _keep(allowed, mapping, row["d"], row["machine_code"])
    ]
    available = {"a": sum(r["available_time"] or 0 for r in machine_days), "n": len(machine_days)}
    overrides = _overrides(conn)
    rows = []
    for row in conn.execute(
        f"""
        SELECT s.production_date, s.machine_code, m.knitting_structure, s.stop_code, s.stop_description,
               s.stop_time, s.stop_count
        FROM knitting_stop_details s
        JOIN knitting_machine_daily m ON m.production_date = s.production_date AND m.machine_code = s.machine_code
        WHERE {where}
        ORDER BY s.stop_time DESC, s.production_date, s.machine_code
        """,
        params,
    ):
        if not _keep(allowed, mapping, row["production_date"], row["machine_code"]):
            continue
        row_category, _ = resolve_category(row["stop_code"], row["stop_description"], overrides)
        if category == TOTAL or row_category == category:
            program = program_lookup.program_of(mapping, row["production_date"], row["machine_code"])
            rows.append({**dict(row), "category": row_category, "program": program["program"], "greige_id": program["greige_id"]})
    stop_total = sum(r["stop_time"] or 0 for r in rows)
    return {
        "period": period,
        "category": category,
        "from_date": cell_filters["from_date"],
        "to_date": cell_filters["to_date"],
        "available": available["a"] or 0,
        "machine_days": available["n"] or 0,
        "stop_time": stop_total,
        "pct": _pct(stop_total, available["a"] or 0),
        "rows": rows,
    }


# ---------------------------------------------------------------------------
# Export Excel
# ---------------------------------------------------------------------------

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="24292F", end_color="24292F", fill_type="solid")
_BAD_FONT = Font(color="9C0006")
_BAD_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
_BOLD = Font(bold=True)


def _write_header(sheet: Any, headers: list[str]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL


def _write_pivot(sheet: Any, report: dict[str, Any], as_pct: bool) -> None:
    _write_header(sheet, ["Downtime (%)" if as_pct else "Stop Time", "Before", "Target", *report["periods"], "Total"])
    for row in report["rows"] + [report["total_row"]]:
        values = row["pct"] if as_pct else row["stop_time"]
        total = row["total_pct"] if as_pct else row["total_stop_time"]
        to_ratio = (lambda v: None if v is None else v / 100) if as_pct else (lambda v: v)
        sheet.append([row["category"], to_ratio(row["before"]), to_ratio(row["target"]), *[to_ratio(v) for v in values], to_ratio(total)])
        excel_row = sheet.max_row
        for col in range(2, sheet.max_column + 1):
            cell = sheet.cell(row=excel_row, column=col)
            cell.number_format = "0.0%" if as_pct or col <= 3 else "#,##0.00"
            if col > 3 and as_pct and row["target"] is not None and cell.value is not None and cell.value * 100 > row["target"] + 1e-9:
                cell.font, cell.fill = _BAD_FONT, _BAD_FILL
            if not as_pct and col <= 3 and cell.value is not None:
                cell.value = cell.value / 100
        if row["category"] == TOTAL:
            sheet.cell(row=excel_row, column=1).font = _BOLD
    if not as_pct:
        sheet.append(["Plan PRD (Available)", None, None, *report["available"], report["total_available"]])
        for col in range(4, sheet.max_column + 1):
            sheet.cell(row=sheet.max_row, column=col).number_format = "#,##0.00"
    sheet.column_dimensions["A"].width = 24
    for col in range(2, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 11
    sheet.freeze_panes = "D2"


def export_excel(conn: Any, filters: dict[str, Any]) -> bytes:
    report = build_report(conn, filters)
    workbook = Workbook()
    _write_pivot(workbook.active, report, as_pct=True)
    workbook.active.title = "Downtime %"
    _write_pivot(workbook.create_sheet("Stop Time"), report, as_pct=False)

    data = workbook.create_sheet("Data")
    _write_header(data, ["Production Date", "M/c Code", "Knitting Structure", "Greige ID", "Program", "Stop Code", "Stop Description", "Category", "Stop Time", "Stop #"])
    overrides = _overrides(conn)
    where, params = _machine_where(filters)
    allowed = program_lookup.allowed_program_keys(conn, filters["programs"], filters["core_only"])
    mapping = program_lookup.machine_day_programs(conn, filters["from_date"], filters["to_date"], filters["machines"] or None)
    for row in conn.execute(
        f"""
        SELECT s.production_date, s.machine_code, m.knitting_structure, s.stop_code, s.stop_description, s.stop_time, s.stop_count
        FROM knitting_stop_details s
        JOIN knitting_machine_daily m ON m.production_date = s.production_date AND m.machine_code = s.machine_code
        WHERE {where}
        ORDER BY s.production_date, s.machine_code, s.stop_code
        """,
        params,
    ):
        if not _keep(allowed, mapping, row["production_date"], row["machine_code"]):
            continue
        category, _ = resolve_category(row["stop_code"], row["stop_description"], overrides)
        program = program_lookup.program_of(mapping, row["production_date"], row["machine_code"])
        data.append([row["production_date"], row["machine_code"], row["knitting_structure"], program["greige_id"], program["program"],
                     row["stop_code"], row["stop_description"], category, row["stop_time"], row["stop_count"]])
    if data.max_row > 1:
        data.auto_filter.ref = data.dimensions
    data.freeze_panes = "A2"
    for col, width in zip("ABCDEFGHIJ", (14, 11, 22, 16, 30, 10, 26, 22, 11, 8)):
        data.column_dimensions[col].width = width

    info = workbook.create_sheet("Filters")
    _write_header(info, ["Filter", "Value"])
    for label, value in (
        ("From Date", filters["from_date"]), ("To Date", filters["to_date"]),
        ("Group By", {"date": "Day", "week": "Week", "month": "Month"}[filters["group_by"]]),
        ("M/c Code", ", ".join(filters["machines"]) or "All"),
        ("Knitting Structure", ", ".join(filters["structures"]) or "All"),
        ("Program", " | ".join(filters["programs"]) or "All"),
        ("Core program only", "Yes" if filters["core_only"] else "No"),
        ("Program rule", "Program of a machine-day = Greige ID of Piece Produced rolls covering that day (most Available minutes); no roll or Greige not in program list = blank"),
        ("Formula", "% = Σ Stop Time of category / Σ Available (Plan PRD) of all filtered machine-days in the period"),
        ("Exported at", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ):
        info.append([label, value])
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 90

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()

"""
tests/test_batch_matrix_brand_fabric_filters.py
--------------------------------------------------
Verify bộ lọc của báo cáo Batch/Day - Fabric/Color Matrix. 2026-10-06: bộ lọc GIỐNG HỆT tab
Batch/Day Trend (Capacity, Brand Program, Tank Type, Group By Day/Week/Month, ReDye = 0; bỏ
filter Fabric Type), mẻ Normal theo quy tắc Trend (Dyelot *0, SapLot 1*/3*, ReDye = 0).

Brand Program filter: RỦI RO CAO — 1 máy có thể chạy NHIỀU brand_program trong CÙNG 1 ngày,
mẫu số phải là giờ CẢ NGÀY của tập máy (không phải chỉ giờ các mẻ thuộc brand đó) và không
được cộng trùng giờ máy. Test bẫy ĐÚNG loại bug đó: M1 chạy 2 mẻ CÙNG ngày, CÙNG
Fabric/Color/Capacity, nhưng 2 Brand Program khác nhau.

Chạy: python tests/test_batch_matrix_brand_fabric_filters.py
"""
from __future__ import annotations

import io
import os
import shutil
import sqlite3
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from flask import Flask  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from core.database import close_db  # noqa: E402
from modules.dyeing.engines.batch_matrix.service import build_matrix, export_matrix_excel, get_cell_batches  # noqa: E402

DAY1, DAY2 = "2026-09-01", "2026-09-02"  # cùng tuần ISO W36


def _make_temp_app(db_path: str) -> Flask:
    app = Flask(__name__)
    app.config["DATABASE_PATH"] = db_path
    app.config["SQLITE_PRAGMAS"] = {}
    return app


def _init_db(db_path: str) -> None:
    conn = sqlite3.connect(db_path)
    conn.execute("""
        CREATE TABLE availability_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            batch TEXT, batch_ref_no TEXT, fabric_type TEXT, machine TEXT, capacity_kg REAL,
            start_time TEXT, end_time TEXT
        )
    """)
    conn.execute(
        "CREATE TABLE batch_details (id INTEGER PRIMARY KEY AUTOINCREMENT, dyelot TEXT NOT NULL, shade TEXT, colour_no TEXT, "
        "batch_type TEXT, greige_code TEXT, machine TEXT, start_time TEXT, end_time TEXT, sap_lot TEXT, redye REAL DEFAULT 0)"
    )
    conn.execute("CREATE TABLE machines (machine_code TEXT, machine_id TEXT, capacity_kg REAL, tank_type TEXT)")
    conn.execute("INSERT INTO machines (machine_code, capacity_kg, tank_type) VALUES ('M1', 600, 'J tank'), ('M2', 300, NULL)")
    # (batch, fabric, machine, capacity, start, end) — M1 9h/ngày (UQ 4h + Nike 5h), M2 6h ngày 1.
    conn.executemany(
        "INSERT INTO availability_logs (batch, fabric_type, machine, capacity_kg, start_time, end_time) VALUES (?, ?, ?, ?, ?, ?)",
        [
            ("M1-UQ0", "CVC", "M1", 600, "2026-09-01 07:00:00", "2026-09-01 11:00:00"),
            ("M1-NIKE0", "CVC", "M1", 600, "2026-09-01 11:00:00", "2026-09-01 16:00:00"),
            # Capacity của M2 trong Availability (999) CỐ TÌNH khác Machine Master (300) — báo cáo
            # phải dùng Machine Master.
            ("M2-UQ0", "CVC", "M2", 999, "2026-09-01 07:00:00", "2026-09-01 13:00:00"),
            # Ngày 2: M2 mẻ ReDye = 1 (8h) — chỉ là Normal khi TẮT checkbox ReDye = 0.
            ("M2-RD0", "CVC", "M2", 999, "2026-09-02 07:00:00", "2026-09-02 15:00:00"),
            # Mẻ Rework (Dyelot không kết thúc bằng 0) — có giờ máy, không được đếm.
            ("M2-RW1", "CVC", "M2", 999, "2026-09-02 15:00:00", "2026-09-02 16:00:00"),
        ],
    )
    conn.executemany(
        "INSERT INTO batch_details (dyelot, shade, colour_no, batch_type, greige_code, sap_lot, redye) VALUES (?, 'Dark', '091-NAVY', 'Normal', ?, '1000', ?)",
        [("M1-UQ0", "G1", 0), ("M1-NIKE0", "G2", 0), ("M2-UQ0", "G1", 0), ("M2-RD0", "G1", 1), ("M2-RW1", "G1", 0)],
    )
    conn.commit()
    conn.close()


def _check(label: str, actual, expected, failures: list[str], tolerance: float = 0.0) -> None:
    ok = actual is not None and abs(actual - expected) <= tolerance if tolerance else actual == expected
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}: actual={actual} expected={expected}")
    if not ok:
        failures.append(label)


def _cvc_dark(data: dict) -> dict:
    return next(r for r in data["rows"] if r["row_type"] == "data" and r["fabric_type"] == "CVC" and r["color_group"] == "Dark")


def main() -> int:
    failures: list[str] = []
    tmp_dir = tempfile.mkdtemp()
    try:
        db_path = os.path.join(tmp_dir, "test.db")
        _init_db(db_path)
        with _make_temp_app(db_path).app_context():
            from core.brand_program_importer import ensure_brand_program_table
            conn = sqlite3.connect(db_path)
            conn.row_factory = sqlite3.Row
            ensure_brand_program_table(conn)
            conn.execute("INSERT INTO brand_program_mapping (greige_code, brand, brand_program) VALUES ('G1', 'UQ', 'Ht Fleece')")
            conn.execute("INSERT INTO brand_program_mapping (greige_code, brand, brand_program) VALUES ('G2', 'Nike', 'Performance')")
            conn.commit()
            conn.close()

            print("=== Kịch bản 1: không lọc (Group By Day, ReDye = 0 bật) ===")
            data = build_matrix(group_by="date")
            # Ngày 1: 3 mẻ, tập máy {M1 9h, M2 6h} = 15h -> 4.8. Ngày 2: M2-RD0 (ReDye 1) + Rework -> không có mẻ.
            _check("ngày 1 CVC/Dark = 3*24/15", _cvc_dark(data)["values"][DAY1], 4.8, failures, tolerance=0.01)
            _check("ngày 2 trống (ReDye 1 + Rework không đếm)", _cvc_dark(data)["values"].get(DAY2), None, failures)
            _check("available_brand_programs", data["available_brand_programs"], ["Nike - Performance", "UQ - Ht Fleece"], failures)
            _check("available_tank_types", data["available_tank_types"], ["J tank", "Unclassified"], failures)
            _check("available_capacities lấy từ Machine Master (không có 999)", data["available_capacities"], [300.0, 600.0], failures)
            _check("không còn Fabric Type filter", "available_fabric_types" in data, False, failures)

            print("\n=== Kịch bản 2: Brand Program — mẫu số = giờ CẢ NGÀY của tập máy ===")
            # UQ: 2 mẻ (M1-UQ0, M2-UQ0), tập máy {M1, M2} = 9h + 6h -> 2*24/15 = 3.2 (M1 tính đủ 9h).
            _check("brand=UQ", _cvc_dark(build_matrix(brand_programs="UQ - Ht Fleece"))["values"][DAY1], 3.2, failures, tolerance=0.01)
            _check("brand=Nike (1*24/9)", _cvc_dark(build_matrix(brand_programs="Nike - Performance"))["values"][DAY1], 2.67, failures, tolerance=0.01)

            print("\n=== Kịch bản 3: Tank Type + Capacity ===")
            _check("tank=J tank: chỉ M1 (2*24/9)", _cvc_dark(build_matrix(tank_types="J tank"))["values"][DAY1], 5.33, failures, tolerance=0.01)
            _check("capacity=300: chỉ M2 (1*24/6)", _cvc_dark(build_matrix(capacities="300"))["values"][DAY1], 4.0, failures, tolerance=0.01)

            print("\n=== Kịch bản 4: ReDye = 0 TẮT + Group By Week ===")
            any_redye = build_matrix(group_by="date", require_redye_zero=False)
            _check("tắt ReDye: ngày 2 M2-RD0 được đếm, M2 = 9h (8h + 1h Rework) -> 1*24/9", _cvc_dark(any_redye)["values"][DAY2], 2.67, failures, tolerance=0.01)
            week = build_matrix(group_by="week")
            _check("week: 1 cột W36", week["period_keys"], ["2026-W36"], failures)
            _check("week (ReDye bật) = 3*24/15", _cvc_dark(week)["values"]["2026-W36"], 4.8, failures, tolerance=0.01)
            week_any = build_matrix(group_by="week", require_redye_zero=False)
            _check("week (ReDye tắt) = 4*24/(15+9)", _cvc_dark(week_any)["values"]["2026-W36"], 4.0, failures, tolerance=0.01)

            print("\n=== Kịch bản 5: drill-down + Export ===")
            _check("ô ngày 1 brand=UQ -> 2 mẻ", len(get_cell_batches("CVC", "Dark", DAY1, brand_programs="UQ - Ht Fleece")), 2, failures)
            _check("ô Total (period rỗng) -> 3 mẻ", len(get_cell_batches("CVC", "Dark", None)), 3, failures)
            _check("ô tuần ReDye tắt -> 4 mẻ", len(get_cell_batches("CVC", "Dark", "2026-W36", group_by="week", require_redye_zero=False)), 4, failures)
            workbook = load_workbook(io.BytesIO(export_matrix_excel(group_by="date")))
            _check("Excel sheets", workbook.sheetnames, ["Matrix", "Batches", "Filters"], failures)
            _check("Excel Batches: 3 mẻ Normal", workbook["Batches"].max_row - 1, 3, failures)
            matrix_row = next(r for r in workbook["Matrix"].iter_rows(min_row=2, values_only=True) if r[1] == "Dark")
            _check("Excel Matrix ô ngày 1", matrix_row[4], 4.8, failures, tolerance=0.01)
            close_db()
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    if failures:
        print(f"\n{len(failures)} FAILURE(S): {failures}")
        return 1
    print("\nAll checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())

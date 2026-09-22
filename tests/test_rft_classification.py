"""
tests/test_rft_classification.py
------------------------------------
Verify Engine `rft` sau khi có nguồn dữ liệu thật (`rft_dye_results`, import từ file
"RFT report.xlsx"): (1) `classify_rft_category()` map đúng 6 giá trị Stage, Stage lạ trả
`None`; (2) `parse_rft_file()` nhận diện đúng file mẫu thật; (3) công thức KPI mới
(rate_pct = OK / tổng mẻ CỦA CHÍNH TAB, không phải tỷ trọng so với 6 tab); (4) tab
Rework/Adjust Color chỉ tính máy >=500kg; (5) mẻ không khớp `availability_logs` rơi vào
bucket "Unknown Date" khi không lọc ngày, bị loại khi có lọc ngày tường minh.

Dùng DB TẠM (file SQLite tạm + Flask app context tạm, KHÔNG đụng DB thật) — cùng pattern
`tests/test_batch_matrix_formula.py`.

Chạy: python tests/test_rft_classification.py
"""
from __future__ import annotations

import os
import sqlite3
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from flask import Flask  # noqa: E402

from core.database import close_db  # noqa: E402
from core.excel_importer import detect_file_type_from_headers  # noqa: E402
from core.rft_importer import _ensure_rft_dye_results_table, parse_rft_file, sync_rft_results, RFT_FIELDS  # noqa: E402
from modules.dyeing.engines.rft.service import classify_rft_category, get_rft_pivot_data  # noqa: E402

FIXTURE_PATH = Path(__file__).resolve().parent / "fixtures" / "sample_imports" / "RFT report.xlsx"


def _make_temp_app(db_path: str) -> Flask:
    app = Flask(__name__)
    app.config["DATABASE_PATH"] = db_path
    app.config["SQLITE_PRAGMAS"] = {}
    return app


def _init_schema(db_path: str) -> None:
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    conn.execute("""
        CREATE TABLE availability_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            batch TEXT, capacity_kg REAL, fabric_type TEXT, start_time TEXT, end_time TEXT
        )
    """)
    # Bảng `batch_details` TỐI GIẢN (chỉ đủ cột `_rft_rows()` cần) — `rft/service.py` LEFT
    # JOIN bảng này (khoá dyelot=dyelot) làm nguồn fallback fabric_type khi
    # `availability_logs` không khớp (xem docstring `_rft_rows()`).
    conn.execute("""
        CREATE TABLE batch_details (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            dyelot TEXT NOT NULL, fabric_type TEXT, end_time TEXT
        )
    """)
    _ensure_rft_dye_results_table(conn)
    conn.commit()
    conn.close()


def _insert_rft_row(conn: sqlite3.Connection, **overrides: str) -> None:
    record = {field: "" for field in RFT_FIELDS}
    record.update(overrides)
    columns = ",".join(RFT_FIELDS)
    placeholders = ",".join("?" for _ in RFT_FIELDS)
    conn.execute(f"INSERT INTO rft_dye_results ({columns}) VALUES ({placeholders})", tuple(record[f] for f in RFT_FIELDS))


def _check(label: str, actual, expected, failures: list[str]) -> None:
    ok = actual == expected
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}: actual={actual!r} expected={expected!r}")
    if not ok:
        failures.append(label)


def _scenario_stage_mapping(failures: list[str]) -> None:
    print("\n=== Kịch bản 1: classify_rft_category() map đúng 6 Stage + Stage lạ -> None ===")
    cases = [
        ("Lab to Lab", "Lab to Lab"), (" lab to lab ", "Lab to Lab"),
        ("Lab to Bulk", "Lab to Bulk"), ("LAB TO BULK", "Lab to Bulk"),
        ("Bulk to Bulk", "Bulk to Bulk"),
        ("2nd batch", "2nd Batch"), ("2ND BATCH", "2nd Batch"),
        ("Rework", "Rework"),
        ("Adjust Color", "Adjust Color"), ("adjust color", "Adjust Color"),
        ("Something Else", None), ("", None),
    ]
    for stage_value, expected in cases:
        _check(f"stage={stage_value!r}", classify_rft_category({"stage": stage_value}), expected, failures)


def _scenario_parse_real_fixture(failures: list[str]) -> None:
    print("\n=== Kịch bản 2: parse_rft_file() + detect_file_type_from_headers() trên file mẫu thật ===")
    if not FIXTURE_PATH.exists():
        print(f"  [SKIP] Không tìm thấy fixture {FIXTURE_PATH}")
        return
    file_bytes = FIXTURE_PATH.read_bytes()
    result = parse_rft_file(file_bytes)
    _check("total_records", result["total_records"], 31, failures)
    _check("số dòng lỗi", len(result["errors"]), 0, failures)
    ok_count = sum(1 for row in result["rows"] if row["result_dye"] == "OK")
    ng_count = sum(1 for row in result["rows"] if row["result_dye"] == "NG")
    _check("số dòng ResultDYE=OK", ok_count, 21, failures)
    _check("số dòng ResultDYE=NG", ng_count, 10, failures)

    import openpyxl
    workbook = openpyxl.load_workbook(FIXTURE_PATH, read_only=True)
    headers = [str(cell.value).strip() for cell in next(workbook.worksheets[0].iter_rows(min_row=1, max_row=1))]
    workbook.close()
    _check("detect_file_type_from_headers() nhận đúng RFT", detect_file_type_from_headers(headers), "RFT", failures)


def _scenario_kpi_formula_and_machine_restriction(failures: list[str]) -> None:
    print("\n=== Kịch bản 3: công thức KPI OK/tổng-trong-tab + giới hạn >=500kg cho Rework/Adjust Color ===")
    fd, db_path = tempfile.mkstemp(suffix=".db")
    os.close(fd)
    try:
        _init_schema(db_path)
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row

        # Tab "Lab to Bulk": 3 mẻ, 2 OK 1 NG -> rate 66.7%. Cả 3 đều khớp availability_logs
        # (có production_date) để không rơi vào bucket "Unknown Date".
        conn.executemany(
            "INSERT INTO availability_logs (batch, capacity_kg, fabric_type, start_time, end_time) VALUES (?, ?, ?, ?, ?)",
            [
                ("LB-1", 600, "Cotton", "2026-09-01 08:00:00", "2026-09-01 10:00:00"),
                ("LB-2", 600, "Cotton", "2026-09-01 08:00:00", "2026-09-01 10:00:00"),
                ("LB-3", 600, "Cotton", "2026-09-01 08:00:00", "2026-09-01 10:00:00"),
            ],
        )
        _insert_rft_row(conn, dyelot="LB-1", stage="lab to bulk", result_dye="OK", machine_type=">=500kg")
        _insert_rft_row(conn, dyelot="LB-2", stage="lab to bulk", result_dye="OK", machine_type=">=500kg")
        _insert_rft_row(conn, dyelot="LB-3", stage="lab to bulk", result_dye="NG", machine_type=">=500kg")

        # Tab "Rework": 2 mẻ máy lớn (1 OK, 1 NG) + 1 mẻ Small Machine (phải bị LOẠI khỏi
        # KPI/tổng của tab này theo yêu cầu "chỉ tính máy >=500kg").
        conn.executemany(
            "INSERT INTO availability_logs (batch, capacity_kg, fabric_type, start_time, end_time) VALUES (?, ?, ?, ?, ?)",
            [
                ("RW-BIG-OK", 600, "Cotton", "2026-09-02 08:00:00", "2026-09-02 10:00:00"),
                ("RW-BIG-NG", 600, "Cotton", "2026-09-02 08:00:00", "2026-09-02 10:00:00"),
            ],
        )
        _insert_rft_row(conn, dyelot="RW-BIG-OK", stage="rework", result_dye="OK", machine_type=">=500kg")
        _insert_rft_row(conn, dyelot="RW-BIG-NG", stage="rework", result_dye="NG", machine_type=">=500kg")
        _insert_rft_row(conn, dyelot="RW-SMALL", stage="rework", result_dye="OK", machine_type="Small Machine")

        # 1 mẻ Stage lạ (không thuộc 6 nhóm) — phải đếm vào other_stage_count, KHÔNG rơi vào
        # tab nào. Cần khớp availability_logs CÙNG ngày với kịch bản Rework (2026-09-02) để
        # không bị loại bởi from_date/to_date filter của chính kịch bản đó.
        conn.execute(
            "INSERT INTO availability_logs (batch, capacity_kg, fabric_type, start_time, end_time) VALUES (?, ?, ?, ?, ?)",
            ("UNK-1", 600, "Cotton", "2026-09-02 08:00:00", "2026-09-02 10:00:00"),
        )
        _insert_rft_row(conn, dyelot="UNK-1", stage="mystery stage", result_dye="OK", machine_type=">=500kg")

        conn.commit()
        conn.close()

        app = _make_temp_app(db_path)
        with app.app_context():
            lab_to_bulk = get_rft_pivot_data(category="Lab to Bulk", from_date="2026-09-01", to_date="2026-09-01")
            rework_default = get_rft_pivot_data(category="Rework", from_date="2026-09-02", to_date="2026-09-02")
            rework_small_only = get_rft_pivot_data(category="Rework", machine_types="Small Machine", from_date="2026-09-02", to_date="2026-09-02")
            close_db()

        _check("Lab to Bulk: total_batches", lab_to_bulk["kpis"]["total_batches"], 3, failures)
        _check("Lab to Bulk: ok_batches", lab_to_bulk["kpis"]["ok_batches"], 2, failures)
        _check("Lab to Bulk: rate_pct (2/3)", lab_to_bulk["kpis"]["rate_pct"], 66.7, failures)

        _check("Rework (mặc định, không lọc Machine Type): total_batches CHỈ tính máy >=500kg (loại RW-SMALL)", rework_default["kpis"]["total_batches"], 2, failures)
        _check("Rework: ok_batches (chỉ RW-BIG-OK)", rework_default["kpis"]["ok_batches"], 1, failures)
        _check("Rework: rate_pct (1/2)", rework_default["kpis"]["rate_pct"], 50.0, failures)
        _check("Rework: other_stage_count đếm được UNK-1", rework_default["other_stage_count"], 1, failures)

        _check("Rework lọc Machine Type=Small Machine: vẫn 0 (bị business rule >=500kg loại trước)", rework_small_only["kpis"]["total_batches"], 0, failures)
    finally:
        os.unlink(db_path)


def _scenario_unknown_date_bucket(failures: list[str]) -> None:
    print("\n=== Kịch bản 4: mẻ không khớp availability_logs -> bucket 'Unknown Date' ===")
    fd, db_path = tempfile.mkstemp(suffix=".db")
    os.close(fd)
    try:
        _init_schema(db_path)
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        conn.execute(
            "INSERT INTO availability_logs (batch, capacity_kg, fabric_type, start_time, end_time) VALUES (?, ?, ?, ?, ?)",
            ("BB-KNOWN", 600, "Cotton", "2026-09-03 08:00:00", "2026-09-03 10:00:00"),
        )
        _insert_rft_row(conn, dyelot="BB-KNOWN", stage="bulk to bulk", result_dye="OK", machine_type=">=500kg")
        _insert_rft_row(conn, dyelot="BB-UNKNOWN", stage="bulk to bulk", result_dye="NG", machine_type="Small Machine")
        conn.commit()
        conn.close()

        app = _make_temp_app(db_path)
        with app.app_context():
            no_date_filter = get_rft_pivot_data(category="Bulk to Bulk")
            with_date_filter = get_rft_pivot_data(category="Bulk to Bulk", from_date="2026-09-03", to_date="2026-09-03")
            close_db()

        _check("Không lọc ngày: total_batches = 2 (giữ cả dòng không khớp availability_logs)", no_date_filter["kpis"]["total_batches"], 2, failures)
        _check("Không lọc ngày: có period 'Unknown Date'", "Unknown Date" in no_date_filter["periods"], True, failures)
        _check("Có lọc ngày tường minh: total_batches = 1 (loại dòng không xác định được ngày)", with_date_filter["kpis"]["total_batches"], 1, failures)
    finally:
        os.unlink(db_path)


def _scenario_sync_rft_results_end_to_end(failures: list[str]) -> None:
    """`sync_rft_results()` (import THẬT, UPSERT vào DB) trên file mẫu thật — DB tạm, KHÔNG
    đụng DB thật. Import 2 LẦN liên tiếp để verify UPSERT idempotent (không nhân đôi dòng)."""
    print("\n=== Kịch bản 5: sync_rft_results() end-to-end trên file mẫu thật (UPSERT idempotent) ===")
    if not FIXTURE_PATH.exists():
        print(f"  [SKIP] Không tìm thấy fixture {FIXTURE_PATH}")
        return
    fd, db_path = tempfile.mkstemp(suffix=".db")
    os.close(fd)
    try:
        _init_schema(db_path)
        file_bytes = FIXTURE_PATH.read_bytes()
        app = _make_temp_app(db_path)
        with app.app_context():
            result1 = sync_rft_results(file_bytes, imported_by="tester", filename="RFT report.xlsx")
            result2 = sync_rft_results(file_bytes, imported_by="tester", filename="RFT report.xlsx")
            close_db()

        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        total_rows = conn.execute("SELECT COUNT(*) AS c FROM rft_dye_results").fetchone()["c"]
        import_logs_count = conn.execute("SELECT COUNT(*) AS c FROM import_logs").fetchone()["c"]
        sample = conn.execute("SELECT result_dye, stage FROM rft_dye_results WHERE dyelot='C260612220'").fetchone()
        conn.close()

        _check("Lần 1: rows_imported = 31", result1["rows_imported"], 31, failures)
        _check("Lần 1: status = completed", result1["status"], "completed", failures)
        _check("Lần 2 (re-import CÙNG file): rows_imported vẫn = 31", result2["rows_imported"], 31, failures)
        _check("Sau 2 lần import: rft_dye_results KHÔNG nhân đôi (vẫn 31 dòng, UPSERT theo dyelot)", total_rows, 31, failures)
        _check("import_logs ghi đủ 2 lần import (audit trail)", import_logs_count, 2, failures)
        _check("Dòng mẫu C260612220 lưu đúng result_dye/stage", (sample["result_dye"], sample["stage"]), ("NG", "2nd batch"), failures)
    finally:
        os.unlink(db_path)


def _scenario_fabric_type_fallback_batch_details(failures: list[str]) -> None:
    """`fabric_type` PHẢI fallback sang `batch_details` (khoá dyelot=dyelot) khi
    `availability_logs` không khớp — xem docstring `_rft_rows()` (2026-09-19, sửa sau khi
    người dùng report báo cáo chia-3-loại-vải trống trơn vì file "RFT report.xlsx" không có
    cột Fabric Type và nhiều dyelot không khớp `availability_logs`)."""
    print("\n=== Kịch bản 6: fabric_type fallback sang batch_details khi availability_logs không khớp ===")
    fd, db_path = tempfile.mkstemp(suffix=".db")
    os.close(fd)
    try:
        _init_schema(db_path)
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row

        # FB-MATCH: khớp availability_logs (fabric_type=CVC) — batch_details CÓ dòng khác
        # (Polyester) nhưng KHÔNG được dùng, availability_logs vẫn ưu tiên trước.
        conn.execute(
            "INSERT INTO availability_logs (batch, capacity_kg, fabric_type, start_time, end_time) VALUES (?, ?, ?, ?, ?)",
            ("FB-MATCH", 600, "CVC", "2026-09-05 08:00:00", "2026-09-05 10:00:00"),
        )
        conn.execute("INSERT INTO batch_details (dyelot, fabric_type) VALUES ('FB-MATCH', 'Polyester')")
        # FB-FALLBACK: KHÔNG khớp availability_logs -> phải lấy fabric_type từ batch_details.
        conn.execute("INSERT INTO batch_details (dyelot, fabric_type) VALUES ('FB-FALLBACK', 'Cotton')")
        # FB-NONE: KHÔNG khớp cả 2 nguồn -> fabric_type rỗng, không rơi vào dòng vải nào.
        for dyelot in ("FB-MATCH", "FB-FALLBACK", "FB-NONE"):
            _insert_rft_row(conn, dyelot=dyelot, stage="lab to lab", result_dye="OK", machine_type=">=500kg")
        conn.commit()
        conn.close()

        app = _make_temp_app(db_path)
        with app.app_context():
            data = get_rft_pivot_data(category="Lab to Lab")
            close_db()

        rows_by_fabric = {row["fabric_type"]: row for row in data["rows"]}
        _check("KPI tổng vẫn đếm CẢ 3 dyelot (không đổi hành vi cũ)", data["kpis"]["total_batches"], 3, failures)
        _check("Cotton total = 1 (FB-FALLBACK, lấy fabric_type từ batch_details)", rows_by_fabric["Cotton"]["total"], 100.0, failures)
        _check("CVC total = 1 (FB-MATCH, ưu tiên availability_logs dù batch_details khác)", rows_by_fabric["CVC"]["total"], 100.0, failures)
        _check("Polyester total = 0 (FB-NONE không khớp nguồn nào, không bị gán nhầm)", rows_by_fabric["Polyester"]["total"], 0.0, failures)
    finally:
        os.unlink(db_path)


def _scenario_duplicate_dyelot_no_double_count(failures: list[str]) -> None:
    """1 Dyelot có 2 dòng `batch_details` (mẻ gốc + mẻ redye, xem điều tra C260659920 trong
    memory-bank/activeContext.md), KHÔNG khớp `availability_logs` nào — đúng population dễ bị
    lỗi nhân đôi nhất (fallback fabric_type). PHẢI vẫn đếm ĐÚNG 1 mẻ RFT (không nhân đôi thành
    2 dù JOIN theo dyelot khớp cả 2 dòng batch_details), và chọn fabric_type của dòng end_time
    MỚI NHẤT (fallback khi không có availability_logs để so khớp end_time chính xác)."""
    print("\n=== Kịch bản 7: Dyelot có 2 dòng batch_details (mẻ gốc + redye) — RFT KHÔNG đếm trùng ===")
    fd, db_path = tempfile.mkstemp(suffix=".db")
    os.close(fd)
    try:
        _init_schema(db_path)
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        conn.execute("INSERT INTO batch_details (dyelot, fabric_type, end_time) VALUES ('FB-DUP', 'Cotton', '2026-09-01 08:00:00')")
        conn.execute("INSERT INTO batch_details (dyelot, fabric_type, end_time) VALUES ('FB-DUP', 'CVC', '2026-09-08 06:00:00')")
        _insert_rft_row(conn, dyelot="FB-DUP", stage="lab to lab", result_dye="OK", machine_type=">=500kg")
        conn.commit()
        conn.close()

        app = _make_temp_app(db_path)
        with app.app_context():
            data = get_rft_pivot_data(category="Lab to Lab")
            close_db()

        rows_by_fabric = {row["fabric_type"]: row for row in data["rows"]}
        _check("KPI tổng vẫn đếm ĐÚNG 1 mẻ (KHÔNG nhân đôi thành 2)", data["kpis"]["total_batches"], 1, failures)
        _check("CVC total = 1 (chọn dòng end_time MỚI NHẤT khi không có availability_logs để so khớp)", rows_by_fabric.get("CVC", {}).get("total"), 100.0, failures)
        _check("Cotton total = 0 (dòng cũ hơn KHÔNG được chọn)", rows_by_fabric.get("Cotton", {}).get("total", 0.0), 0.0, failures)
    finally:
        os.unlink(db_path)


def main() -> int:
    failures: list[str] = []
    _scenario_stage_mapping(failures)
    _scenario_parse_real_fixture(failures)
    _scenario_kpi_formula_and_machine_restriction(failures)
    _scenario_unknown_date_bucket(failures)
    _scenario_sync_rft_results_end_to_end(failures)
    _scenario_fabric_type_fallback_batch_details(failures)
    _scenario_duplicate_dyelot_no_double_count(failures)
    print(f"\n{'='*60}\nKẾT QUẢ: {'TẤT CẢ KHỚP' if not failures else f'{len(failures)} CASE LỆCH'}\n{'='*60}")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())

"""
tests/test_knitting_downtime_import.py
-----------------------------------------
Engine `knitting.downtime` (khung sườn) — import file "CET-Stop Reason Analysis by Machine":

  1. Production Date đọc từ tên file (kể cả hậu tố "(1)" khi tải trùng); khoảng không trọn 1
     ngày sản xuất / tên file sai dạng bị từ chối.
  2. Parse file mẫu thật: 86 máy, 238 dòng dừng; số có dấu phẩy nghìn; Stop Color ARGB -> hex.
  3. Import vào DB tạm: đúng số dòng, import lại cùng ngày THAY THẾ (không nhân đôi, mã dừng
     biến mất ở bản sau cũng biến mất), import_logs ghi đúng.
  4. Báo cáo % Downtime trên file mẫu: mapping mặc định không còn mã "Unmapped", Total = tổng
     Stop Time của file, drill-down khớp ô, ghi đè nhóm / Target, Export Excel.
  5. Đối chiếu bảng của người dùng (2026-10-09): dựng lại pivot Stop Time 01–08/10 + Plan PRD,
     % theo ngày phải khớp bảng "Downtime (%)" theo ngày (Excel làm tròn 2 lần: 2 chữ số rồi 1 chữ
     số — giá trị gốc khớp tuyệt đối, chỉ khác hiển thị ở 7/112 ô), tuần W41 = 17.0%,
     nhãn "01-Oct" / "W41-Oct" / "Oct", Total Target = 16.3%.
  6. Bộ lọc Program (Greige ID của Piece Produced report quyết định Program):
     - file thật: Knitting program.xlsx (Core = 4 program, Greige trùng -> dòng cuối, gộp hoa/thường),
       Piece Produced tháng 7 (10915 cuộn, 42 máy), phủ 982 máy-ngày tháng 7, nhận diện loại file;
     - dữ liệu dựng tay: cuộn phủ nhiều ngày, nhiều Greige/ngày -> Greige nhiều phút nhất, Greige
       không có trong danh mục / máy không có cuộn -> Program TRỐNG (lọc qua "(Blank)"), Core only, lọc theo tên không
       phân biệt hoa/thường; mẫu số Available lọc cùng máy-ngày; drill-down + Excel có cột Program.

Chạy: python tests/test_knitting_downtime_import.py
"""
from __future__ import annotations

import io
import os
import sqlite3
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from modules.knitting.engines.downtime import programs, report  # noqa: E402
from modules.knitting.engines.excel_import import importer, program_importer, service  # noqa: E402

SAMPLE_DIR = Path(__file__).resolve().parent / "fixtures" / "sample_imports"
PROGRAM_FIXTURE = SAMPLE_DIR / "Copy of Knitting program.xlsx"
PIECE_FIXTURE = SAMPLE_DIR / "CET-Piece Produced report - 01-07-2026 07;00;00 - 01-08-2026 07;00;00.csv"
FIXTURE = (
    Path(__file__).resolve().parent / "fixtures" / "sample_imports"
    / "CET-Stop Reason Analysis by Machine - 14-09-2026 07;00;00 - 15-09-2026 07;00;00.csv"
)



# Pivot "Sum of Stop Time Details" 01–08/10 người dùng gửi (nhóm gốc của pivot -> mã/mô tả giả lập
# đi qua đúng mapping thật: mã đã chốt hoặc từ khoá mô tả).
PIVOT: dict[tuple[str, str], list[float]] = {
    ("BM", "Bad Material"): [8.28, 0, 0, 0, 0, 0, 3.82, 0],
    ("257", "Doffing"): [7322.3, 7188.23, 6588.74, 4408.84, 6984.86, 7926.41, 8690.66, 7984.1],
    ("9", "9 Un/LoadingYarn"): [699.24, 787.34, 545.15, 341.38, 880.73, 815.4, 963.07, 720.91],
    ("ADJ", "M/c Adjustment"): [21.88, 4.76, 54.88, 0, 61.21, 220.59, 637.14, 188.35],
    ("2", "2 M/c Part Broken"): [838.21, 445.13, 377.25, 225.62, 535.38, 528.6, 1137.8, 705.95],
    ("6", "6 M/c Set Up"): [2647.95, 2350.18, 3.58, 0, 2832.66, 1395.96, 0, 79.5],
    ("1", "1 Needle Broken"): [563.55, 500.17, 532.25, 355.47, 973, 785.05, 644.05, 745.98],
    ("12", "12 Needle cleaning"): [482.3, 506.32, 29.43, 0, 998.6, 1085.3, 946.66, 845.1],
    ("NM", "No Material"): [184.06, 2765.81, 5770.06, 2418.41, 4229.68, 1399.98, 571.66, 425.75],
    ("18", "18 Others"): [346.22, 284.93, 2588.3, 142.42, 1552.44, 2791.23, 4001.32, 9120.14],
    ("256", "Safe Door"): [1110.64, 795.98, 1817.19, 1460.65, 1726.17, 1186.1, 1196.92, 1490.53],
    ("3", "3 Yarn Broken"): [4706.31, 4399.28, 4579.79, 2208.32, 4230.84, 4919.26, 5035.74, 4437.44],
    ("CS", "Cleaning (Scheduled)"): [0, 0.66, 7551.22, 0, 0, 0, 0, 0],
    ("DS", "Drop Stitch"): [683.37, 690.58, 1168.75, 409.59, 808.56, 1287.95, 837.84, 777.12],
}
PLAN = [141923.65, 143983.63, 143943.85, 73697.45, 146596.6, 151680.65, 151047.55, 151715.35]
# Bảng "Downtime (%)" theo ngày người dùng gửi, cột 01-Oct .. 08-Oct.
EXPECTED_DAILY: dict[str, list[float]] = {
    "Doffing + Cleaning": [5.2, 5.0, 4.6, 6.0, 4.8, 5.2, 5.8, 5.3],
    "Yarn Broken": [3.3, 3.1, 3.2, 3.0, 2.9, 3.2, 3.3, 2.9],
    "No Material": [0.1, 1.9, 4.0, 3.3, 2.9, 0.9, 0.4, 0.3],
    "Loading/Unloading Yarn": [0.5, 0.6, 0.4, 0.5, 0.6, 0.5, 0.6, 0.5],
    "Needle Broken": [0.4, 0.4, 0.4, 0.5, 0.7, 0.5, 0.4, 0.5],
    "MC Part Broken": [0.6, 0.3, 0.3, 0.3, 0.4, 0.4, 0.8, 0.5],
    "MC Adjustment": [0.0, 0.0, 0.0, 0.0, 0.0, 0.2, 0.4, 0.1],
    "Needle Cleaning": [0.3, 0.4, 0.0, 0.0, 0.7, 0.7, 0.6, 0.6],
    "Safe Door": [0.8, 0.6, 1.3, 2.0, 1.2, 0.8, 0.8, 1.0],
    "MC Set Up": [1.9, 1.6, 0.0, 0.0, 1.9, 0.9, 0.0, 0.1],
    "Cleaning (Scheduled)": [0.0, 0.0, 5.3, 0.0, 0.0, 0.0, 0.0, 0.0],
    "Drop stitch": [0.5, 0.5, 0.8, 0.6, 0.6, 0.9, 0.6, 0.5],
    # 01-Oct: bảng cũ 0.2% chưa gộp Bad Material (8.28); quy tắc 2026-10-09 gộp Bad Material vào
    # Others -> (346.22 + 8.28) / 141923.65 = 0.2498% -> 0.25 -> 0.3 (khác biệt ĐÚNG theo quy tắc mới).
    "Others": [0.3, 0.2, 1.8, 0.2, 1.1, 1.8, 2.7, 6.0],
    "Total": [13.8, 14.4, 22.0, 16.2, 17.6, 16.1, 16.3, 18.1],
}

def _check(label: str, actual, expected, failures: list[str]) -> None:
    ok = actual == expected
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}: actual={actual!r} expected={expected!r}")
    if not ok:
        failures.append(label)


def _expect_error(label: str, func, failures: list[str]) -> None:
    try:
        func()
        _check(label, "accepted", "ValueError", failures)
    except ValueError:
        _check(label, "ValueError", "ValueError", failures)


def main() -> int:
    failures: list[str] = []
    os.environ.pop("DATABASE_URL", None)

    print("1. Production Date từ tên file")
    _check("file mẫu", importer.production_date_from_filename(FIXTURE.name)[0], "2026-09-14", failures)
    _check(
        "hậu tố (1)",
        importer.production_date_from_filename("CET-Stop Reason Analysis by Machine - 31-12-2026 07;00;00 - 01-01-2027 07;00;00 (1).csv")[0],
        "2026-12-31", failures,
    )
    _expect_error("không có ngày", lambda: importer.production_date_from_filename("stop.csv"), failures)
    _expect_error("1 ca (07->19)", lambda: importer.production_date_from_filename("X - 14-09-2026 07;00;00 - 14-09-2026 19;00;00.csv"), failures)
    _expect_error("không bắt đầu 07:00", lambda: importer.production_date_from_filename("X - 14-09-2026 06;00;00 - 15-09-2026 06;00;00.csv"), failures)

    print("2. Parse file mẫu")
    content = FIXTURE.read_bytes()
    parsed = importer.parse_stop_reason_file(content, FIXTURE.name)
    _check("số máy", len(parsed["machines"]), 86, failures)
    _check("số dòng dừng", len(parsed["stops"]), 238, failures)
    _check("không lỗi dòng", parsed["errors"], [], failures)
    ko21 = next(m for m in parsed["machines"] if m["machine_code"] == "KO0021")
    _check("KO0021 # Rev (dấu phẩy nghìn)", ko21["revolutions"], 1523.0, failures)
    _check("KO0021 Run", ko21["run_time"], 80.25, failures)
    stop = next(s for s in parsed["stops"] if s["machine_code"] == "KO0021" and s["stop_code"] == "259")
    _check("Stop Color int", stop["stop_color"], -65536, failures)
    _check("ARGB -> hex", importer.argb_to_hex(stop["stop_color"]), "#FF0000", failures)
    _check("ARGB -> hex (19 Ready to run)", importer.argb_to_hex(-16727872), "#00C0C0", failures)
    _expect_error("không phải file Stop Reason", lambda: importer.parse_stop_reason_file(b"a,b\n1,2\n", FIXTURE.name), failures)

    print("3. Import vào DB tạm")
    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    conn = sqlite3.connect(tmp.name, isolation_level=None)
    conn.row_factory = sqlite3.Row
    try:
        result = service.import_stop_reason_file(conn, content, FIXTURE.name, "tester")
        _check("status", result["status"], "completed", failures)
        _check("production_date", result["production_date"], "2026-09-14", failures)
        _check("lần đầu không thay thế", result["replaced_existing"], False, failures)
        count = lambda table: conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]  # noqa: E731
        _check("knitting_machine_daily", count("knitting_machine_daily"), 86, failures)
        _check("knitting_stop_details", count("knitting_stop_details"), 238, failures)
        _check("import_log_rows", count("import_log_rows"), 238, failures)

        # Bản xuất lại của cùng ngày, bỏ 1 dòng dừng -> phải thay thế, không nhân đôi.
        lines = content.decode("utf-8-sig").splitlines()
        trimmed = "\n".join(line for line in lines if not line.startswith("KO0021,") or ",259," not in line).encode("utf-8")
        again = service.import_stop_reason_file(conn, trimmed, FIXTURE.name, "tester")
        _check("lần 2 báo thay thế", again["replaced_existing"], True, failures)
        _check("máy sau thay thế", count("knitting_machine_daily"), 86, failures)
        _check("dòng dừng sau thay thế", count("knitting_stop_details"), 237, failures)
        _check(
            "mã dừng bị bỏ đã biến mất",
            conn.execute("SELECT COUNT(*) FROM knitting_stop_details WHERE machine_code='KO0021' AND stop_code='259'").fetchone()[0],
            0, failures,
        )
        _check("import_logs 2 lần", count("import_logs"), 2, failures)
        _check("danh sách ngày", [d["production_date"] for d in service.list_imported_days(conn)], ["2026-09-14"], failures)

        print("4. Báo cáo trên file mẫu")
        f = report.normalize_filters("2026-09-14", "2026-09-14", "date")
        rep = report.build_report(conn, f)
        stop_sum, avail_sum = conn.execute("SELECT (SELECT SUM(stop_time) FROM knitting_stop_details), (SELECT SUM(available_time) FROM knitting_machine_daily)").fetchone()
        _check("không có mã Unmapped", rep["unmapped_codes"], [], failures)
        _check("13 nhóm", [r["category"] for r in rep["rows"]], list(report.CATEGORIES), failures)
        _check("Total stop = tổng file", round(rep["total_row"]["total_stop_time"], 6), round(stop_sum, 6), failures)
        _check("Total % = Σstop/Σavailable", round(rep["total_row"]["pct"][0], 9), round(stop_sum / avail_sum * 100, 9), failures)
        _check("Target Total = 16.3", round(rep["total_row"]["target"], 6), 16.3, failures)
        _check("Before Total = 25.6", round(rep["total_row"]["before"], 6), 25.6, failures)
        others = next(r for r in rep["rows"] if r["category"] == "Others")
        expected_others = conn.execute("SELECT SUM(stop_time) FROM knitting_stop_details WHERE stop_code IN ('18','14','25','17','271')").fetchone()[0]
        _check("Others = 18+14+25+17+271", round(others["stop_time"][0], 6), round(expected_others, 6), failures)
        cell = report.get_cell_details(conn, f, "2026-09-14", "Doffing + Cleaning")
        doffing = next(r for r in rep["rows"] if r["category"] == "Doffing + Cleaning")
        _check("drill-down khớp ô", round(cell["stop_time"], 6), round(doffing["stop_time"][0], 6), failures)
        _check("drill-down Total (ALL) khớp", round(report.get_cell_details(conn, f, "ALL", "Total")["stop_time"], 6), round(stop_sum, 6), failures)
        f_ko = report.normalize_filters("2026-09-14", "2026-09-14", "date", machines="KO0021")
        _check("lọc M/c Code (KO0021 sau khi bỏ mã 259)", round(report.build_report(conn, f_ko)["total_row"]["pct"][0], 6), round((5.31 - 2.28) / 85.57 * 100, 6), failures)
        f_st = report.normalize_filters("2026-09-14", "2026-09-14", "date", structures="Interlock")
        n_inter = conn.execute("SELECT SUM(available_time) FROM knitting_machine_daily WHERE knitting_structure='Interlock'").fetchone()[0]
        _check("lọc Knitting Structure (mẫu số)", round(report.build_report(conn, f_st)["total_available"], 6), round(n_inter, 6), failures)

        report.set_stop_category(conn, "271", "MC Set Up", "tester")
        moved = report.build_report(conn, f)
        _check("ghi đè nhóm áp ngay (không import lại)", round(next(r for r in moved["rows"] if r["category"] == "Others")["stop_time"][0], 6), round(expected_others - 0.35, 6), failures)
        _check("nguồn manual", next(c for c in report.list_stop_codes(conn) if c["stop_code"] == "271")["source"], "manual", failures)
        report.set_stop_category(conn, "271", None, "tester")
        _check("bỏ ghi đè -> default", next(c for c in report.list_stop_codes(conn) if c["stop_code"] == "271")["source"], "default", failures)
        _expect_error("nhóm sai", lambda: report.set_stop_category(conn, "271", "Abc", "tester"), failures)
        report.set_target(conn, "Safe Door", 5.0, 1.2, "tester")
        _check("sửa Target", report.load_targets(conn)["Safe Door"]["target"], 1.2, failures)
        report.set_target(conn, "Safe Door", 5.0, 1.5, "tester")
        _expect_error("From > To", lambda: report.normalize_filters("2026-09-15", "2026-09-14", "date"), failures)

        wb = load_workbook(io.BytesIO(report.export_excel(conn, f)))
        _check("Excel sheets", wb.sheetnames, ["Downtime %", "Stop Time", "Achievement", "Data", "Filters"], failures)
        _check("Excel Data rows (sau thay thế)", wb["Data"].max_row - 1, 237, failures)
        _check("Excel Total %", round(wb["Downtime %"].cell(row=15, column=4).value, 9), round(stop_sum / avail_sum, 9), failures)
    finally:
        conn.close()
        os.unlink(tmp.name)

    print("5. Đối chiếu bảng của người dùng (01–08/10)")
    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    conn = sqlite3.connect(tmp.name, isolation_level=None)
    conn.row_factory = sqlite3.Row
    try:
        service.ensure_tables(conn)
        days = [f"2026-10-0{d}" for d in range(1, 9)]
        for day, available in zip(days, PLAN):
            conn.execute(
                "INSERT INTO knitting_machine_daily (production_date, machine_code, knitting_structure, available_time, period_start, period_end) VALUES (?, 'M1', 'Interlock', ?, '', '')",
                (day, available),
            )
        for (code, description), values in PIVOT.items():
            for day, value in zip(days, values):
                if value:
                    conn.execute(
                        "INSERT INTO knitting_stop_details (production_date, machine_code, stop_code, stop_description, stop_time) VALUES (?, 'M1', ?, ?, ?)",
                        (day, code, description, value),
                    )
        conn.commit()
        daily = report.build_report(conn, report.normalize_filters("2026-10-01", "2026-10-08", "date"))
        _check("nhãn ngày", daily["periods"][:2], ["01-Oct", "02-Oct"], failures)
        _check("không Unmapped", daily["unmapped_codes"], [], failures)
        mismatches = []
        for row in daily["rows"] + [daily["total_row"]]:
            expected = EXPECTED_DAILY[row["category"]]
            # Excel người dùng làm tròn 2 LẦN (bảng % Dt 2 chữ số -> hiển thị 1 chữ số): 0.547 -> 0.55 -> 0.6.
            actual = [round(round(v, 2) + 1e-9, 1) for v in row["pct"]]
            if actual != expected:
                mismatches.append((row["category"], actual, expected))
        _check("13 nhóm + Total x 8 ngày khớp bảng người dùng (làm tròn kiểu Excel)", mismatches, [], failures)
        weekly = report.build_report(conn, report.normalize_filters("2026-10-05", "2026-10-08", "week"))
        _check("nhãn tuần", weekly["periods"], ["W41-Oct"], failures)
        _check("W41 Total = 17.0%", round(weekly["total_row"]["pct"][0], 1), 17.0, failures)
        monthly = report.build_report(conn, report.normalize_filters("2026-10-01", "2026-10-08", "month"))
        _check("nhãn tháng", monthly["periods"], ["Oct"], failures)
        _check("tháng = Σstop/Σplan (KHÔNG trung bình %)", round(monthly["total_row"]["pct"][0], 2), 16.86, failures)
    finally:
        conn.close()
        os.unlink(tmp.name)

    print("6. Bộ lọc Program")
    _check("nhận diện Stop Reason", program_importer.detect_file_type(FIXTURE.name, FIXTURE.read_bytes()), importer.FILE_TYPE, failures)
    _check("nhận diện Piece Produced", program_importer.detect_file_type(PIECE_FIXTURE.name, PIECE_FIXTURE.read_bytes()), program_importer.PIECE_FILE_TYPE, failures)
    _check("nhận diện Knitting program", program_importer.detect_file_type(PROGRAM_FIXTURE.name, PROGRAM_FIXTURE.read_bytes()), program_importer.PROGRAM_FILE_TYPE, failures)
    _expect_error("CSV lạ", lambda: program_importer.detect_file_type("x.csv", b"a,b\n1,2"), failures)

    prog = program_importer.parse_program_file(PROGRAM_FIXTURE.read_bytes())
    _check("Core program (4, bỏ cột H/I)", [c["program"] for c in prog["core_programs"]],
           ["Washed Boxy T-Shirt", "Graphic Tee-20S/1", "Interlock-Body", "Airism (50S/1 165G)"], failures)
    by_greige = {r["greige_code"]: r for r in prog["greige_programs"]}
    _check("số Greige code", len(by_greige), 224, failures)
    _check("Greige trùng -> dòng cuối (New Balance/Kuhl)", by_greige["9V00-0000240"]["program"] in {"New Balance", "Kuhl"}, True, failures)
    _check("gộp hoa/thường GEL/Gel", by_greige["9V00-0001151"]["program_key"], "gel", failures)
    _check("có cảnh báo Greige nhiều Program khác hẳn", any(c["greige_code"] == "9V00-0000240" for c in prog["conflicts"]), True, failures)
    _expect_error("xlsx không phải Knitting program", lambda: program_importer.parse_program_file(b"not an xlsx"), failures)

    piece = program_importer.parse_piece_produced_file(PIECE_FIXTURE.read_bytes())
    _check("Piece Produced: số cuộn", len(piece["rolls"]), 10915, failures)
    _check("Piece Produced: số máy", piece["machines"], 42, failures)
    _check("Piece Produced: không lỗi", piece["errors"], [], failures)
    _check("Record End range", (piece["date_from"], piece["date_to"]), ("2026-07-01", "2026-08-01"), failures)

    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    conn = sqlite3.connect(tmp.name, isolation_level=None)
    conn.row_factory = sqlite3.Row
    try:
        res = service.import_program_file(conn, PROGRAM_FIXTURE.read_bytes(), PROGRAM_FIXTURE.name, "tester")
        _check("import program: Greige", res["greige_codes"], 224, failures)
        res = service.import_piece_produced_file(conn, PIECE_FIXTURE.read_bytes(), PIECE_FIXTURE.name, "tester")
        _check("import Piece Produced: cuộn", res["rolls"], 10915, failures)
        _check("Greige không có trong danh mục", res["greige_not_in_program_list"], ["9VDK250145N2"], failures)
        again = service.import_piece_produced_file(conn, PIECE_FIXTURE.read_bytes(), PIECE_FIXTURE.name, "tester")
        _check("import lại không nhân đôi (UPSERT Roll No)", conn.execute("SELECT COUNT(*) FROM knitting_piece_rolls").fetchone()[0], 10915, failures)
        _check("import lại status", again["status"], "completed", failures)
        july = programs.machine_day_programs(conn, "2026-07-01", "2026-07-31")
        _check("phủ máy-ngày tháng 7 (Start->End)", len(july), 982, failures)
        sources = service.get_program_sources(conn)
        _check("sources: core", sources["core_programs"], sorted(c["program"] for c in prog["core_programs"]), failures)
        options = programs.list_program_options(conn)
        _check("options có (Blank)", options["programs"][-1], programs.BLANK_OPTION, failures)
    finally:
        conn.close()
        os.unlink(tmp.name)

    tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    tmp.close()
    conn = sqlite3.connect(tmp.name, isolation_level=None)
    conn.row_factory = sqlite3.Row
    try:
        service.ensure_tables(conn)
        conn.executemany("INSERT INTO knitting_greige_programs (greige_code, program, program_key) VALUES (?, ?, ?)",
                         [("G1", "Interlock-Body", "interlock-body"), ("G2", "Gel", "gel")])
        conn.execute("INSERT INTO knitting_core_programs (program_key, program) VALUES ('interlock-body', 'Interlock-Body')")
        rolls = [  # roll_no, machine, greige, available, start, end
            ("R1", "M1", "G1", 2000, "2026-07-01", "2026-07-02"),   # phủ 2 ngày
            ("R2", "M2", "G2", 1000, "2026-07-01", "2026-07-01"),
            ("R3", "M2", "G1", 300, "2026-07-01", "2026-07-01"),    # M2 ngày 01: G2 nhiều phút hơn
            ("R4", "M3", "G9", 500, "2026-07-01", "2026-07-01"),    # Greige không có trong danh mục
        ]
        conn.executemany(
            "INSERT INTO knitting_piece_rolls (roll_no, machine_code, greige_id, available, record_start, record_end) VALUES (?, ?, ?, ?, ?, ?)",
            rolls,
        )
        for day in ("2026-07-01", "2026-07-02"):
            for machine, available in (("M1", 100), ("M2", 200), ("M3", 400), ("M4", 800)):
                conn.execute(
                    "INSERT INTO knitting_machine_daily (production_date, machine_code, available_time, period_start, period_end) VALUES (?, ?, ?, '', '')",
                    (day, machine, available),
                )
                conn.execute(
                    "INSERT INTO knitting_stop_details (production_date, machine_code, stop_code, stop_description, stop_time) VALUES (?, ?, '257', 'Doffing', 10)",
                    (day, machine),
                )
        conn.commit()

        def total(**kw):
            return report.build_report(conn, report.normalize_filters("2026-07-01", "2026-07-01", "date", **kw))["total_row"]["pct"][0]

        mapping = programs.machine_day_programs(conn, "2026-07-01", "2026-07-02")
        _check("M1 ngày 02 (cuộn phủ 2 ngày)", programs.program_of(mapping, "2026-07-02", "M1")["program"], "Interlock-Body", failures)
        _check("M2 nhiều Greige -> nhiều phút nhất", programs.program_of(mapping, "2026-07-01", "M2")["program"], "Gel", failures)
        _check("M3 Greige ngoài danh mục -> trống", programs.program_of(mapping, "2026-07-01", "M3")["program"], "", failures)
        _check("M3 vẫn giữ Greige ID", programs.program_of(mapping, "2026-07-01", "M3")["greige_id"], "G9", failures)
        _check("M4 không có cuộn -> trống", programs.program_of(mapping, "2026-07-01", "M4")["program"], "", failures)
        _check("không lọc: 40/1500", round(total(), 6), round(40 / 1500 * 100, 6), failures)
        _check("Core only: chỉ M1 (10/100)", round(total(core_only="1"), 6), 10.0, failures)
        _check("Program 'GEL' (hoa/thường): chỉ M2 (10/200)", round(total(programs="GEL"), 6), 5.0, failures)
        _check("(Blank): M3 + M4 (20/1200)", round(total(programs="(Blank)"), 6), round(20 / 1200 * 100, 6), failures)
        _check("Gel + (Blank): M2 + M3 + M4 (30/1400)", round(total(programs="Gel|(Blank)"), 6), round(30 / 1400 * 100, 6), failures)
        _check("Core only + Gel: không còn máy nào", report.build_report(conn, report.normalize_filters("2026-07-01", "2026-07-01", "date", programs="Gel", core_only=True))["period_keys"], [], failures)
        f = report.normalize_filters("2026-07-01", "2026-07-02", "date", core_only=True)
        cell = report.get_cell_details(conn, f, "ALL", "Total")
        _check("drill-down Core only: chỉ M1", sorted({r["machine_code"] for r in cell["rows"]}), ["M1"], failures)
        _check("drill-down mẫu số", cell["available"], 200, failures)
        _check("drill-down có Program", cell["rows"][0]["program"], "Interlock-Body", failures)
        wb = load_workbook(io.BytesIO(report.export_excel(conn, f)))
        data_rows = list(wb["Data"].iter_rows(values_only=True))
        _check("Excel Data có cột Program", data_rows[0][4], "Program", failures)
        _check("Excel Data lọc Core only", sorted({r[1] for r in data_rows[1:]}), ["M1"], failures)

        print("7. Standard Achievement (% Downtime máy-ngày <= Target) + chi tiết theo ngày")
        # Mỗi máy-ngày dừng Doffing 10; Target Doffing + Cleaning 5.5%:
        # M1 10/100 = 10% (KHÔNG đạt), M2 5%, M3 2.5%, M4 1.25% (đạt) -> 3/4 = 75% mỗi ngày.
        rep7 = report.build_report(conn, report.normalize_filters("2026-07-01", "2026-07-02", "date"))
        doff = next(r for r in rep7["achievement"]["rows"] if r["category"] == "Doffing + Cleaning")
        _check("Achievement Doffing theo ngày", doff["values"], [75.0, 75.0], failures)
        _check("Achievement Doffing tổng", (doff["total_evaluated"], doff["total_passed"], doff["rate_pct"]), (8, 6, 75.0), failures)
        safe = next(r for r in rep7["achievement"]["rows"] if r["category"] == "Safe Door")
        _check("nhóm không dừng -> 0% <= Target -> đạt", safe["rate_pct"], 100.0, failures)
        _check("Total so Target 16.3% (M1 10% vẫn đạt)", rep7["achievement"]["total_row"]["rate_pct"], 100.0, failures)
        _check("KPI achievement = dòng Total", rep7["kpis"]["achievement_pct"], 100.0, failures)
        _check("KPI downtime %", round(rep7["kpis"]["downtime_pct"], 6), round(80 / 3000 * 100, 6), failures)
        cell7 = report.get_cell_details(conn, report.normalize_filters("2026-07-01", "2026-07-02", "week"), "ALL", "Doffing + Cleaning")
        _check("drill: 2 dòng ngày", [d["production_date"] for d in cell7["daily"]], ["2026-07-01", "2026-07-02"], failures)
        _check("drill: ngày 01 Plan / Stop / %", (cell7["daily"][0]["available"], cell7["daily"][0]["stop_time"], round(cell7["daily"][0]["pct"], 6)), (1500, 40, round(40 / 1500 * 100, 6)), failures)
        _check("drill: ngày 01 đạt 3/4", (cell7["daily"][0]["evaluated"], cell7["daily"][0]["passed"]), (4, 3), failures)
        _check("drill: máy-ngày không đạt", [(m["production_date"], m["machine_code"]) for m in cell7["machine_day_rows"] if m["achieved"] is False],
               [("2026-07-01", "M1"), ("2026-07-02", "M1")], failures)
        _check("drill: tổng achievement", cell7["achievement_pct"], 75.0, failures)
        _check("drill: Σ ngày = Σ dòng dừng", sum(d["stop_time"] for d in cell7["daily"]), sum(r["stop_time"] for r in cell7["rows"]), failures)
        wb7 = load_workbook(io.BytesIO(report.export_excel(conn, report.normalize_filters("2026-07-01", "2026-07-02", "date"))))
        ach = {r[0]: r for r in wb7["Achievement"].iter_rows(min_row=2, values_only=True)}
        _check("Excel Achievement Doffing", ach["Doffing + Cleaning"][2:4], (0.75, 0.75), failures)
    finally:
        conn.close()
        os.unlink(tmp.name)

    print()
    print("ALL PASS" if not failures else f"FAILED: {failures}")
    return 0 if not failures else 1


if __name__ == "__main__":
    sys.exit(main())

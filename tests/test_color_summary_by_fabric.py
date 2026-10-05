"""
tests/test_color_summary_by_fabric.py
---------------------------------------
2026-10-05 — bảng "Normal Dyeing Batches by Colour (per day)" của Batch Per Day by Machine tách
thêm theo loại vải Cotton -> CVC -> Polyester (`color_summary.fabric_groups`). Kiểm tra:

  1. Số theo (loại vải, màu, ngày) đúng; tổng các nhóm = `by_day` tổng = KPI "Normal Dyeing",
     ở CẢ 2 trạng thái checkbox ReDye = 0.
  2. Mẻ FabricType "Unknown" (mẻ CM cuối máy, không có mẻ kế tiếp để fill-up) KHÔNG vào bảng.
  3. Lọc Fabric Type = Cotton -> chỉ còn nhóm Cotton.
  4. Excel export sheet "Color Summary": cột Fabric merge dọc, dòng Subtotal mỗi nhóm, dòng
     Total cuối khớp KPI.

Chạy: python tests/test_color_summary_by_fabric.py
"""
from __future__ import annotations

import io
import os
import shutil
import sqlite3
import sys
import tempfile
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from flask import Flask  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from core.database import close_db  # noqa: E402
from modules.dyeing.engines.reports import cleaning_matrix  # noqa: E402

# (dyelot, machine, fabric_type, sap_lot, redye, shade, start, end)
BATCHES = [
    ("C100000", "M1", "Cotton", "1000000001", 0, "Dark", "2026-09-01 08:00:00", "2026-09-01 10:00:00"),
    ("C100010", "M1", "Cotton", "1000000002", 0, "Light", "2026-09-01 11:00:00", "2026-09-01 13:00:00"),
    ("C100020", "M2", "CVC", "3000000003", 0, "Dark", "2026-09-01 08:00:00", "2026-09-01 10:00:00"),
    # Polyester ReDye = 1 -> Rework khi checkbox BẬT, Normal khi TẮT.
    ("C100030", "M2", "Polyester", "1000000004", 1, "Dark", "2026-09-01 11:00:00", "2026-09-01 13:00:00"),
    ("C100040", "M2", "Polyester", "1000000005", 0, "Light", "2026-09-02 08:00:00", "2026-09-02 10:00:00"),
    # Mẻ CM, FabricType trống và là mẻ cuối của M3 -> "Unknown"; không được vào bảng.
    ("C100050-WA", "M3", "", "", 0, "", "2026-09-02 08:00:00", "2026-09-02 09:00:00"),
]
DAYS = ["2026-09-01", "2026-09-02"]


def _check(label: str, actual, expected, failures: list[str]) -> None:
    ok = actual == expected
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}: actual={actual!r} expected={expected!r}")
    if not ok:
        failures.append(label)


def _init_db(db_path: str) -> None:
    conn = sqlite3.connect(db_path)
    conn.execute("""
        CREATE TABLE batch_details (
            id INTEGER PRIMARY KEY AUTOINCREMENT, dyelot TEXT NOT NULL, sap_lot TEXT, redye REAL DEFAULT 0,
            shade TEXT, colour_no TEXT, recipe_no TEXT, customer_color TEXT, batch_type TEXT,
            is_rework INTEGER DEFAULT 0, run_time REAL DEFAULT 0, greige_code TEXT,
            machine TEXT, fabric_type TEXT, start_time TEXT, end_time TEXT
        )
    """)
    conn.execute("CREATE TABLE machines (machine_id TEXT, machine_code TEXT, mc_brand TEXT, tank_type TEXT, mc_quantity INTEGER, tube_no INTEGER, capacity_kg REAL, domain TEXT)")
    conn.executemany(
        "INSERT INTO machines (machine_id, machine_code, capacity_kg, domain) VALUES (?, ?, 600, 'dyeing')",
        [("M1", "M1"), ("M2", "M2"), ("M3", "M3")],
    )
    conn.executemany(
        "INSERT INTO batch_details (dyelot, machine, fabric_type, sap_lot, redye, shade, start_time, end_time) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        BATCHES,
    )
    conn.commit()
    conn.close()


def _groups(data: dict) -> dict[str, dict[str, list[int]]]:
    return {group["fabric"]: group["by_day"] for group in data["color_summary"]["fabric_groups"]}


def _scenario_counts(failures: list[str]) -> None:
    # labels = (Dark, Light, Medium, Black, White)
    for require_redye_zero, poly_day1 in ((True, [0, 0, 0, 0, 0]), (False, [1, 0, 0, 0, 0])):
        print(f"\n=== Kịch bản 1: require_redye_zero={require_redye_zero} ===")
        data = cleaning_matrix.get_cleaning_matrix(DAYS[0], DAYS[1], require_redye_zero=require_redye_zero)
        summary = data["color_summary"]
        groups = _groups(data)
        _check("thứ tự nhóm Cotton -> CVC -> Polyester, không có Unknown", list(groups), ["Cotton", "CVC", "Polyester"], failures)
        _check("Cotton 09-01 = 1 Dark + 1 Light", groups["Cotton"]["2026-09-01"], [1, 1, 0, 0, 0], failures)
        _check("CVC 09-01 = 1 Dark", groups["CVC"]["2026-09-01"], [1, 0, 0, 0, 0], failures)
        _check("Polyester 09-01 theo checkbox ReDye", groups["Polyester"]["2026-09-01"], poly_day1, failures)
        _check("Polyester 09-02 = 1 Light (mẻ CM Unknown không tính)", groups["Polyester"]["2026-09-02"], [0, 1, 0, 0, 0], failures)
        for day in DAYS:
            per_group = [sum(values) for values in zip(*(groups[fabric][day] for fabric in groups))]
            _check(f"{day}: tổng các nhóm = by_day tổng", per_group, summary["by_day"][day], failures)
        grand = sum(sum(counts) for counts in summary["by_day"].values())
        _check("tổng bảng = KPI Normal Dyeing", grand, data["kpis"]["normal_batches"], failures)


def _scenario_filter(failures: list[str]) -> None:
    print("\n=== Kịch bản 2: lọc Fabric Type = Cotton ===")
    data = cleaning_matrix.get_cleaning_matrix(DAYS[0], DAYS[1], fabric_types=["Cotton"])
    _check("chỉ còn nhóm Cotton", list(_groups(data)), ["Cotton"], failures)


def _scenario_excel(failures: list[str]) -> None:
    print("\n=== Kịch bản 3: Excel export sheet 'Color Summary' ===")
    content = cleaning_matrix.export_cleaning_matrix_excel(DAYS[0], DAYS[1], require_redye_zero=False)
    sheet = load_workbook(io.BytesIO(content))["Color Summary"]
    rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    _check("header", rows[0], ["Fabric", "Colour", *DAYS, "Total"], failures)
    # 3 nhóm x (5 màu + Subtotal) + header + Total = 20 dòng.
    _check("số dòng", len(rows), 1 + 3 * 6 + 1, failures)
    _check("cột A mỗi nhóm", [rows[1][0], rows[7][0], rows[13][0]], ["Cotton", "CVC", "Polyester"], failures)
    _check("Cotton Subtotal", rows[6][1:], ["Subtotal", 2, 0, 2], failures)
    _check("Polyester Subtotal (ReDye tắt)", rows[18][1:], ["Subtotal", 1, 1, 2], failures)
    _check("dòng Total", [rows[19][0], *rows[19][2:]], ["Total", 4, 1, 5], failures)
    merged = sorted(str(rng) for rng in sheet.merged_cells.ranges)
    _check("cột Fabric merge dọc theo nhóm", merged, ["A14:A19", "A2:A7", "A8:A13"], failures)


def main() -> int:
    failures: list[str] = []
    tmp_dir = tempfile.mkdtemp()
    try:
        db_path = os.path.join(tmp_dir, "test.db")
        _init_db(db_path)
        conn = sqlite3.connect(db_path)
        conn.row_factory = sqlite3.Row
        cleaning_matrix.recompute_all([date(2026, 9, 1), date(2026, 9, 2)], conn)
        conn.commit()
        unknown = [tuple(r) for r in conn.execute("SELECT dyelot_ref, badge FROM cleaning_mc_daily_summary WHERE fabric_type = 'Unknown'")]
        conn.close()
        _check("fixture: mẻ CM cuối M3 có FabricType Unknown", unknown, [("C100050-WA", "CM")], failures)

        app = Flask(__name__)
        app.config["DATABASE_PATH"] = db_path
        app.config["SQLITE_PRAGMAS"] = {}
        with app.app_context():
            _scenario_counts(failures)
            _scenario_filter(failures)
            _scenario_excel(failures)
            close_db()
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    print(f"\n{'=' * 60}\nKẾT QUẢ: {'TẤT CẢ KHỚP' if not failures else f'{len(failures)} CASE LỆCH'}\n{'=' * 60}")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
